"""Core pipeline regression tests.

Covers two real bugs found and fixed in this codebase (a label collision
across dimensions corrupting Gate 1's share-sum check, and the ME writer
silently leaving a mismatched template's stale data in the "generated"
output) plus the region/Global rollup math the Assembler performs.

Run with: pytest tests/
"""
from __future__ import annotations

import pytest
from openpyxl import Workbook, load_workbook

from me_engine.assembly.assembler import Assembler
from me_engine.domain.drivers import AspDrivers, DriverSet, GeographyDrivers, SegmentationDrivers
from me_engine.domain.runtime_taxonomy import RuntimeTaxonomy, build_taxonomy_from_workbook
from me_engine.domain.series import Series
from me_engine.domain.taxonomy import YEARS, Band, Dimension, GeographyIndex, _g
from me_engine.generation.sheet_writer import write_input_sheet
from me_engine.generation.skeleton import Skeleton
from me_engine.io.me_builder import MEWorkbookBuilder
from me_engine.io.me_writer import MEWorkbookWriter, TemplateMismatchError
from me_engine.validation.gate1 import InputValidator


def _flat_series(base: float, growth: float) -> Series:
    values = {}
    v = base
    for i, year in enumerate(YEARS):
        values[year] = base if i == 0 else values[YEARS[i - 1]] * (1 + growth)
    return Series(values)


# --------------------------------------------------------------------------- #
# Regression: a segment label shared across dimensions (e.g. "Others") must
# not corrupt another dimension's share-sum. Fixed in gate1.py / workbook_inspector.py.
# --------------------------------------------------------------------------- #
def test_gate1_shares_sum_survives_shared_label_across_dimensions(tmp_path):
    skeleton = Skeleton(
        market_name="Test Market",
        geographies=["U.S.", "Canada"],
        regions={"North America": ["U.S.", "Canada"]},
        nesting={
            "By Product Type": {"__root__": ["A", "B", "Others"]},
            "By Application": {"__root__": ["X", "Y", "Others"]},
        },
        priced_dimension="By Product Type",
    )
    state = {
        "skeleton": skeleton,
        "shares": {
            "By Product Type": {"__global__": {"__root__": {"A": 0.3, "B": 0.5, "Others": 0.2}}},
            "By Application": {"__global__": {"__root__": {"X": 0.1, "Y": 0.3, "Others": 0.6}}},
        },
        "asp": {"A": 10.0, "B": 20.0, "Others": 5.0},
        "regional": {"North America": {"value_share": 1.0, "cagr": 0.05}},
        "geo": {"North America": {
            "U.S.": {"share_of_region": 0.7, "cagr": 0.05},
            "Canada": {"share_of_region": 0.3, "cagr": 0.05},
        }},
        "global_anchor_2025": 1000.0,
        "global_cagr": 0.05,
    }
    out_path = write_input_sheet(state, tmp_path / "input.xlsx", flat=True)

    taxonomy = build_taxonomy_from_workbook(out_path)
    report = InputValidator(taxonomy).validate(out_path)

    assert report.passed(), str(report)
    assert not any(f.check == "shares-sum" for f in report.failures)


# --------------------------------------------------------------------------- #
# Shared fixture-style helper for the Assembler/writer tests below.
# --------------------------------------------------------------------------- #
def _build_taxonomy_and_result():
    dim = Dimension(title="By Type", segments=("A", "B"))
    geographies = GeographyIndex(_g("Global", _g("North America", _g("U.S."), _g("Canada"))))
    taxonomy = RuntimeTaxonomy(
        market_name="Test Market",
        geographies=geographies,
        segmentation_dimensions=(dim,),
        priced_dimension=dim,
    )

    def geo_drivers(name: str, base_value: float, share_a: float) -> GeographyDrivers:
        return GeographyDrivers(
            name=name,
            value=_flat_series(base_value, 0.05),
            segmentation=SegmentationDrivers(shares={
                "By Type": {
                    "A": _flat_series(share_a, 0.0),
                    "B": _flat_series(1 - share_a, 0.0),
                },
            }),
            asp=AspDrivers(asp={"A": _flat_series(10.0, 0.01), "B": _flat_series(20.0, 0.01)}),
        )

    drivers = DriverSet(market_name="Test Market", geographies={
        "U.S.": geo_drivers("U.S.", 100.0, 0.6),
        "Canada": geo_drivers("Canada", 50.0, 0.5),
    })
    result = Assembler(taxonomy).assemble(drivers)
    return taxonomy, result


def test_assembler_region_rollup_is_additive_and_asp_is_derived():
    taxonomy, result = _build_taxonomy_and_result()

    assert set(result.geographies) == {"U.S.", "Canada", "North America", "Global"}

    us = result.geographies["U.S."].bands[Band.VALUE].total
    ca = result.geographies["Canada"].bands[Band.VALUE].total
    na = result.geographies["North America"].bands[Band.VALUE].total
    glob = result.geographies["Global"].bands[Band.VALUE].total

    for year in YEARS:
        assert na.at(year) == pytest.approx(us.at(year) + ca.at(year))
        assert glob.at(year) == pytest.approx(na.at(year))

    # ASP is never averaged/summed — it must equal value/volume*1000 exactly,
    # the same identity used at the leaf level.
    na_value_a = result.geographies["North America"].bands[Band.VALUE].rows_by_label["A"].series
    na_volume_a = result.geographies["North America"].bands[Band.VOLUME].rows_by_label["A"].series
    na_asp_a = result.geographies["North America"].bands[Band.ASP].rows_by_label["A"].series
    for year in YEARS:
        expected = na_value_a.at(year) / na_volume_a.at(year) * 1000.0
        assert na_asp_a.at(year) == pytest.approx(expected)


def test_me_builder_writes_every_geography_level_without_a_template(tmp_path):
    taxonomy, result = _build_taxonomy_and_result()
    out = MEWorkbookBuilder(taxonomy).build(result, tmp_path / "me.xlsx")

    wb = load_workbook(out, data_only=True)
    assert set(wb.sheetnames) == {"Global", "North America", "U.S.", "Canada"}
    ws = wb["North America"]
    # Value & CAGR header row: label in col C, value in col D.
    assert ws.cell(1, 3).value == "Value & CAGR"
    assert ws.cell(1, 4).value == pytest.approx(
        result.geographies["North America"].bands[Band.VALUE].total.anchor
    )


def test_me_writer_rejects_a_mismatched_template_instead_of_copying_it(tmp_path):
    _taxonomy, result = _build_taxonomy_and_result()

    unrelated = Workbook()
    unrelated.active.title = "Nothing To Do With This Market"
    template_path = tmp_path / "unrelated_template.xlsx"
    unrelated.save(template_path)

    with pytest.raises(TemplateMismatchError):
        MEWorkbookWriter(template_path).write(result, tmp_path / "out.xlsx")
