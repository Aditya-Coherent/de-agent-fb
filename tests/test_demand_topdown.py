"""Shared top-down segmentation layer + approach 2A/2B segmentation phases.

All offline: a fake LLM returns canned JSON and web search is stubbed.
"""
from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook

from me_engine.demand_topdown import regions, segments, split
from me_engine.demand_topdown.writer import format_terminal_summary


class FakeLLM:
    """Answers the two prompts segments.py sends: dimension proposal and shares."""

    def __init__(self, dimensions, shares_by_dim, evidence=None):
        self.dimensions = dimensions
        self.shares_by_dim = shares_by_dim
        self.evidence = evidence or []
        self.calls = []

    def complete_json(self, system, user, *, retries=3):
        self.calls.append(user)
        if "Propose the 2-4 segmentation dimensions" in user:
            return {"dimensions": self.dimensions}
        for dim, shares in self.shares_by_dim.items():
            if f"Dimension: {dim}\n" in user:
                return {"shares": shares, "evidence": self.evidence, "reasoning": "test"}
        raise AssertionError(f"unexpected prompt: {user[:120]}")


class NoDbConfig:
    neon_url = None


# --------------------------------------------------------------------------- #
# regions
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("country,region", [
    ("United States", regions.NORTH_AMERICA),
    ("U.S.", regions.NORTH_AMERICA),
    ("Brazil", regions.LATIN_AMERICA),
    ("Germany", regions.EUROPE),
    ("Rest of Europe", regions.EUROPE),
    ("India", regions.ASIA_PACIFIC),
    ("Rest of Asia Pacific", regions.ASIA_PACIFIC),
    ("Saudi Arabia", regions.MIDDLE_EAST_AFRICA),
    ("Rest of Middle East & Africa", regions.MIDDLE_EAST_AFRICA),
    ("Atlantis", regions.REST_OF_WORLD),
])
def test_region_for(country, region):
    assert regions.region_for(country) == region


def test_rollup_regions_keeps_every_country_and_orders_regions():
    totals = {"India": 30.0, "United States": 50.0, "Germany": 20.0, "Atlantis": 1.0}
    rolled = regions.rollup_regions(totals)
    assert list(rolled) == [regions.NORTH_AMERICA, regions.EUROPE, regions.ASIA_PACIFIC, regions.REST_OF_WORLD]
    assert sum(rolled.values()) == pytest.approx(101.0)


# --------------------------------------------------------------------------- #
# split
# --------------------------------------------------------------------------- #
def test_normalize_shares_accepts_percentages_and_fills_missing():
    shares, equal = split.normalize_shares({"A": 60, "B": 40}, ["A", "B", "C"])
    assert not equal
    assert shares == {"A": 0.6, "B": 0.4, "C": 0.0}


def test_normalize_shares_falls_back_to_equal_split():
    shares, equal = split.normalize_shares({}, ["A", "B"])
    assert equal
    assert shares == {"A": 0.5, "B": 0.5}


def test_split_top_down_adds_back_exactly_to_country_totals():
    totals = {"India": 100.0, "Germany": 33.3333}
    shares = {"India": {"A": 1 / 3, "B": 1 / 3, "C": 1 / 3}, "Germany": {"A": 0.5, "B": 0.25, "C": 0.25}}
    values = split.split_top_down(totals, shares, ["A", "B", "C"])
    for country, total in totals.items():
        assert sum(values[country].values()) == pytest.approx(total, abs=1e-9)
    table = split.build_dimension_table("By X", "top_down", ["A", "B", "C"], values)
    checks = split.consistency_checks(totals, [table])
    assert checks[0]["passed"]
    assert table["global_total"] == pytest.approx(133.3333, abs=1e-4)


def test_consistency_check_catches_a_dimension_that_does_not_add_up():
    totals = {"India": 100.0}
    table = split.build_dimension_table("By X", "bottom_up", ["A"], {"India": {"A": 90.0}})
    assert not split.consistency_checks(totals, [table])[0]["passed"]


# --------------------------------------------------------------------------- #
# segments: draft / approve / finalize
# --------------------------------------------------------------------------- #
TOTALS = {"United States": 60.0, "Germany": 25.0, "India": 15.0}
DIMS = [
    {"dimension": "By Product Type", "segments": ["Liquid", "Powder"]},
    {"dimension": "By Application", "segments": ["Food", "Cleaning"]},
    {"dimension": "By Region", "segments": ["NA", "EU"]},  # geography: must be dropped
]
SHARES = {
    "By Product Type": {c: {"Liquid": 0.8, "Powder": 0.2} for c in TOTALS},
    "By Application": {c: {"Food": 0.7, "Cleaning": 0.3} for c in TOTALS},
}


def _draft(bottom_up=None, evidence=None, snippets=None):
    llm = FakeLLM(DIMS, SHARES, evidence=evidence)
    draft = segments.build_segments_draft(
        llm=llm, config=NoDbConfig(), approach="test", market_name="Test Market",
        scope="all of it", base_year=2025, country_totals=TOTALS,
        bottom_up_dimensions=bottom_up, search=lambda queries: list(snippets or []),
    )
    return draft, llm


def test_draft_drops_geography_and_keeps_application_when_no_bottom_up():
    draft, _ = _draft()
    names = [d["dimension"] for d in draft["top_down_dimensions"]]
    assert names == ["By Product Type", "By Application"]
    assert all(d["evidence_status"] == "assumption" for d in draft["top_down_dimensions"])


def test_draft_excludes_application_when_it_is_built_bottom_up():
    bottom_up = [{"dimension": "By Application (bottom-up)", "segments": ["Food"], "country_values": {}}]
    draft, _ = _draft(bottom_up=bottom_up)
    assert [d["dimension"] for d in draft["top_down_dimensions"]] == ["By Product Type"]


def test_evidence_status_requires_a_verbatim_quote_with_a_number():
    snippet = "Industry report: liquid products held 78% of the market in 2024."
    evidence = [{"segment": "Liquid", "quote": "liquid products held 78% of the market"}]
    draft, _ = _draft(evidence=evidence, snippets=[snippet])
    assert draft["top_down_dimensions"][0]["evidence_status"] == "web_evidence_backed"

    fabricated = [{"segment": "Liquid", "quote": "liquid products held 95% of the market"}]
    draft, _ = _draft(evidence=fabricated, snippets=[snippet])
    assert draft["top_down_dimensions"][0]["evidence_status"] == "assumption"


def test_approve_rejects_shares_that_do_not_sum_to_one():
    draft, _ = _draft()
    draft["top_down_dimensions"][0]["shares"]["India"] = {"Liquid": 0.5, "Powder": 0.2}
    with pytest.raises(ValueError, match="sum to"):
        segments.approve_segments(draft)


def test_approve_rejects_missing_country_and_unknown_segment():
    draft, _ = _draft()
    del draft["top_down_dimensions"][0]["shares"]["India"]
    with pytest.raises(ValueError, match="no shares for"):
        segments.approve_segments(draft)

    draft, _ = _draft()
    draft["top_down_dimensions"][0]["shares"]["India"] = {"Liquid": 0.5, "Gel": 0.5}
    with pytest.raises(ValueError, match="unknown segment"):
        segments.approve_segments(draft)


def test_approve_accepts_percentages_after_human_edit():
    draft, _ = _draft()
    draft["top_down_dimensions"][0]["shares"]["India"] = {"Liquid": 70, "Powder": 30}
    approved = segments.approve_segments(draft)
    assert approved["top_down_dimensions"][0]["shares"]["India"] == {"Liquid": 0.7, "Powder": 0.3}


def test_finalize_report_numbers_add_up_across_every_dimension():
    draft, _ = _draft()
    report = segments.finalize_report(segments.approve_segments(draft), "test")
    assert report["geography"]["global"] == pytest.approx(100.0)
    assert report["geography"]["regions"] == {
        regions.NORTH_AMERICA: 60.0, regions.EUROPE: 25.0, regions.ASIA_PACIFIC: 15.0,
    }
    product = report["segmentation"][0]
    assert product["global_values"] == {"Liquid": pytest.approx(80.0), "Powder": pytest.approx(20.0)}
    assert all(c["passed"] for c in report["consistency_checks"])
    assert "Global: $100.0M" in format_terminal_summary(report)


# --------------------------------------------------------------------------- #
# approach 2B: Application bottom-up from 02_sizing.json, rest top-down
# --------------------------------------------------------------------------- #
def _sizing_2b():
    rows = []
    values = {
        "United States": {"Food": 40.0, "Cleaning": 20.0},
        "Germany": {"Food": 15.0, "Cleaning": 10.0},
        "India": {"Food": 10.0, "Cleaning": 5.0},
    }
    for c, apps in values.items():
        for a, v in apps.items():
            rows.append({"country": c, "usage_bucket": a, "value_usd_mn": v})
        rows.append({"country": c, "usage_bucket": "Strategic reserve", "value_usd_mn": 99.0})
    return {
        "market_name": "Test Market",
        "market_setup": {"scope": "all of it", "base_year": 2025},
        "usage_buckets": [
            {"name": "Food", "bucket_category": "physical_consumption"},
            {"name": "Cleaning", "bucket_category": "physical_consumption"},
            {"name": "Strategic reserve", "bucket_category": "inventory_reserve"},
        ],
        "country_totals_usd_mn": {c: sum(a.values()) for c, a in values.items()},
        "build_up_rows": rows,
    }


def test_2b_segments_report_end_to_end(tmp_path, monkeypatch):
    from me_engine.demandUseCaseCountry_approach2B import checkpoints, segmentation

    monkeypatch.setattr(checkpoints, "_OUTPUT_ROOT", tmp_path / "demand_usage")
    sizing = _sizing_2b()
    checkpoints.write_json(checkpoints.sizing_output_prefix("Test Market").with_suffix(".json"), sizing)

    llm = FakeLLM(DIMS, SHARES)
    draft = segmentation.run_segments_phase(
        "Test Market", NoDbConfig(), search=lambda q: [], llm=llm,
    )
    bottom_up = draft["bottom_up_dimensions"][0]
    assert bottom_up["segments"] == ["Food", "Cleaning"]  # inventory bucket excluded
    assert [d["dimension"] for d in draft["top_down_dimensions"]] == ["By Product Type"]

    segmentation.approve_segments(draft)
    report, json_path, xlsx_path = segmentation.run_report_phase("Test Market")

    app_table, product_table = report["segmentation"]
    assert app_table["method"] == "bottom_up"
    assert app_table["global_values"] == {"Food": 65.0, "Cleaning": 35.0}
    assert product_table["method"] == "top_down"
    assert product_table["global_total"] == pytest.approx(100.0)
    assert all(c["passed"] for c in report["consistency_checks"])

    assert json.loads(json_path.read_text(encoding="utf-8"))["approach"] == "2B"
    wb = load_workbook(xlsx_path)
    assert wb.sheetnames[0] == "Summary" and wb.sheetnames[-1] == "Method & Evidence"
    assert len(wb.sheetnames) == 4


# --------------------------------------------------------------------------- #
# approach 2A: driver gate, region rollup, every dimension top-down
# --------------------------------------------------------------------------- #
def test_2a_driver_gate_requires_exactly_one_total_demand_line(tmp_path, monkeypatch):
    from me_engine.demandCountry_approach2A import checkpoints, driver

    monkeypatch.setattr(checkpoints, "_OUTPUT_ROOT", tmp_path / "demand_country_2a")
    bucket = driver.total_demand_bucket("Vinegar", "litres")
    assert bucket["bucket_category"] == "physical_consumption"
    good = {"market_name": "Vinegar", "usage_buckets": [bucket], "drivers": {bucket["name"]: {"driver_name": "x"}}}
    assert driver.approve_driver(good)["approved"]

    two = dict(good, usage_buckets=[bucket, dict(bucket, name="Other")])
    with pytest.raises(ValueError, match="exactly ONE"):
        driver.approve_driver(two)
    with pytest.raises(ValueError, match="no demand driver"):
        driver.approve_driver(dict(good, drivers={}))


def test_2a_region_rollup_on_sizing_state():
    from me_engine.demandCountry_approach2A.graph import add_region_rollup

    out = add_region_rollup({"country_totals": TOTALS, "log": []})
    assert out["region_totals"][regions.NORTH_AMERICA] == 60.0
    assert out["country_region"]["India"] == regions.ASIA_PACIFIC
    assert out["log"][-1].startswith("[2A] country -> region rollup")


def test_2a_segments_report_end_to_end(tmp_path, monkeypatch):
    from me_engine.demandCountry_approach2A import checkpoints, segmentation

    monkeypatch.setattr(checkpoints, "_OUTPUT_ROOT", tmp_path / "demand_country_2a")
    checkpoints.write_json(
        checkpoints.sizing_output_prefix("Test Market").with_suffix(".json"),
        {"market_setup": {"scope": "all", "base_year": 2025}, "country_totals_usd_mn": TOTALS},
    )
    draft = segmentation.run_segments_phase(
        "Test Market", NoDbConfig(), search=lambda q: [], llm=FakeLLM(DIMS, SHARES),
    )
    assert draft["bottom_up_dimensions"] == []
    assert [d["dimension"] for d in draft["top_down_dimensions"]] == ["By Product Type", "By Application"]

    segmentation.approve_segments(draft)
    report, _json_path, xlsx_path = segmentation.run_report_phase("Test Market")
    assert {t["method"] for t in report["segmentation"]} == {"top_down"}
    app = next(t for t in report["segmentation"] if t["dimension"] == "By Application")
    assert app["global_values"] == {"Food": pytest.approx(70.0), "Cleaning": pytest.approx(30.0)}
    assert all(c["passed"] for c in report["consistency_checks"])
    assert xlsx_path.exists()


def test_2a_report_requires_approved_segments(tmp_path, monkeypatch):
    from me_engine.demandCountry_approach2A import checkpoints, segmentation

    monkeypatch.setattr(checkpoints, "_OUTPUT_ROOT", tmp_path / "demand_country_2a")
    with pytest.raises(FileNotFoundError):
        segmentation.run_report_phase("Nothing Here")


def test_report_db_taxonomy_rejected_when_it_describes_another_market(monkeypatch):
    db_dims = [{"dimension": "By Type", "segments": ["AC Type", "DC Type"], "taxonomy_source": "cmi_report_db"}]
    monkeypatch.setattr(segments, "dimensions_from_report_db", lambda m, c: db_dims)

    class Judge(FakeLLM):
        def __init__(self, fits):
            super().__init__(DIMS, SHARES)
            self.fits = fits

        def complete_json(self, system, user, *, retries=3):
            if "Candidate segmentation taken from a stored report" in user:
                return {"fits": self.fits, "reason": "test"}
            return super().complete_json(system, user)

    rejected = segments.propose_dimensions(Judge(False), NoDbConfig(), "Crude Oil", "crude oil", False)
    assert [d["taxonomy_source"] for d in rejected] == ["llm_proposed", "llm_proposed"]

    accepted = segments.propose_dimensions(Judge(True), NoDbConfig(), "Crude Oil", "crude oil", False)
    assert accepted == db_dims
