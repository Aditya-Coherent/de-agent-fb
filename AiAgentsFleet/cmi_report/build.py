"""Fills the CMI whitepaper template (the drone-market .docx) with one
market's model, prose, agent tables and charts.

The template is cloned element by element (body indices below refer to
`Updated_Drone Market_Whitepaper Template_CMI.docx`); each element keeps its
styling and only its content changes. Figures and tables are renumbered in
order. Removed on purpose: the 'Data status' box, the Contact section (with
its email/phone placeholders) and the Disclaimer, plus the 'draft template'
wording in the footer and the assumptions list.
"""
from __future__ import annotations

import copy
import re
from datetime import datetime
from pathlib import Path

import docx
from docx.oxml.ns import qn
from PIL import Image

from AiAgentsFleet.cmi_report import fmt
from AiAgentsFleet.cmi_report.guard import ClaimGuard, NumberGuard, clean_text, vendor_clean
from AiAgentsFleet.cmi_report.model import Y0, Y1
from AiAgentsFleet.cmi_report.guard import sentences
from AiAgentsFleet.cmi_report.groups import canonical, short
from AiAgentsFleet.cmi_report.prose import (absent_names, base_facts, claim_groups, region_country_fallback,
                                            seg_fallback)

_COUNT = {1: "One", 2: "Two", 3: "Three", 4: "Four", 5: "Five", 6: "Six"}


def _sentence(text: str) -> str:
    text = (text or "").strip()
    return text if not text or text[-1] in ".!?" else text + "."


_MID_THOUGHT = re.compile(r"^\s*(?:within|also|the region|this market|it|its|they|these|this)\b"
                          r"|^\W*(?:\w+\W+){0,5}?(?:also|too)\b", re.I)


def open_chapter(text: str | None, must: str, lead: str, avoid: str | None = None) -> str:
    """A chapter's first sentence names the chapter's own subject (`must`) and does not start
    mid-thought ('Within Africa, ...', 'X is also ...'); otherwise that sentence goes and a
    data-based opener (`lead`) comes first. `avoid`: a subject that must not lead this chapter
    (the other sub-region)."""
    sents = sentences(text or "")
    if not sents:
        return lead
    first = sents[0]
    wrong_subject = bool(avoid) and re.search(avoid, first) and \
        (not re.search(must, first) or re.search(avoid, first).start() < re.search(must, first).start())
    if re.search(must, first) and not _MID_THOUGHT.search(first) and not wrong_subject:
        return " ".join(sents)
    rest = sents[1:] if (_MID_THOUGHT.search(first) or wrong_subject) else sents
    return " ".join([lead] + rest)


def complete_rows(rows: list[list[str]], log: list[str], what: str) -> list[list[str]]:
    """Agent-written rows with an empty cell are left out rather than shown with '-'."""
    keep = [r for r in rows if all(str(c).strip() and str(c).strip() not in ("-", "–") for c in r)]
    if len(keep) < len(rows):
        log.append(f"{what}: {len(rows) - len(keep)} incomplete row(s) left out")
    return keep

REMOVE = {22, 76, 77, 78, 79, 80, 393, 394, 395, 396}
PROFILE_BLOCKS = {  # first element of each country profile block (12 elements each)
    "Middle East": [169, 182, 195, 208],
    "Africa": [249, 261, 273, 285],
}
# what each chapter's closing sentence must set up: (topic, keyword the closing must contain)
NEXT_TOPIC = {
    "ch2": ("the market trends, drivers and restraints behind these numbers", r"trend|driver|restraint"),
    "ch3": ("the global market size, scenarios, segments and regions", r"global|size|scenario|segment"),
    "ch4": ("the Middle East market opportunity", r"Middle East"),
    "ch5": ("the Africa market opportunity", r"Africa"),
    "ch6": ("the regulatory and policy landscape", r"regulat|polic"),
    "ch7": ("the vendor landscape", r"vendor|compet|player|supplier"),
    "ch8": ("the strategic recommendations", r"recommend|strateg|action"),
    "ch9": ("the research methodology behind these estimates", r"method"),
    None: ("the publisher behind this report", r"CMI|Coherent"),
}
PROFILE_STARTS = {v[0]: k for k, v in PROFILE_BLOCKS.items()}           # 169 -> Middle East, 249 -> Africa
PROFILE_RANGE = set(range(169, 221)) | set(range(249, 297))               # all template profile elements
COUNTRY_TABLES = {  # (heading-2, intro paragraph, table) per region in section 3.5
    "North America": (114, 115, 117), "Latin America": (119, 120, 122), "Europe": (124, 125, 127),
    "Asia Pacific": (130, 131, 133), "Middle East": (135, 136, 138), "Africa": (140, 141, 143),
}


# --------------------------------------------------------------------------- #
# XML helpers
# --------------------------------------------------------------------------- #
def _runs(p):
    return [r for r in p.findall(qn("w:r")) if r.find(qn("w:drawing")) is None and r.find(qn("w:fldChar")) is None
            and r.find(qn("w:instrText")) is None]


def _is_bold(r) -> bool:
    rpr = r.find(qn("w:rPr"))
    b = rpr.find(qn("w:b")) if rpr is not None else None
    return b is not None and b.get(qn("w:val")) not in ("0", "false")


def _text_run(proto, text: str, bold: bool | None = None):
    r = copy.deepcopy(proto)
    for ch in list(r):
        if ch.tag != qn("w:rPr"):
            r.remove(ch)
    if bold is not None:
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            rpr = r.makeelement(qn("w:rPr"), {})
            r.insert(0, rpr)
        for b in rpr.findall(qn("w:b")) + rpr.findall(qn("w:bCs")):
            rpr.remove(b)
        if bold:
            rpr.insert(0, rpr.makeelement(qn("w:b"), {}))
    t = r.makeelement(qn("w:t"), {})
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    return r


def set_text(p, text: str, bold: bool | None = None):
    """Replaces a paragraph's text, keeping its first plain run's formatting
    (or the first run's) and the paragraph properties (style, bullets)."""
    runs = _runs(p)
    plain = [r for r in runs if not _is_bold(r)]
    proto = (plain or runs or [p.makeelement(qn("w:r"), {})])[0]
    for r in runs:
        p.remove(r)
    p.append(_text_run(proto, text, bold))
    return p


def set_lead(p, lead: str, text: str):
    """Bold lead-in + plain text (template bullet/box pattern)."""
    runs = _runs(p)
    proto = (runs or [p.makeelement(qn("w:r"), {})])[0]
    for r in runs:
        p.remove(r)
    p.append(_text_run(proto, lead.rstrip() + " ", True))
    p.append(_text_run(proto, text, False))
    return p


def paras(el):
    return list(el.iter(qn("w:p")))


def _text_runs(p):
    return [r for r in _runs(p) if r.find(qn("w:t")) is not None]


def set_h1(p, num: str | None = None, title: str | None = None):
    """Chapter heading = number run ('03   ', orange) + title runs (white):
    each is rewritten in place so the two colours survive."""
    rs = _text_runs(p)
    if num is not None:
        rs[0].find(qn("w:t")).text = num
    if title is not None:
        for k, r in enumerate(rs[1:]):
            for t in r.findall(qn("w:t")):
                t.text = title if k == 0 else ""
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    return p


def _style(el) -> str:
    s = el.find(".//" + qn("w:pStyle"))
    return s.get(qn("w:val")) if s is not None else ""


# Optional blocks of the template (body indices) and the section flag that keeps them.
def skipped(m: dict) -> tuple[set[int], list[str]]:
    sec = m["sections"]
    skip: set[int] = set()
    notes = []

    def drop(name: str, idx):
        skip.update(idx)
        notes.append(f"removed {name}: {sec[name.split(':')[0]]['why']}")

    profile_starts = [s for st in PROFILE_BLOCKS.values() for s in st]
    if not sec["type"]["show"]:
        drop("type: 'by type' sections", list(range(83, 88)) + list(range(158, 163)) + list(range(238, 243))
             + [s + o for s in profile_starts for o in range(2, 7)])
    if not sec["demand"]["show"]:
        drop("demand: 'by demand segment' sections", list(range(88, 94)) + list(range(163, 168))
             + list(range(243, 248)) + [s + o for s in profile_starts for o in range(7, 12)])
    if not sec["type"]["show"] and not sec["demand"]["show"]:
        skip.update({81, 82})
    if not sec["use_cases_me"]["show"]:
        drop("use_cases_me: Middle East use-case heat map", range(221, 225))
    if not sec["use_cases_af"]["show"]:
        drop("use_cases_af: Africa use-case heat map", range(298, 302))
    # vendor shares (PLAN-05): a donut only from sourced shares; otherwise the positioning table
    if not sec["vendor_global"]["show"] and not sec["positioning"]["show"]:
        drop("vendor_global: global vendor share", range(332, 337))
    if not sec["vendor_mea"]["show"] and not sec["positioning"]["show"]:
        drop("vendor_mea: MEA vendor share", range(337, 343))
    if not sec["capabilities"]["show"]:
        drop("capabilities: vendor capability map", range(343, 347))
    if not sec["dashboard"]["show"]:
        drop("dashboard: vendor dashboard", range(347, 351))
    if not sec["milestones"]["show"]:
        drop("milestones: regulatory milestones timeline", range(305, 308))
    if not sec["regulation"]["show"]:
        drop("regulation: Regulatory & Policy Outlook chapter", range(303, 325))
    return skip, notes


def to_verify_items(b: "Builder", ag: dict, m: dict) -> list[str]:
    """Every [to verify] item exactly as printed in the finished document (text
    and tables), plus the milestones marked on the timeline chart -- so the
    appendix count matches the report."""
    items = []
    for el in b.out:
        if el.tag == qn("w:tbl"):
            for tr in el.iter(qn("w:tr")):
                cells = ["".join(t.text or "" for t in tc.iter(qn("w:t"))) for tc in tr.findall(qn("w:tc"))]
                if any("[to verify]" in c for c in cells):
                    items.append(" | ".join(c for c in cells if c)[:200])
        else:
            text = "".join(t.text or "" for t in el.iter(qn("w:t")))
            items.extend(x[:200] for x in sentences(text) if "[to verify]" in x)
    if m["sections"]["regulation"]["show"] and m["sections"]["milestones"]["show"]:
        items += [f"Timeline chart: {x.get('date')} {x.get('jurisdiction', '')} -- {x.get('event')} [to verify]"
                  for x in (ag.get("regulation") or {}).get("milestones") or [] if x.get("verified") is False]
    return list(dict.fromkeys(items))


POSITIONING_HEADER = ["Vendor", "HQ", "Parent group", "Business type", "Role in this market",
                      "MEA presence (CMI assessment)"]
_POSITIONING_WIDTHS = [1700, 1100, 1500, 1400, 2306, 1500]      # the dashboard table's 9,506 twips


def _widen(tbl, widths: list[int]):
    """The template's 4-column vendor dashboard table with len(widths) columns
    (cells cloned from each row's last cell)."""
    grid = tbl.find(qn("w:tblGrid"))
    for gc in grid.findall(qn("w:gridCol")):
        grid.remove(gc)
    for w in widths:
        gc = grid.makeelement(qn("w:gridCol"), {})
        gc.set(qn("w:w"), str(w))
        grid.append(gc)
    for tr in tbl.findall(qn("w:tr")):
        while len(tr.findall(qn("w:tc"))) < len(widths):
            last = tr.findall(qn("w:tc"))[-1]
            last.addnext(copy.deepcopy(last))
        for tc, w in zip(tr.findall(qn("w:tc")), widths):
            tcw = tc.find(qn("w:tcPr") + "/" + qn("w:tcW"))
            if tcw is not None:
                tcw.set(qn("w:w"), str(w))
    return tbl


def short_note(note: str | None) -> str:
    """A confirmed ownership note beside a vendor name ('now part of Huber'); a long
    sentence is not a name annotation and is left out."""
    note = str(note or "").strip().rstrip(".")
    return note if note and len(note.split()) <= 10 else ""


def positioning_rows(m: dict, vsafe) -> list[list[str]]:
    """vendor | HQ | parent group | business type | role | MEA presence -- confirmed vendors only, in
    alphabetical order; no tier or ranking (there are no sourced shares to rank by). The role is the one
    the product evidence line states; MEA presence is a CMI assessment, never a percentage."""
    rows = []
    for r in m["vendor_positioning"]:
        own = short_note(r.get("ownership_note"))
        rows.append([r["vendor"] + (f" ({own})" if own and own.lower() not in r["vendor"].lower() else ""),
                     r["hq"], r["parent"], r["business_type"], vsafe(r["role"]).rstrip("."),
                     r["mea_presence"] or "Not assessed"])
    return rows


def vendor_shares(b: "Builder", i: int, m: dict, s, root: str, charts: dict, vsafe) -> None:
    """Sections 7.2 / 7.3 (template elements 332-341): a share donut where at least
    3 shares are sourced, else the positioning table (once, in the first section
    without a donut) and the section's paragraph (PLAN-05)."""
    sec, vy = m["sections"], m["vendor_year"]
    show_g, show_m, pos = sec["vendor_global"]["show"], sec["vendor_mea"]["show"], sec["positioning"]["show"]
    pos_at = 332 if not show_g else 337           # where the positioning table goes
    geo, show, label = ("global", show_g, "Global") if i < 337 else ("MEA", show_m, "MEA")
    k = i - (332 if i < 337 else 337)              # 0 heading, 1 caption, 2 figure, 3 source, 4 paragraph
    if k == 0:
        if show:
            b.keep(i)
        else:
            b.para(i, f"7.{2 if i < 337 else 3} Vendor positioning ({label})")
    elif k == 4:
        b.para(i, s("global_share_para" if i < 337 else "mea_share_para"))
    elif show:
        if k == 1:
            b.caption(i, "Figure", f"Vendor share of {geo} {root.lower()} market value, {vy} (published figures)")
        elif k == 2:
            b.image(i, charts[f"vendor_{geo.lower()}"])
        else:
            pubs = dict.fromkeys(x.get("domain") or "" for x in m["vendor_sources"][geo.lower()])
            b.para(i, "Source: " + ", ".join(p for p in pubs if p) + "; shares as stated by each source")
    elif pos and (i - k) == pos_at:
        if k == 1:
            b.caption(i, "Table", f"Vendor positioning in the {root.lower()} market (published shares are not "
                                  f"available for enough vendors; vendors in alphabetical order; MEA presence is a CMI assessment)")
        elif k == 2:
            b.src.append(_widen(b.el(349), _POSITIONING_WIDTHS))     # the dashboard table, 6 columns
            b.table(len(b.src) - 1, complete_rows(positioning_rows(m, vsafe), b.log, "vendor positioning"),
                    header=POSITIONING_HEADER)
        else:
            b.para(i, "Source: company websites, filings and news (HQ, parent group, business type and role, "
                      "each confirmed in the evidence); CMI assessment (MEA presence)")


def add_appendix(b: "Builder", m: dict, ag: dict, prose: dict, removed_notes: list[str],
                 verification: dict | None = None, register: list[dict] | None = None) -> None:
    """Appendix 'Consistency & Verification Log': master data sheet, rankings
    register, [to verify] list and the corrections made while building."""
    from AiAgentsFleet.cmi_report.prose import _rank
    H1, H2, BODYP, BULLET, TBL5, TBL3, CAP, SRC = 385, 387, 392, 381, 117, 313, 40, 42
    br = b.el(181)                                   # page break paragraph
    b.out.append(br)
    b.out.append(set_h1(b.el(H1), num="11   ", title="Appendix: Consistency & Verification Log"))
    b.para(BODYP, "Every figure in this report is drawn from the master data sheet below; every ranking from the "
                  "rankings register. Values are CMI estimates (to 2026) and forecasts (2027 onward).")
    # contents entry for the appendix (cloned from chapter 10's entry, placed right after it)
    toc_el = b.el(18)
    if 10 in b.toc:
        b.out.insert(b.out.index(b.toc[10]) + 1, toc_el)
    b.toc[11] = toc_el

    def row(label, v):
        return [label, fmt.usd(v["v26"]), fmt.usd(v["v33"]), fmt.pct(v["cagr"]), fmt.incr(v["incr"])]

    b.out.append(set_text(b.el(H2), "A.1 Master data sheet"))
    rows = [row("Global", m["global"])] + [row(r, v) for r, v in m["regions"].items()] + \
           [row(f"{r} (sub-region)", v) for r, v in m["subregions"].items()]
    for r, d in m["countries"].items():
        rows += [row(f"{r}: {c}", v) for c, v in d.items()]
    for geo in ("Global", "Middle East", "Africa"):
        for dim in ("type", "demand"):
            if m["sections"][dim]["show"]:
                lab = m["dimensions"][dim]["singular"]
                rows += [row(f"{geo} {lab}: {n}", v) for n, v in m["segments"][geo][dim].items()]
    b.caption(CAP, "Table", "Master data sheet (USD; CMI estimates and forecasts)")
    b.table(TBL5, rows, header=["Line item", f"{Y0}E", f"{Y1}F", "CAGR", "Incremental $"])
    b.keep(SRC)

    # A.2 Number register: every DERIVED fact (shares, gaps, multiples, scenarios, vendor and value-chain
    # shares, published figures) with its ID; the full register (incl. every segment value) is the CSV
    # saved beside this report
    b.out.append(set_text(b.el(H2), "A.2 Number register (derived facts)"))
    b.para(BODYP, "Each derived figure used in the report, with its register ID. The complete register, including "
                  "every segment value, is supplied with this report as a CSV file.")
    derived = [e for e in register or [] if e["metric"] in ("share_parent", "share_global", "gap_pp", "multiple",
                                                            "share", "published")
               and not e["id"].startswith("S-")]
    derived += [e for e in register or [] if e["id"].startswith("SC-")]
    b.caption(CAP, "Table", "Number register: derived facts")
    b.table(TBL3, [[e["id"], e["fact"], e["display"] + (" pp" if e["metric"] == "gap_pp" else "")]
                   for e in derived], header=["ID", "Fact (exactly what it measures)", "Value"])
    b.keep(SRC)

    b.out.append(set_text(b.el(H2), "A.3 Rankings register"))
    groups = {"Regions": m["regions"], "Middle East vs Africa": m["subregions"],
              **{f"{r} countries": d for r, d in m["countries"].items()}}
    for geo in ("Global", "Middle East", "Africa"):
        for dim in ("type", "demand"):
            if m["sections"][dim]["show"]:
                groups[f"{geo} {m['dimensions'][dim]['plural'].lower()}"] = m["segments"][geo][dim]
    labels = {f"largest_in_{Y0}": f"Largest in {Y0}E", f"largest_in_{Y1}": f"Largest in {Y1}F",
              f"smallest_in_{Y0}": f"Smallest in {Y0}E", "fastest_growing": "Fastest-growing (CAGR)",
              "slowest_growing": "Slowest-growing (CAGR)", "largest_absolute_opportunity": "Largest absolute opportunity"}
    rrows = []
    for gname, d in groups.items():
        for k, v in _rank(d).items():
            rrows.append([f"{gname} ({len(d)})", labels.get(k, k), v])
    b.caption(CAP, "Table", "Rankings register (comparison set and size in brackets)")
    b.table(TBL3, rrows, header=["Comparison set", "Ranking", "Result"])
    b.keep(SRC)

    b.out.append(set_text(b.el(H2), "A.4 Items marked [to verify]"))
    tv = to_verify_items(b, ag, m)
    b.bullets(BULLET, tv or ["No item in this report is marked [to verify]."])

    b.out.append(set_text(b.el(H2), "A.5 Corrections made during the self-audit"))
    corr = list(m.get("corrections") or []) + list(removed_notes)
    dropped = prose.get("dropped") or {}
    n_drop = sum(len(v) for v in dropped.values())
    if n_drop:
        why = {"numbers not in model": 0, "names": 0, "repeat": 0, "headline restated": 0, "caveat": 0, "claim": 0}
        for v in dropped.values():
            for x in v:
                key = ("numbers not in model" if "numbers not in model" in x else "names" if "not in the data" in x
                       else "repeat" if x.startswith("[repeat]") else "headline restated" if "headline" in x
                       else "caveat" if "caveat" in x else "claim")
                why[key] += 1
        corr.append(f"{n_drop} generated sentence(s) removed before publication: " + ", ".join(
            f"{v} {k}" for k, v in {"with a number not in the master sheet": why["numbers not in model"],
                                    "with a ranking the register does not support": why["claim"],
                                    "naming a category absent from the data": why["names"],
                                    "repeating an earlier sentence": why["repeat"],
                                    "restating the headline outside the summary": why["headline restated"],
                                    "with template or draft wording": why["caveat"]}.items() if v))
    corr += [x for x in b.log if "incomplete row" in x]
    b.bullets(BULLET, corr or ["No corrections were needed."])

    b.out.append(set_text(b.el(H2), "A.6 Consistency & Verification Log"))
    v = verification or {}
    log = [
        "Every value, share, gap and ranking was derived from the master data sheet (A.1), the number "
        "register (A.2) and the rankings register (A.3) before any text was written; text, tables and chart labels print them in one format. "
        + fmt.PRECISION_RULE,
        "Checked: children sum to parents for every year (countries to sub-regions, sub-regions to regions, "
        "regions to global, segments to totals); every CAGR recomputed from its start and end values; every "
        "incremental value equals end minus start; every share uses its own denominator; every ranking and tie "
        "is taken from the rankings register; share sets sum to 100%.",
    ]
    if v:
        log.append(f"Number consistency sweep of the finished report: {v.get('checked', 0)} numbers checked; "
                   f"{v.get('first_sweep', 0)} text block(s) with a number not matching the register in the first "
                   f"sweep; {v.get('fixed', 0)} sentence(s) removed or corrected; {v.get('final_sweep', 0)} "
                   f"mismatch(es) remaining after the final sweep.")
    log.append(f"{len(tv)} item(s) are marked [to verify] (A.4).")
    b.bullets(BULLET, log)


def size_basis_text(m: dict, root: str) -> str:
    """Section 9.1: where the global base comes from (cmi_report/verify_size.py). Names the
    source of a sourced headline; otherwise calls it a CMI estimate. Never explains gaps
    to other figures, and an unsourced headline gets no published-comparison wording."""
    g, chk = m["global"], m.get("size_check") or {}
    base = f"CMI's global {root.lower()} base of {fmt.usd(g['v25'])} for {Y0 - 1}"
    src = chk.get("source") if chk.get("status") == "sourced" else None
    if not src:
        return (f"{base} is a CMI estimate built from the desk research for this report; no third-party "
                f"figure for this market could be confirmed in the evidence gathered, so none is cited.")
    quoted = src.get("derivation") == "as quoted"
    if chk.get("source_is_cmi"):
        text = (f"{base} is CMI's published estimate." if quoted else
                f"{base} is carried forward from CMI's published {src['year']} estimate, at the growth rate it states.")
    else:
        how = "is the figure reported by" if quoted else f"is carried forward from the {src['year']} figure reported by"
        tail = "" if quoted else ", at the growth rate that source states"
        who = src["domain"] if src["publisher"] == src["domain"] else f"{src['publisher']} ({src['domain']})"
        text = f"{base} {how} {who}{tail}."
    if not chk.get("cagr_sourced"):
        text += " The forecast growth rate is a CMI estimate."
    if m.get("published_2025"):
        text += " Other published figures are shown below for reference only."
    return text


def renumber(b: "Builder", toc_titles: dict[int, str], toc_pages: dict[int, int] | None) -> None:
    """Chapters and numbered sub-sections are renumbered in order (a removed
    section leaves no gap), and the contents page lists exactly the chapters
    present, with their new numbers."""
    new_of: dict[int, int] = {}
    ch = sub = 0
    for el in b.out:
        st = _style(el)
        if st == "Heading1":
            rs = _text_runs(el)
            orig = int(re.match(r"\s*(\d+)", rs[0].find(qn("w:t")).text).group(1))
            ch, sub = ch + 1, 0
            new_of[orig] = ch
            set_h1(el, num=f"{ch:02d}   ")
        elif st == "Heading2":
            text = el.xpath("string(.)")
            mm = re.match(r"\s*\d+\.\d+\s+(.*)", text)
            if mm and ch:
                sub += 1
                set_text(el, f"{ch}.{sub} {mm.group(1)}")
    for orig, el in b.toc.items():
        if orig not in new_of:
            b.out.remove(el)
            continue
        n = new_of[orig]
        rs = _text_runs(el)
        rs[0].find(qn("w:t")).text = f" {n}    "
        for k, r in enumerate(rs[1:-1]):
            for t in r.findall(qn("w:t")):
                t.text = toc_titles[orig] if k == 0 else ""
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        if toc_pages and n in toc_pages:
            rs[-1].find(qn("w:t")).text = str(toc_pages[n])


# --------------------------------------------------------------------------- #
# builder
# --------------------------------------------------------------------------- #
class Builder:
    def __init__(self, template: Path):
        self.doc = docx.Document(str(template))
        self.body = self.doc.element.body
        self.src = [copy.deepcopy(e) for e in self.body.iterchildren()]
        self.out: list = []
        self.fig = self.tbl = 0
        self.pic_id = 9000
        self.log: list[str] = []
        self.toc: dict[int, object] = {}      # original chapter number -> its contents-page entry

    def el(self, i: int):
        return copy.deepcopy(self.src[i])

    def keep(self, i: int):
        self.out.append(self.el(i))

    def para(self, i: int, text: str | None, bold: bool | None = None):
        if text:
            self.out.append(set_text(self.el(i), text, bold))
        else:
            self.log.append(f"element {i}: no text -- left out")

    def lead(self, i: int, lead: str, text: str):
        self.out.append(set_lead(self.el(i), lead, text))

    def bullets(self, i: int, items: list[str]):
        for it in items:
            self.out.append(set_text(self.el(i), it, False))

    def caption(self, i: int, kind: str, text: str):
        if kind == "Figure":
            self.fig += 1
            n = self.fig
        else:
            self.tbl += 1
            n = self.tbl
        self.out.append(set_text(self.el(i), f"{kind} {n}: {text}", True))

    def image(self, i: int, path: Path):
        el = self.el(i)
        r_id, _ = self.doc.part.get_or_add_image(str(path))
        with Image.open(path) as im:
            w, h = im.size
        for blip in el.iter(qn("a:blip")):
            blip.set(qn("r:embed"), r_id)
        for tag in (qn("wp:extent"), qn("a:ext")):
            for ext in el.iter(tag):
                cx = int(ext.get("cx"))
                cy = int(cx * h / w)
                if cy > 5_200_000:                      # very tall: shrink to fit the page
                    cx, cy = int(5_200_000 * w / h), 5_200_000
                ext.set("cx", str(cx))
                ext.set("cy", str(cy))
        for tag in (qn("wp:docPr"), qn("pic:cNvPr")):
            for node in el.iter(tag):
                self.pic_id += 1
                node.set("id", str(self.pic_id))
        self.out.append(el)

    def table(self, i: int, rows: list[list[str]], total: list[str] | None = None, header: list[str] | None = None):
        """Header row kept; data rows cloned from the template's first data
        row (and its last row for a bold total). A cell whose value the
        template also uses (High / Medium / Low) takes that template cell's
        shading, so rating colours follow the value, not the row position."""
        tbl = self.el(i)
        trs = tbl.findall(qn("w:tr"))
        proto, last = trs[1], trs[-1]
        shading, para_proto = {}, {}
        for tr in trs[1:]:
            for tc in tr.findall(qn("w:tc")):
                pr = tc.find(qn("w:tcPr"))
                shd = pr.find(qn("w:shd")) if pr is not None else None
                key = tc.xpath("string(.)").strip()
                shading.setdefault(key, shd)
                para_proto.setdefault(key, tc.find(qn("w:p")))
        for tr in trs[1:]:
            tbl.remove(tr)
        if header:
            for tc, val in zip(trs[0].findall(qn("w:tc")), header):
                set_text(tc.find(qn("w:p")), val, True)
        for row in rows + ([total] if total else []):
            tr = copy.deepcopy(last if (total and row is total) else proto)
            for tc, val in zip(tr.findall(qn("w:tc")), row):
                ps = tc.findall(qn("w:p"))
                set_text(ps[0], str(val), True if (total and row is total) else None)
                for extra in ps[1:]:
                    tc.remove(extra)
                if str(val) in shading and str(val) in ("High", "Medium", "Low"):
                    # the template's own cell for this rating: its text colour goes with its shading
                    tc.replace(tc.find(qn("w:p")), copy.deepcopy(para_proto[str(val)]))
                    pr = tc.find(qn("w:tcPr"))
                    if pr is not None:
                        for old in pr.findall(qn("w:shd")):
                            pr.remove(old)
                        if shading[str(val)] is not None:
                            pr.append(copy.deepcopy(shading[str(val)]))
            tbl.append(tr)
        self.out.append(tbl)

    def box(self, i: int, title: str | None, items: list, numbered: bool = False):
        """Single-cell callout box: title paragraph + item paragraphs."""
        tbl = self.el(i)
        ps = paras(tbl)
        cell = ps[0].getparent()
        if title:
            set_text(ps[0], title, True)
        proto = ps[1]
        for p in ps[1:]:
            cell.remove(p)
        for k, it in enumerate(items, start=1):
            p = copy.deepcopy(proto)
            if isinstance(it, tuple):
                set_lead(p, it[0], it[1])
            elif numbered:
                # the box numbers its items; drop a number the text already carries ("1. 1. ...")
                set_lead(p, f"{k}.", re.sub(r"^\s*(?:\d+[.)]\s*)+", "", str(it)))
            else:
                set_text(p, it, False)
            cell.append(p)
        self.out.append(tbl)

    def save(self, path: Path, header: str, footer_label: str):
        for c in list(self.body.iterchildren()):
            self.body.remove(c)
        for e in self.out:
            self.body.append(e)
        for s in self.doc.sections:
            for p in s.header.paragraphs:
                for r in p.runs:
                    if "Outlook" in (r.text or ""):
                        r.text = header
            for p in s.footer.paragraphs:
                for r in p.runs:
                    if "Confidential" in (r.text or "") or "template" in (r.text or "").lower():
                        r.text = footer_label
        xml = self.doc.element.xml
        for r_id, rel in list(self.doc.part.rels.items()):
            if "image" in rel.reltype and f'"{r_id}"' not in xml:
                del self.doc.part.rels[r_id]
        self.doc.save(str(path))


# --------------------------------------------------------------------------- #
# content
# --------------------------------------------------------------------------- #
def _root(market: str) -> str:
    return re.sub(r"^\s*global\s+|\s+market\s*$", "", market, flags=re.I).strip()


def build(template: Path, out_path: Path, m: dict, prose: dict, ag: dict, charts: dict[str, Path],
          definition: dict | None, toc_pages: dict[int, int] | None = None,
          verification: dict | None = None) -> list[str]:
    b = Builder(template)
    S = prose.get("slots", {})
    root = _root(m["market"])
    tdim, ddim = m["dimensions"]["type"], m["dimensions"]["demand"]
    tl, dl = tdim["singular"].lower(), ddim["singular"].lower()
    g, mea = m["global"], m["regions"]["Middle East & Africa"]
    me, af = m["subregions"]["Middle East"], m["subregions"]["Africa"]

    from AiAgentsFleet.cmi_report.register import FactGuard, build_register
    register = build_register(m, ag)
    ng = FactGuard(register)                      # agent-written cells obey the same register
    from AiAgentsFleet.cmi_report.prose import make_guard, set_template
    set_template(template)
    cg = make_guard(m, None)
    dropped: list[str] = []
    absent = absent_names(m)

    def safe(text: str | None) -> str:
        """Agent-written table/box text through the same guards as the prose."""
        return clean_text(text or "", ng, cg, dropped, absent) if text else ""

    def vsafe(text: str | None) -> str:
        """Vendor-chapter text: the same guards, then the vendor data (guard.vendor_clean)."""
        return vendor_clean(safe(text), m.get("vendor_text"), dropped) if text else ""

    def s(key: str, fallback: str | None = None):
        v = S.get(key)
        return v if v else fallback

    def lst(key: str) -> list:
        v = S.get(key)
        return v if isinstance(v, list) else []

    def obj(key: str) -> dict:
        v = S.get(key)
        return v if isinstance(v, dict) else {}

    def pairs(key: str) -> list[tuple[str, str]]:
        """(lead, text) items from a slot, whatever shape the model returned:
        {lead, text} objects, 'Lead: text' strings; anything else is skipped."""
        out = []
        for x in lst(key):
            if isinstance(x, dict) and x.get("lead") and x.get("text"):
                out.append((str(x["lead"]).strip(), str(x["text"]).strip()))
            elif isinstance(x, str) and ":" in x[:60]:
                lead, text = x.split(":", 1)
                out.append((lead.strip(), text.strip()))
            else:
                b.log.append(f"{key}: an item without lead/text was left out")
        return [(l, t) for l, t in out if l and t]

    seg = m["segments"]
    sec = m["sections"]
    skip, removed_notes = skipped(m)
    # the scorecard's "fastest-growing" row uses a dimension the report actually shows
    fdim = "type" if sec["type"]["show"] else "demand" if sec["demand"]["show"] else None
    fastest = {geo: max(seg[geo][fdim], key=lambda n: seg[geo][fdim][n]["cagr"]) for geo in seg} if fdim else {}
    fdim_label = (tdim if fdim == "type" else ddim)["singular"].lower() if fdim else ""
    toc_titles: dict[int, str] = {}
    size_unsourced = (m.get("size_check") or {}).get("status") != "sourced"
    from AiAgentsFleet.cmi_report.model import regulatory_countries
    model_reg_countries = regulatory_countries(ag)

    def seg_para(geo: str, kind: str) -> str:
        d = seg[geo]["type" if kind == "type" else "demand"]
        dim = tdim["singular"] if kind == "type" else ddim["singular"]
        return s(f"{kind}_para__{geo}") or seg_fallback(d, dim, geo)

    def country_rows(region: str) -> tuple[list[list[str]], list[str]]:
        d = m["countries"][region]
        rest = [k for k in d if k.startswith("Rest of")]
        order = sorted((k for k in d if k not in rest), key=lambda k: -d[k]["v26"]) + rest
        unit = fmt.unit_for([d[k]["v26"] for k in d])
        rows = [[k, fmt.usd_in(d[k]["v26"], unit), fmt.usd_in(d[k]["v33"], unit), fmt.pct(d[k]["cagr"]),
                 fmt.incr(d[k]["incr"], unit)] for k in order]
        tot = m["subregions"].get(region) or m["regions"].get(region)
        total = [f"{region} total", fmt.usd_in(tot["v26"], unit), fmt.usd_in(tot["v33"], unit),
                 fmt.pct(tot["cagr"]), fmt.incr(tot["incr"], unit)]
        return rows, total

    month = datetime.now().strftime("%B %Y").upper()
    H1_CH = {23: "ch1", 44: "ch2", 64: "ch3", 146: "ch4", 226: "ch5", 303: "ch6", 325: "ch7", 352: "ch8",
             362: "ch9", 385: None}
    current_ch = None
    for i in range(len(b.src)):
        if i in REMOVE or i in skip:
            continue
        if i in H1_CH:
            # close the previous chapter with the sentence that sets up this one
            if current_ch and H1_CH[i] is not None:
                # every chapter ends by setting up the chapter that actually follows (H1_CH[i])
                key = f"{current_ch}.closing_{current_ch}" if current_ch in ("ch4", "ch5") else f"closing_{current_ch}"
                nxt = H1_CH[i]
                topic, kw = NEXT_TOPIC.get(nxt, ("the publisher behind this report", r"CMI|about"))
                closing = s(key)
                # what follows "the next chapter" must be about the next chapter, and only about it
                tail = re.split(r"\b(?:next|following) chapter\b", closing or "", maxsplit=1, flags=re.I)
                tail = tail[1] if len(tail) > 1 else (closing or "")
                others = [k for c, (t, k) in NEXT_TOPIC.items() if c not in (nxt, None) and k != kw]
                # a transition sets the next chapter up without pre-stating its figures (no repetition)
                if not closing or re.search(r"\d", closing) or not re.search(kw, tail, re.I) \
                        or any(re.search(k, tail, re.I) for k in others):
                    closing = f"The next chapter turns to {topic}."
                b.para(24, closing)
            current_ch = H1_CH[i]
        # ------------------------------------------------------------ cover / TOC / boxes
        if i == 0:
            b.image(i, charts["cover"])
        elif i == 1:
            el = b.el(i)
            rs = _runs(el)
            for r in rs:
                el.remove(r)
            el.append(_text_run(rs[0], f"WHITEPAPER  |  {month}", True))
            b.out.append(el)
        elif i == 2:
            b.para(i, f"{root} Market Outlook {Y0}-{Y1}")
        elif i == 6:
            b.para(i, s("practice", "Industry Research Practice"))
        elif 9 <= i <= 18:
            # contents entry for original chapter i-8; numbers, titles and pages are
            # filled by renumber() once we know which chapters this report has
            titles = {9: "Executive Summary: Global vs Middle East & Africa",
                      10: "Current Market Trends, Growth Factors, Restraints and Impact",
                      11: f"Global {root} Market Size: Scenario Analysis by Segments, Region & Top Countries",
                      12: f"Middle East {root} Market Size: Opportunity Analysis by Segments & Country",
                      13: f"Africa {root} Market Size: Opportunity Analysis by Segments & Country"}
            el = b.el(i)
            texts = _text_runs(el)
            toc_titles[i - 8] = titles.get(i) or "".join(r.xpath("string(.)") for r in texts[1:-1])
            b.toc[i - 8] = el
            b.out.append(el)
        elif i == 20:
            # A1: the market definition comes first and everything else obeys it; inclusions and
            # exclusions are itemised once, in the Methodology scope table
            dfn = m["definition"]
            what = _sentence(dfn["definition"]) if dfn.get("definition") else ""
            # what is counted and at which value-chain level, verbatim from the definition (PLAN-04)
            measures = _sentence(dfn["measures"]) if dfn.get("measures") and dfn["measures"] != dfn.get("definition") \
                else ""
            level = f"Value-chain level: {dfn['value_chain_level']}." if dfn.get("value_chain_level") else ""
            units =(f"Market values are CMI estimates in USD (current prices): {Y0 - 1} is the base year, {Y0}E the "
                     f"estimate year and {Y0 + 1}F-{Y1}F are forecasts. Scores, ratings and shares of value are CMI "
                     f"assessments unless attributed to a named source; items marked [to verify] could not be "
                     f"confirmed from public sources.")
            scope = " ".join(x for x in sentences(s("about_scope", "") or "")
                             if not re.search(r"market values|base year|forecasts?\b|excludes|scope covers", x, re.I))
            b.box(i, None, [x for x in (f"Market definition: {what}" if what else "",
                                        f"What is measured: {measures}" if measures else "", level,
                                        scope or f"This whitepaper sizes the global {root.lower()} market and goes "
                                                 f"deep on the Middle East and Africa.", units) if x])
        # ------------------------------------------------------------ chapter 1
        elif i == 24:
            headline = (f"CMI estimates the global {root.lower()} market at {fmt.usd(g['v25'])} in "
                        f"{Y0 - 1}, rising to {fmt.usd(g['v26'])} in {Y0} and {fmt.usd(g['v33'])} by "
                        f"{Y1}, a {fmt.pct(g['cagr'])} CAGR.")
            opener = s("exec_para", headline)
            # an unsourced headline is always called a CMI estimate (PLAN-02)
            if size_unsourced and not re.search(r"\bCMI estimat", opener):
                opener = headline
            b.para(i, ((s("thesis") or "") + " " if s("thesis") else "") + opener)
        elif i == 25:
            b.caption(i, "Figure", "Headline market metrics: Global vs Middle East & Africa")
        elif i == 26:
            b.image(i, charts["kpi"])
        elif i in (29, 33):
            b.bullets(i, lst("global_bullets" if i == 29 else "mea_bullets"))
        elif i in (30, 31, 34, 35):
            continue
        elif i == 36:
            sh26, sh33 = mea["v26"] / g["v26"], mea["v33"] / g["v33"]
            cap = ("MEA outgrows the global market and gains share through " if mea["cagr"] > g["cagr"] and sh33 > sh26
                   else "MEA grows more slowly than the global market and loses share through "
                   if mea["cagr"] < g["cagr"] else "MEA grows in line with the global market through ")
            b.caption(i, "Figure", cap + str(Y1))
        elif i == 37:
            b.image(i, charts["cagr_share"])
        elif i == 40:
            b.caption(i, "Table", "Global vs MEA scorecard")
        elif i == 41:
            dyn = ag.get("dynamics") or {}
            fleet_dr = m.get("fleet_drivers") or {}

            def top(lst_, key):
                return max(lst_, key=lambda x: abs(float(x.get(key) or 0)))["name"] if lst_ else None

            def largest_pool(geo):
                d = seg[geo]["demand"] if sec["demand"]["show"] and seg[geo]["demand"] else \
                    seg[geo]["type"] if sec["type"]["show"] else {}
                if not d:
                    return None
                n = max(d, key=lambda k: d[k]["v26"])
                return f"{n} ({fmt.usd(d[n]['v26'])} in {Y0}E)"

            # data-derived values: a scorecard cell is never blank
            fallback = {
                "largest_demand_pool": (largest_pool("Global"), largest_pool("Middle East") and
                                        f"Middle East: {largest_pool('Middle East')}; Africa: {largest_pool('Africa')}"),
                "primary_catalyst": (top(fleet_dr.get("drivers") or [], "cagr_contribution_bps"),
                                     top(dyn.get("mea_drivers") or [], "impact_pp")),
                "primary_constraint": (top(fleet_dr.get("restraints") or [], "cagr_drag_bps"),
                                       top(dyn.get("mea_restraints") or [], "impact_pp")),
            }

            def two(key):
                o = obj(key)
                fg, fm = fallback.get(key, (None, None))
                return [safe(o.get("global")) or fg or "Not applicable (no data identified)",
                        safe(o.get("mea")) or fm or "Not applicable (no data identified)"]
            rows = [[f"Market size {Y0 - 1}", fmt.usd(g["v25"]), fmt.usd(mea["v25"])],
                    [f"Market size {Y0}E", fmt.usd(g["v26"]), fmt.usd(mea["v26"])],
                    [f"Market size {Y1}F", fmt.usd(g["v33"]), fmt.usd(mea["v33"])],
                    [f"CAGR {Y0}-{Y1}", fmt.pct(g["cagr"]), fmt.pct(mea["cagr"])],
                    [f"Incremental revenue {Y0}-{str(Y1)[2:]}", fmt.usd(g["incr"]), fmt.usd(mea["incr"])],
                    ["Largest demand pool", *two("largest_demand_pool")]]
            if fdim:
                ft = seg["Global"][fdim][fastest["Global"]]
                fme, faf = fastest["Middle East"], fastest["Africa"]
                cme, caf = seg["Middle East"][fdim][fme]["cagr"], seg["Africa"][fdim][faf]["cagr"]
                mea_fast = (f"{fme} (Middle East {fmt.pct(cme)}; Africa {fmt.pct(caf)})" if fme == faf else
                            f"{fme} (Middle East {fmt.pct(cme)}); {faf} (Africa {fmt.pct(caf)})")
                rows.append([f"Fastest-growing {fdim_label}", f"{fastest['Global']} ({fmt.pct(ft['cagr'])})", mea_fast])
            rows += [["Primary growth catalyst", *two("primary_catalyst")],
                     ["Primary constraint", *two("primary_constraint")]]
            b.table(i, rows)
        elif i == 43:
            b.box(i, "Key conclusions", lst("key_conclusions"), numbered=True)
        # ------------------------------------------------------------ chapter 2
        elif i == 45:
            b.para(i, open_chapter(s("trends_intro"), r"(?i)trend|force|shift",
                                   f"Six forces are reshaping the {root.lower()} market, and each one shows up in the "
                                   f"regional and segment numbers that follow."))
        elif i == 46:
            b.caption(i, "Figure", f"Six trends shaping the {root.lower()} market")
        elif i == 47:
            b.image(i, charts["trends"])
        elif i == 50:
            # every trend shown in the figure is explained in the text
            got = {l.lower().rstrip("."): t for l, t in pairs("trend_paras")}
            for tr in ((ag.get("dynamics") or {}).get("trends") or [])[:6]:
                title = tr.get("title", "")
                text = got.get(title.lower()) or next((t for l, t in got.items() if title.lower()[:12] in l), None) \
                    or _sentence(tr.get("tagline", ""))
                b.lead(i, title.rstrip(".") + ".", text)
        elif i in (51, 52):
            continue
        elif i == 54:
            b.para(i, s("drivers_para"))
        elif i == 56:
            b.caption(i, "Figure", "Impact of growth factors and restraints on MEA CAGR (percentage points)")
        elif i == 57:
            b.image(i, charts["drivers"])
        elif i == 60:
            b.para(i, s("matrix_para"))
        elif i == 61:
            b.caption(i, "Table", "Impact assessment of market factors")
        elif i == 62:
            rows = [[safe(r.get("factor")), r.get("type", ""), r.get("impact", ""), r.get("timeframe", ""),
                     safe(r.get("most_exposed"))] for r in (ag.get("dynamics") or {}).get("impact_matrix") or []]
            b.table(i, rows)
        # ------------------------------------------------------------ chapter 3
        elif i == 64:
            b.out.append(set_h1(b.el(i), title=f"Global {root} Market Size: Scenario Analysis"))
        elif i == 65:
            b.para(i, s("sizing_para", f"CMI's base case values the global {root.lower()} market at "
                                       f"{fmt.usd(g['v26'])} in {Y0}E, reaching {fmt.usd(g['v33'])} in {Y1}F."))
        elif i == 66:
            b.caption(i, "Figure", f"Global {root.lower()} market by region, {Y0 - 1}-{Y1} "
                                   f"(USD {'billion' if g['v33'] >= 1000 else 'million'})")
        elif i == 67:
            b.image(i, charts["regions_stacked"])
        elif i == 70:
            b.caption(i, "Figure", f"Global {root.lower()} market scenarios: conservative, base and optimistic")
        elif i == 71:
            b.image(i, charts["scenarios"])
        elif i == 73:
            b.caption(i, "Table", "Scenario definitions")
        elif i == 74:
            cond = obj("scenario_conditions")
            drv = (m.get("fleet_drivers") or {}).get("drivers") or []
            rst = (m.get("fleet_drivers") or {}).get("restraints") or []
            d0 = max(drv, key=lambda x: x.get("cagr_contribution_bps", 0))["name"].lower() if drv else "core demand"
            r0 = max(rst, key=lambda x: x.get("cagr_drag_bps", 0))["name"].lower() if rst else "structural headwinds"
            sdef = {"Conservative": f"Restraints such as {r0} weigh more heavily than in the base case.",
                    "Base": f"Current drivers, led by {d0}, continue at their present pace.",
                    "Optimistic": f"Faster momentum from {d0}, with easing pressure from {r0}."}
            rows = [[n, fmt.pct(v["cagr"]), fmt.usd(v["series"][Y1]), cond.get(n) or sdef[n]]
                    for n, v in m["scenarios"].items()]
            b.table(i, rows)
        elif i == 81:
            b.para(i, f"3.2 Global {root} Market Size by Segments")
        elif i == 83:
            b.para(i, f"Global Opportunity by {tdim['singular']}")
        elif i == 84:
            b.caption(i, "Figure", f"Global market by {tl}, {Y0}E vs {Y1}F")
        elif i == 85:
            b.image(i, charts["type_Global"])
        elif i == 86:
            b.keep(i)
            for rule in m.get("allocation_rules") or []:          # MECE footnote for overlapping taxonomies
                b.para(27, f"Allocation rule: {_sentence(rule)}")
        elif i == 87:
            b.para(i, seg_para("Global", "type"))
        elif i == 386:
            el = b.el(i)
            lead = "Coherent Market Insights (CMI)"
            rest = b.src[i].xpath("string(.)").strip()[len(lead):].strip()
            practice = s("practice", "Industry Research Practice")
            b.out.append(set_lead(el, lead, f"{rest} This report was prepared by CMI's {practice}."))
        elif i == 88:
            b.para(i, f"Global Opportunity by {ddim['singular']}")
        elif i == 89:
            b.caption(i, "Figure", f"Global market by {dl}, {Y0}E vs {Y1}F (USD)")
        elif i == 90:
            b.image(i, charts["demand_Global"])
        elif i == 92:
            b.para(i, seg_para("Global", "demand"))
        elif i == 94:
            b.para(i, f"3.3 Global {root} Market Size by Region")
        elif i == 95:
            b.caption(i, "Figure", f"Regional share of the global market, {Y0}E vs {Y1}F")
        elif i == 97:
            b.para(i, s("share_para"))
        elif i == 98:
            b.image(i, charts["region_donuts"])
        elif i == 100:
            b.para(i, f"Global: Absolute dollar opportunity by region, {Y0}-{Y1}", True)
        elif i == 102:
            b.para(i, s("abs_para"))
        elif i == 103:
            b.caption(i, "Figure", f"Global: Absolute dollar opportunity by region, {Y0}-{Y1}")
        elif i == 104:
            b.image(i, charts["region_opportunity"])
        elif i == 107:
            hl = obj("highlights")
            for r in ["North America", "Europe", "Asia Pacific", "Latin America", "Middle East & Africa"]:
                d = m["regions"][r]
                text = re.sub(rf"^\s*{re.escape(r)}[.:]\s*", "", hl.get(r) or "") or \
                    f"{fmt.usd(d['v26'])} in {Y0}E, reaching {fmt.usd(d['v33'])} by {Y1}F at a {fmt.pct(d['cagr'])} CAGR."
                b.lead(i, f"{r}.", text)
        elif i in (108, 109, 110, 111):
            continue
        elif i == 112:
            b.para(i, s("country_bridge"))
        elif any(i in t for t in COUNTRY_TABLES.values()):
            region = next(r for r, t in COUNTRY_TABLES.items() if i in t)
            head, intro, tbl = COUNTRY_TABLES[region]
            if i == head:
                b.para(i, f"{region} {root} Market by Country")
            elif i == intro:
                b.para(i, obj("country_intro").get(region) or region_country_fallback(region, m["countries"][region]))
            else:
                b.table(i, country_rows(region)[0])
        elif any(i == t[2] + 1 for t in COUNTRY_TABLES.values()):
            # the table's source line, then a short profile of every country in the table
            region = next(r for r, t in COUNTRY_TABLES.items() if i == t[2] + 1)
            b.keep(i)
            if region in m["profiled"]:
                # full profiles with segment charts follow in Chapters 4/5
                b.para(24, f"Every country in this table, and the Rest of {region}, is profiled in full, with its "
                           f"segment breakdown, in the {region} chapter.")
                continue
            profiles = obj("country_profiles")
            why = m.get("geography_why") or {}
            for c, v in m["countries"][region].items():
                size_line = (f"{fmt.usd(v['v26'])} in {Y0}E (CMI estimate), forecast to grow at a {fmt.pct(v['cagr'])} "
                             f"CAGR to {fmt.usd(v['v33'])} by {Y1}F.")
                if profiles.get(c) and "$" not in profiles[c]:
                    # a profile always states the country's size and growth
                    profiles[c] = size_line + " " + profiles[c]
                text = profiles.get(c) or (
                    f"{fmt.usd(v['v26'])} in {Y0}E, growing at a {fmt.pct(v['cagr'])} CAGR to {fmt.usd(v['v33'])} "
                    f"by {Y1}F (CMI estimate). " + (safe(why.get(c)) or ""))
                b.lead(29, f"{c}:", text.strip())
        # ------------------------------------------------------------ chapters 4 and 5
        elif i in (146, 226):
            sub = "Middle East" if i == 146 else "Africa"
            b.out.append(set_h1(b.el(i), title=f"{sub} {root} Market: Opportunity Landscape"))
        elif i in (147, 227):
            sub, ch = ("Middle East", "ch4") if i == 147 else ("Africa", "ch5")
            d = m["subregions"][sub]
            lead_in = (f"The {sub} {root.lower()} market is estimated at {fmt.usd(d['v26'])} in {Y0}E and is "
                       f"forecast to reach {fmt.usd(d['v33'])} by {Y1}F ({fmt.pct(d['cagr'])} CAGR), CMI estimates.")
            other = "Africa" if sub == "Middle East" else "Middle East"
            # opens on its own sub-region by name (not 'South Africa'), not on the other sub-region
            b.para(i, open_chapter(s(f"{ch}.region_para"),
                                   rf"(?<!South )(?<!Rest of )\b{re.escape(sub)}\b(?! &)", lead_in,
                                   avoid=rf"(?<!South )(?<!Rest of )\b{re.escape(other)}\b(?! &)"))
        elif i in (148, 228):
            sub = "Middle East" if i == 148 else "Africa"
            unit = "billion" if m["subregions"][sub]["v33"] >= 1000 else "million"
            b.caption(i, "Figure", f"{sub} {root.lower()} market by country, {Y0 - 1}-{Y1} (USD {unit})")
        elif i in (149, 229):
            b.image(i, charts["me_stacked" if i == 149 else "af_stacked"])
        elif i in (152, 232):
            b.caption(i, "Table", f"{'Middle East' if i == 152 else 'Africa'} {root.lower()} market by country (USD)")
        elif i in (153, 233):
            rows, total = country_rows("Middle East" if i == 153 else "Africa")
            b.table(i, rows, total)
        elif i in (155, 235):
            b.caption(i, "Figure", "Country attractiveness: growth vs absolute opportunity")
        elif i in (156, 236):
            b.image(i, charts["me_bubble" if i == 156 else "af_bubble"])
        elif i in (158, 238):
            b.para(i, f"{'4.2 Middle East' if i == 158 else '5.2 Africa'} Opportunity by {tl}")
        elif i in (159, 239):
            b.caption(i, "Figure", f"{'Middle East' if i == 159 else 'Africa'} market by {tl}, {Y0}E vs {Y1}F")
        elif i in (160, 240):
            b.image(i, charts["type_Middle East" if i == 160 else "type_Africa"])
        elif i in (162, 242):
            b.para(i, seg_para("Middle East" if i == 162 else "Africa", "type"))
        elif i in (163, 243):
            b.para(i, f"{'4.3 Middle East' if i == 163 else '5.3 Africa'} Opportunity by {dl}")
        elif i in (164, 244):
            b.caption(i, "Figure", f"{'Middle East' if i == 164 else 'Africa'} market by {dl}, {Y0}E vs {Y1}F (USD)")
        elif i in (165, 245):
            b.image(i, charts["demand_Middle East" if i == 165 else "demand_Africa"])
        elif i in (167, 247):
            b.para(i, seg_para("Middle East" if i == 167 else "Africa", "demand"))
        elif i in (168, 248):
            b.para(i, f"{'4.4 Middle East' if i == 168 else '5.4 Africa'} {root} Market: Country landscape")
        elif i in PROFILE_STARTS:
            # one full profile per named country + the 'Rest of' group, each cloned from the
            # template's first profile block (12 elements); the other template blocks are dropped
            sub = PROFILE_STARTS[i]
            for k, unit in enumerate(m["profiled"][sub]):
                if k and sub == "Middle East":
                    b.keep(181)                                     # page break between profiles
                d = m["countries"][sub][unit]
                for off in range(12):
                    j = i + off
                    if off in range(2, 7) and not sec["type"]["show"] or off in range(7, 12) and not sec["demand"]["show"]:
                        continue
                    if off == 0:
                        b.para(j, unit)
                    elif off == 1:
                        size_line = (f"{unit} is a {fmt.usd(d['v26'])} market in {Y0}E (CMI estimate), forecast to "
                                     f"grow at a {fmt.pct(d['cagr'])} CAGR to {fmt.usd(d['v33'])} by {Y1}F.")
                        prof = s(f"profile__{unit}")
                        # a profile always states the country's size and growth
                        b.para(j, prof if prof and "$" in prof else (size_line + (" " + prof if prof else "")))
                    elif off == 2:
                        b.para(j, f"{unit} Opportunity by {tl}")
                    elif off == 3:
                        b.caption(j, "Figure", f"{unit} market by {tl}, {Y0}E vs {Y1}F")
                    elif off == 4:
                        b.image(j, charts[f"type_{unit}"])
                    elif off == 6:
                        b.para(j, seg_para(unit, "type"))
                    elif off == 7:
                        b.para(j, f"{unit} Opportunity by {dl}")
                    elif off == 8:
                        b.caption(j, "Figure", f"{unit} market by {dl}, {Y0}E vs {Y1}F (USD)")
                    elif off == 9:
                        b.image(j, charts[f"demand_{unit}"])
                    elif off == 11:
                        b.para(j, seg_para(unit, "demand"))
                    else:
                        b.keep(j)
        elif i in PROFILE_RANGE:
            continue                                                # replaced by the cloned blocks above
        elif i in (221, 298):
            b.para(i, f"{'4.5 Middle East' if i == 221 else '5.5 Africa'} {root} Market: Use-case opportunity heat map")
        elif i in (222, 299):
            b.caption(i, "Table", "Use-case attractiveness by country (CMI assessment)")
        elif i in (223, 300):
            key = "middle_east" if i == 223 else "africa"
            # columns = this market's named countries (template slots: 4 + Rest for ME, 5 for Africa)
            cols = [short(c) for c in m["geography"]["Middle East"]] + ["Rest of ME"] if i == 223 else \
                [short(c) for c in m["geography"]["Africa"]]
            rows = [[r.get("use_case", "")] + [r.get(c, "") for c in cols]
                    for r in (ag.get("use_cases") or {}).get(key) or []]
            b.table(i, complete_rows(rows, b.log, "use-case heat map"), header=["Use case"] + cols)
        elif i in (225, 302):
            ch = "ch4" if i == 225 else "ch5"
            items = [(lead.rstrip(":") + ":", text) for lead, text in pairs(f"{ch}.where_to_play")]
            b.box(i, f"Where to play in {'the Middle East' if i == 225 else 'Africa'}", items)
        # ------------------------------------------------------------ chapter 6
        elif i == 304:
            b.para(i, open_chapter(s("reg_intro"), r"(?i)regulat|polic|approv|registr|licens|authorit",
                                   f"Regulation decides how quickly {root.lower()} products reach buyers in the Middle "
                                   f"East & Africa: which authority approves them, how long registration takes and "
                                   f"what is changing next."))
        elif i == 305:
            b.caption(i, "Figure", f"Regulatory milestones shaping the {root.lower()} market")
        elif i == 306:
            b.image(i, charts["timeline"])
        elif i == 307:
            b.para(i, "Source: Public regulatory notices; CMI analysis")
        elif i == 308:
            b.caption(i, "Figure", "Regulatory enablement scorecard by country")
        elif i == 309:
            b.image(i, charts["reg_scorecard"])
        elif i == 312:
            b.caption(i, "Table", "Regulatory snapshot")
        elif i == 313:
            rows = [[r.get("country", ""), r.get("regulator", "") + ("" if r.get("verified") else " [to verify]"),
                     safe(r.get("requirements"))]
                    for r in (ag.get("regulation") or {}).get("snapshot") or []
                    if canonical(r.get("country", "")) in model_reg_countries]
            b.table(i, complete_rows(rows, b.log, "regulatory snapshot"))
        elif i == 316:
            b.bullets(i, lst("policy_bullets"))
        elif i in (317, 318, 319):
            continue
        elif i == 320:
            b.para(i, f"6.3 Regulatory watchlist {Y0}-{Y0 + 2}")
        elif i == 321:
            b.caption(i, "Table", "Regulatory watchlist")
        elif i == 322:
            rows = [[r.get("theme", ""), safe(r.get("what_to_watch")), safe(r.get("implication"))]
                    for r in (ag.get("regulation") or {}).get("watchlist") or []]
            b.table(i, complete_rows(rows, b.log, "regulatory watchlist"))
        elif i == 324:
            b.box(i, "Implication for market entry", [s("entry_implication", "")] if s("entry_implication") else [])
        # ------------------------------------------------------------ chapter 7
        elif i == 326:
            b.para(i, open_chapter(s("vendor_intro"), r"(?i)vendor|compan|player|supplier|tier|producer",
                                   f"The {root.lower()} vendor landscape is set out below by value-chain position, "
                                   + ("market share and capability." if sec["vendor_global"]["show"]
                                      else "market position and capability.")))
        elif i == 328:
            b.caption(i, "Figure", f"{root} value chain and indicative share of value pool")
        elif i == 329:
            b.image(i, charts["value_chain"])
        elif i == 331:
            b.para(i, s("value_chain_para"))
        elif 332 <= i <= 341:
            vendor_shares(b, i, m, s, root, charts, vsafe)
        elif i == 344:
            b.caption(i, "Figure", "Capability heat map of leading vendors in MEA")
        elif i == 345:
            b.image(i, charts["capabilities"])
        elif i == 348:
            b.caption(i, "Table", "Vendor dashboard")
        elif i == 349:
            # vendors with a confirmed HQ only; HQ and ownership note from the entity check (PLAN-05)
            rows = []
            for r in m["vendor_dashboard"]:
                own = short_note(r["ownership_note"])
                foot = vsafe(r["mea_footprint"])
                rows.append([r["vendor"] + (f" ({own})" if own and own.lower() not in r["vendor"].lower() else ""),
                             r["hq"], vsafe(r["core_strength"]),
                             (foot + ("" if r["verified"] else " [to verify]")) if foot else ""])
            b.table(i, complete_rows(rows, b.log, "vendor dashboard"))
        elif i == 351:
            b.box(i, "Competitive implications", lst("implications"), numbered=True)
        # ------------------------------------------------------------ chapter 8
        elif i == 353:
            b.para(i, s("reco_intro") or "The recommendations below translate the report's findings into actions for "
                                         "each stakeholder group; every priority action cites the finding that "
                                         "justifies it.")
        elif i == 354:
            b.caption(i, "Figure", "Three-phase way forward for MEA entry and scale-up")
        elif i == 355:
            b.image(i, charts["phases"])
        elif i == 358:
            b.caption(i, "Table", "Strategic recommendations by stakeholder")
        elif i == 359:
            rows = [[r.get("stakeholder", ""), safe(r.get("priority_moves")), safe(r.get("kpis"))]
                    for r in (ag.get("strategy") or {}).get("stakeholders") or []]
            b.table(i, complete_rows(rows, b.log, "stakeholder recommendations"))
        elif i == 361:
            # B9: each action carries the finding it follows from (written after the body)
            traced = [(lead.rstrip(".") + ".", text) for lead, text in pairs("priority_actions")
                      if text.startswith("Finding")]
            acts = traced or [x for x in (safe(a) for a in (ag.get("strategy") or {}).get("priority_actions") or []) if x]
            acts = acts[:5]
            # the title states how many actions survived the guards ('Three priority actions ...')
            b.box(i, f"{_COUNT.get(len(acts), str(len(acts)))} priority actions for {Y0}-{str(Y0 + 1)[2:]}",
                  acts, numbered=True)
        # ------------------------------------------------------------ chapter 9
        elif i == 363:
            b.para(i, open_chapter(s("method_para"), r"(?i)\bCMI\b|estimat|method|model",
                                   f"CMI's {root.lower()} estimates combine top-down sizing from regional and national "
                                   f"totals with bottom-up build-ups by segment and country."))
        elif i == 364:
            b.caption(i, "Figure", "CMI five-step research process")
        elif i == 365:
            b.image(i, charts["process"])
        elif i == 367:
            b.para(i, "9.1 Basis of the headline estimate")
        elif i == 368:
            b.para(i, size_basis_text(m, root))
        elif i in (369, 370, 371):
            pub = m["published_2025"]
            if not pub:
                continue
            if i == 369:
                b.caption(i, "Figure", f"Published {Y0 - 1} global {root.lower()} market estimates vs CMI base")
            elif i == 370:
                b.image(i, charts["published"])
            else:
                b.para(i, "Source: " + ", ".join(dict.fromkeys(p["publisher"] for p in pub)) + "; CMI analysis")
        elif i in (372, 373, 374, 375):
            # reference figures only beside a sourced headline; they never set the CMI numbers
            refs = (ag.get("published") or {}).get("references") or [] if not size_unsourced else []
            if not refs:
                continue
            if i == 373:
                b.caption(i, "Table", "Third-party reference figures (for reference only; not used in the model)")
            elif i == 374:
                b.table(i, [[r["name"], r.get("scope", ""), r["published_value"]] for r in refs])
            else:
                b.keep(i)
        elif i == 377:
            b.caption(i, "Table", "Scope of the study")
        elif i == 378:
            dfn = m["definition"]
            rows = [["Market measured", f"{dfn.get('measures') or dfn['definition']} ({dfn['unit']})"]] \
                if dfn.get("measures") or dfn.get("definition") else []
            if dfn.get("value_chain_level"):
                rows.append(["Value-chain level", dfn["value_chain_level"]])
            rows += [[(tdim if d == "type" else ddim)["singular"], "; ".join(seg["Global"][d])]
                     for d in ("type", "demand") if sec[d]["show"]]
            if dfn.get("included"):
                rows.append(["Inclusions", dfn["included"]])
            if dfn.get("excluded"):
                rows.append(["Exclusions", dfn["excluded"]])
            # exactly the countries covered in the body
            gl = []
            for r in ("North America", "Europe", "Asia Pacific", "Latin America", "Middle East", "Africa"):
                named = [c for c in m["countries"][r] if not c.startswith("Rest of")]
                rest = [c for c in m["countries"][r] if c.startswith("Rest of")]
                gl.append(f"{r} ({', '.join(named + rest)})")
            rows.append(["Geography", "; ".join(gl)])
            for d in ("type", "demand"):
                for n, txt in (m.get("segment_definitions") or {}).get(d, {}).items():
                    if sec[d]["show"] and n in seg["Global"][d]:
                        rows.append([f"{(tdim if d == 'type' else ddim)['singular']}: {n}", txt])
            for rule in m.get("allocation_rules") or []:
                rows.append(["Allocation rule", rule])
            rows.append(["Years", f"Base {Y0 - 1}; estimate {Y0}; forecast {Y0 + 1}-{Y1}"])
            b.table(i, rows)
        elif i == 381:
            b.bullets(i, lst("assumptions")[:4])
        elif i in (382, 383, 384):
            continue
        # ------------------------------------------------------------ captions/sources kept as-is
        else:
            text = b.src[i].xpath("string(.)").strip()
            m_cap = re.match(r"^(Figure|Table) \d+:", text)
            if m_cap:
                b.caption(i, m_cap.group(1), text.split(":", 1)[1].strip())
            else:
                b.keep(i)

    toc_titles[11] = "Appendix: Consistency & Verification Log"
    add_appendix(b, m, ag, prose, removed_notes, verification, register)
    renumber(b, toc_titles, toc_pages)
    b.save(out_path, f"The {root} Market Outlook {Y0}-{Y1}", "Whitepaper")
    return removed_notes + b.log + [f"guard dropped: {d}" for d in dropped]
