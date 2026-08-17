"""Export a generated Input Sheet into a single hierarchical JSON:
Global revenue at the top, then by region/continent, then countries and
segment (sub-segment) shares within each region.

Reads only the already-generated Input Sheet .xlsx (Data + Market Value
sheets) — doesn't touch the live Desk Research pipeline/state, so it works on
any run's output, past or future, as long as the file exists.

Shape:
    {
      "market": "...",
      "global": {"revenue_usd_mn": ..., "cagr": ...},
      "regions": {
        "<Region>": {
          "revenue_usd_mn": ..., "share_of_global": ..., "cagr": ...,
          "countries": {"<Country>": {"revenue_usd_mn": ..., "share_of_region": ..., "cagr": ...}, ...},
          "segments": {"<Dimension>": {"<Segment>": <value-weighted avg share across the region's countries>, ...}, ...}
        },
        ...
      }
    }

Example:
    python export_market_json.py --market "Global Avocado Oil Market"
    python export_market_json.py --input "agentGeneratedInputSheet/Data File - Global Smartphone Market.xlsx"
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
INPUT_DIR = ROOT / "agentGeneratedInputSheet"
OUT_DIR = ROOT / "output" / "market_json"


def _grid(ws) -> list[tuple]:
    return list(ws.iter_rows(values_only=True))


def _find_input_sheet(market: str) -> Path:
    slug = re.sub(r"[^a-z0-9]+", "", market.lower())
    candidates = sorted(INPUT_DIR.glob("Data File*.xlsx"))
    for p in candidates:
        if "[pipeline]" in p.stem:
            continue
        if re.sub(r"[^a-z0-9]+", "", p.stem.lower()).endswith(slug):
            return p
    raise FileNotFoundError(f"No Input Sheet found for market '{market}' in {INPUT_DIR}")


def _parse_market_value(rows: list[tuple]) -> tuple[dict, dict[str, dict]]:
    """Returns (global_info, {region_name: {value, share, cagr}})."""
    header_idx = next(i for i, r in enumerate(rows) if r and r[1] == "Regions")
    regions: dict[str, dict] = {}
    global_info: dict = {}
    for row in rows[header_idx + 1:]:
        if not row or not row[1]:
            continue
        label = str(row[1]).strip()
        if label == "Global":
            global_info = {"revenue_usd_mn": row[2], "cagr": row[4]}
            continue
        regions[label] = {"revenue_usd_mn": row[2], "share_of_global": row[3], "cagr": row[4]}
    return global_info, regions


def _parse_geography_block(rows: list[tuple], region_names: set[str]) -> dict[str, dict[str, dict]]:
    """Returns {region_name: {country_name: {value, share_of_region, cagr}}}."""
    start = next((i for i, r in enumerate(rows) if r and r[0] and "100th Row" in str(r[0])), None)
    if start is None:
        return {}
    out: dict[str, dict[str, dict]] = {}
    current_region: str | None = None
    for row in rows[start + 1:]:
        if not row or not row[1]:
            if row and any(c is not None for c in row):
                continue
            break
        label = str(row[1]).strip()
        if label == "Global":
            continue
        share, cagr, value = row[2], row[3], row[4]
        if label in region_names:
            current_region = label
            out[current_region] = {}
            continue
        if current_region is None:
            continue
        out[current_region][label] = {"revenue_usd_mn": value, "share_of_region": share, "cagr": cagr}
    return out


def _parse_segments(rows: list[tuple]) -> tuple[list[str], dict[str, dict[str, dict[str, float]]]]:
    """Returns (geography_columns, {dimension: {segment: {country: share}}})."""
    header_idx = next(i for i, r in enumerate(rows) if r and r[0] == "Segment Name in B Column Only")
    header = rows[header_idx]
    geos = [str(g) for g in header[2:] if g and str(g) != "Confidence"]

    dims: dict[str, dict[str, dict[str, float]]] = {}
    current_dim: str | None = None
    for row in rows[header_idx + 1:]:
        if not row or row[1] is None:
            current_dim = None
            continue
        label = str(row[1]).strip()
        values = row[2:2 + len(geos)]
        if all(v is None for v in values):
            current_dim = label
            dims[current_dim] = {}
            continue
        if current_dim is None:
            continue
        dims[current_dim][label] = {geo: v for geo, v in zip(geos, values) if v is not None}
    return geos, dims


def _region_segments(
    dims: dict[str, dict[str, dict[str, float]]],
    countries_in_region: dict[str, dict],
) -> dict[str, dict[str, float]]:
    """Value-weighted average of each country's segment share, weighted by
    that country's revenue within the region."""
    weights = {c: info["revenue_usd_mn"] or 0.0 for c, info in countries_in_region.items()}
    total_weight = sum(weights.values())
    out: dict[str, dict[str, float]] = {}
    for dim, segments in dims.items():
        out[dim] = {}
        for seg, by_country in segments.items():
            if total_weight <= 0:
                continue
            weighted = sum(by_country.get(c, 0.0) * w for c, w in weights.items())
            out[dim][seg] = round(weighted / total_weight, 4)
    return out


def export_market_json(input_path: Path) -> dict:
    wb = load_workbook(input_path, data_only=True, read_only=True)
    data_rows = _grid(wb["Data"])
    market_name = data_rows[0][2] if data_rows and len(data_rows[0]) > 2 else input_path.stem

    global_info, mv_regions = _parse_market_value(_grid(wb["Market Value"])) if "Market Value" in wb.sheetnames else ({}, {})
    geo_regions = _parse_geography_block(data_rows, set(mv_regions.keys()))
    _, dims = _parse_segments(data_rows)

    regions_out: dict[str, dict] = {}
    for region, info in mv_regions.items():
        countries = geo_regions.get(region, {})
        regions_out[region] = {
            **info,
            "countries": countries,
            "segments": _region_segments(dims, countries),
        }

    return {
        "market": market_name,
        "global": global_info,
        "regions": regions_out,
    }


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--market", default=None, help="Market name — auto-finds its Input Sheet in agentGeneratedInputSheet/")
    ap.add_argument("--input", default=None, help="Explicit path to an Input Sheet .xlsx")
    ap.add_argument("--save-to", default=None, help="Output path (default: output/market_json/<slug>.json)")
    args = ap.parse_args()

    if args.input:
        input_path = Path(args.input)
    elif args.market:
        input_path = _find_input_sheet(args.market)
    else:
        print("[ERROR] Pass --market or --input.")
        return 1

    if not input_path.exists():
        print(f"[ERROR] Input Sheet not found: {input_path}")
        return 1

    result = export_market_json(input_path)

    slug = re.sub(r"[^a-z0-9]+", "_", result["market"].lower()).strip("_")
    out = Path(args.save_to) if args.save_to else OUT_DIR / f"{slug}.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")

    print(f"Market: {result['market']}")
    print(f"Global revenue: {result['global'].get('revenue_usd_mn')} USD Mn")
    for region, info in result["regions"].items():
        print(f"  {region}: {info['revenue_usd_mn']} USD Mn ({len(info['countries'])} countries, "
              f"{len(info['segments'])} segment dimensions)")
    print(f"\nSaved to {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
