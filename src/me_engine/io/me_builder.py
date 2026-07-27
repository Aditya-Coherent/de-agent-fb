"""Build an ME workbook from scratch — no template/reference file required.

`MEWorkbookWriter` (me_writer.py) can only fill in cells whose (sheet, band,
label) already exist in a borrowed template — silently leaving every other
cell as the template's own stale data when the market doesn't match (see
`TemplateMismatchError`). This builder is the real fix for a genuinely NEW
market: it creates every sheet, row and column directly from the assembled
`MarketResult` + the market's own `RuntimeTaxonomy`, so nothing is copied from
anywhere else.

Two things the Assembler itself doesn't compute are synthesized here, purely
from data `MarketResult` already contains — no assembler change needed:
  - a dimension "total" row (e.g. "By Product Type") — the geography's own
    band total, redisplayed once per dimension before its segment breakdown
    (matches the source convention: the parent total is shown, then split).
  - a "By Country" rollup block for any geography with children in the
    taxonomy's geography tree — one row per direct child, taken from that
    child's OWN Value-band total (all children already live in the same
    `MarketResult.geographies`).
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from ..assembly.model import BandResult, MarketResult, MetricRow
from ..domain.runtime_taxonomy import RuntimeTaxonomy
from ..domain.taxonomy import BASE_YEAR, Band, YEARS
from .layout import COL_CAGR, COL_LABEL, COL_METRIC_START, COL_SHARE_START, COL_YOY_START
from .me_writer import MEWorkbookWriter

_BOLD = Font(bold=True)
_INVALID_SHEET_CHARS = set('[]:*?/\\')


class MEWorkbookBuilder:
    """Builds a fresh ME workbook straight from a MarketResult — no
    pre-existing reference/template file required."""

    def __init__(self, taxonomy: RuntimeTaxonomy) -> None:
        self._taxonomy = taxonomy

    def build(self, result: MarketResult, out_path: Path | str) -> Path:
        wb = Workbook()
        wb.remove(wb.active)  # drop the default blank sheet

        for geo_node in self._taxonomy.geographies.in_order:
            geo = result.geographies.get(geo_node.name)
            if geo is None:
                continue
            ws = wb.create_sheet(title=self._safe_sheet_name(geo_node.name, wb))
            self._write_geography(ws, geo_node.name, result)

        out = Path(out_path)
        wb.save(out)
        return out

    @staticmethod
    def _safe_sheet_name(name: str, wb: Workbook) -> str:
        safe = "".join(c for c in name if c not in _INVALID_SHEET_CHARS)[:31] or "Sheet"
        base, i = safe, 2
        while safe in wb.sheetnames:
            suffix = f" {i}"
            safe = safe[: 31 - len(suffix)] + suffix
            i += 1
        return safe

    # --- per-geography sheet -------------------------------------------------
    def _write_geography(self, ws: Worksheet, geo_name: str, result: MarketResult) -> None:
        geo = result.geographies[geo_name]
        value_band = geo.bands.get(Band.VALUE)
        volume_band = geo.bands.get(Band.VOLUME)

        if value_band is not None:
            self._write_summary_header(ws, 1, "Value & CAGR", value_band.total)
        if volume_band is not None:
            self._write_summary_header(ws, 2, "Volume & CAGR", volume_band.total)

        row = 4
        for band_kind in (Band.VALUE, Band.ASP, Band.VOLUME):
            band = geo.bands.get(band_kind)
            if band is None or not band.rows_by_label:
                continue
            row = self._write_band_header(ws, row, band_kind)
            if band_kind is Band.VALUE:
                row = self._write_value_band(ws, row, band, geo_name, result)
            else:
                row = self._write_priced_band(ws, row, band)
            row += 1  # blank separator row between bands

    @staticmethod
    def _write_summary_header(ws: Worksheet, row: int, label: str, total) -> None:
        ws.cell(row, COL_LABEL, label).font = _BOLD
        ws.cell(row, COL_METRIC_START, total.at(BASE_YEAR))
        ws.cell(row, COL_METRIC_START + 1, total.cagr())

    @staticmethod
    def _write_band_header(ws: Worksheet, row: int, band_kind: Band) -> int:
        ws.cell(row, COL_METRIC_START, band_kind.value).font = _BOLD
        row += 1
        for year in YEARS:
            ws.cell(row, COL_METRIC_START + YEARS.index(year), year)
        for year in YEARS[1:]:
            ws.cell(row, COL_YOY_START + (year - YEARS[1]), year)
        ws.cell(row, COL_CAGR, "CAGR")
        for year in YEARS:
            ws.cell(row, COL_SHARE_START + YEARS.index(year), year)
        return row + 1

    def _write_value_band(self, ws: Worksheet, row: int, band: BandResult,
                          geo_name: str, result: MarketResult) -> int:
        """Each segmentation dimension: a synthesised total row + its segments."""
        for dim in self._taxonomy.segmentation_dimensions:
            present = [s for s in dim.segments if s in band.rows_by_label]
            if not present:
                continue
            row = self._write_row(ws, row, dim.title, MetricRow(dim.title, band.total, None), bold=True)
            for segment in present:
                row = self._write_row(ws, row, segment, band.rows_by_label[segment])
            row += 1

        children = self._taxonomy.geographies.by_name[geo_name].children
        child_rows = [
            (c.name, result.geographies[c.name].bands[Band.VALUE].total)
            for c in children if c.name in result.geographies
        ]
        if child_rows:
            row = self._write_row(ws, row, "By Country", MetricRow("By Country", band.total, None), bold=True)
            for child_name, child_total in child_rows:
                row = self._write_row(ws, row, child_name, MetricRow(child_name, child_total, None))
            row += 1
        return row

    def _write_priced_band(self, ws: Worksheet, row: int, band: BandResult) -> int:
        """ASP/Volume: leaf priced-dimension products only (matches what the
        Assembler computes for these bands), with the same total-row convention."""
        priced = self._taxonomy.priced_dimension
        present = [s for s in priced.segments if s in band.rows_by_label]
        if not present:
            return row
        row = self._write_row(ws, row, priced.title, MetricRow(priced.title, band.total, None), bold=True)
        for segment in present:
            row = self._write_row(ws, row, segment, band.rows_by_label[segment])
        return row

    @staticmethod
    def _write_row(ws: Worksheet, row: int, label: str, metric_row: MetricRow,
                   bold: bool = False) -> int:
        cell = ws.cell(row, COL_LABEL, label)
        if bold:
            cell.font = _BOLD
        MEWorkbookWriter._write_metric(ws, row, metric_row)
        MEWorkbookWriter._write_yoy(ws, row, metric_row)
        MEWorkbookWriter._write_cagr(ws, row, metric_row)
        MEWorkbookWriter._write_share(ws, row, metric_row)
        return row + 1
