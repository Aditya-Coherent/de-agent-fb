"""Build the demonstration web app: an Excel-like view of the generated Input
Sheet where every cell opens a document explaining how the number was made
and why it differs from the human analyst's.

    python build_demo_app.py --tag v10            # -> demo/high-energy-laser-v10/index.html
    python build_demo_app.py --tag v10 --serve    # and open it on http://localhost:8765

Everything shown comes from the run's own files (no new estimation):
    output/hel/Data File - <market> (<tag>, human structure).xlsx   final numbers
    agentGeneratedInputSheet/Data File - <market> (<tag>).xlsx      the generator's raw numbers
    output/hel/regulation_<tag>.json, events_<tag>.json, evaluation_<tag>.json, gen_state_<tag>.json
    output/hel/structure_spec.json, assumptor/wb_drivers.json, the human Data file
"""
from __future__ import annotations

import argparse
import json
import math
import re
import statistics
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from me_engine.generation import assumptor_regulator as ar  # noqa: E402
from me_engine.generation import data_writer as dw  # noqa: E402
from me_engine.generation import evaluator_agent as ev  # noqa: E402
from me_engine.generation.structure_agent import _norm  # noqa: E402

TEMPLATE = ROOT / "demo" / "template.html"


def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def load_json(p: Path):
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else None


def grid(wb, name):
    ws = wb[name]
    return [list(r) for r in ws.iter_rows(values_only=True)], ws.max_row, ws.max_column


def build(market: str, human_data: Path, out_dir: Path, tag: str, gen_tag: str | None = None,
          run_dir: Path | None = None, layout_only: bool = False) -> Path:
    """`layout_only`: `human_data` is a blank template built from an ME's labels
    (run_market_e2e --me-structure) -- its numbers are layout markers, never
    shown or compared."""
    hel = run_dir or ROOT / "output" / "hel"
    spec = load_json(hel / "structure_spec.json")
    final_path = hel / f"Data File - {market} ({tag}, human structure).xlsx"
    gen_path = ROOT / "agentGeneratedInputSheet" / f"Data File - {market} ({gen_tag or tag}).xlsx"
    reg = load_json(hel / f"regulation_{tag}.json") or {}
    events = load_json(hel / f"events_{tag}.json") or {}
    evaluation = load_json(hel / f"evaluation_{tag}.json") or {}
    state = load_json(hel / f"gen_state_{gen_tag or tag}.json") or {}
    usecases = load_json(hel / f"usecases_{tag}.json")
    geos = spec["geographies"]
    regions = spec["regions"]
    region_of = {c: r for r, cs in regions.items() for c in cs}

    hwb = load_workbook(human_data, data_only=True)
    fwb = load_workbook(final_path, data_only=True)
    gen = ev.read_data_file(gen_path, spec)
    gen_mv = dw._market_value(gen_path)
    gen_asp = dw._gen_asp(gen_path, geos)
    hd = ev.read_data_file(human_data, spec)
    if layout_only:
        hd = {**hd, "shares": {}, "region_value": {}}
    final_shares = ev.read_data_file(final_path, spec)["shares"]

    # ---- real-world driver splits the assumptor used (recomputed from its inputs)
    drivers = load_json(hel / "assumptor" / "wb_drivers.json") or {}
    dim = reg.get("dimension")
    seg_values = {s: a["value_mn"] for s, a in (reg.get("segments") or {}).items()}
    driver_alloc, driver_country = {}, {g: 0.0 for g in geos}
    for s, a in (reg.get("segments") or {}).items():
        by_geo = ar.driver_by_geo(drivers, geos, a["driver"]) if drivers else {}
        tot = sum(by_geo.values())
        driver_alloc[s] = {g: (a["value_mn"] * v / tot if tot else 0.0) for g, v in by_geo.items()}
        for g, v in driver_alloc[s].items():
            driver_country[g] += v
    driver_mix = {g: {s: (driver_alloc[s].get(g, 0) / driver_country[g] if driver_country[g] else None)
                      for s in seg_values} for g in geos}
    dct = sum(driver_country.values()) or 1.0

    # human global split of each partition segment (value-weighted by the human's country values)
    hw = {g: hd["region_value"].get(g, 0.0) for g in geos}
    hwt = sum(hw.values()) or 1.0
    human_global_mix = {s: sum(hw[g] * (hd["shares"].get(((dim, (s,)), g)) or 0) for g in geos) / hwt for s in seg_values}
    our_global_mix = {s: v / (sum(seg_values.values()) or 1) for s, v in seg_values.items()}

    # dispersion of shares across countries (human analysts tend to vary shares little)
    def disp(shares, key):
        vals = [shares.get((key, g)) for g in geos if shares.get((key, g)) is not None]
        return round(100 * statistics.pstdev(vals), 2) if len(vals) > 2 else None

    # systematic gap per segment: in how many countries is ours above the analyst's, and by how much
    seg_bias = {}
    for key in {k for (k, g) in hd["shares"]}:
        diffs = [gen_final - hd["shares"][(key, g)] for g in geos
                 if (key, g) in hd["shares"] and (gen_final := final_shares.get((key, g))) is not None]
        if diffs:
            seg_bias[key] = (round(100 * statistics.mean(diffs), 1), sum(d > 0 for d in diffs), len(diffs))

    cons_dims = set()
    for note in (reg.get("consistency") or {}).get("notes", []):
        cons_dims.add(note.split(":")[0])
    llm_cons = (reg.get("consistency") or {}).get("llm") or {}
    guard = reg.get("guard_flags") or []
    ev_scan = events.get("scan") or {}
    ev_applied = events.get("applied") or {}
    ev_changes = {c["geo"]: c for c in ev_applied.get("changes") or []}

    def events_for(geo):
        out = []
        for e in ev_scan.get("events") or []:
            aff = e["affects"]
            if "Global" in aff or geo in aff or region_of.get(geo) in aff or (geo in regions and any(
                    c in aff for c in regions[geo])):
                out.append(e)
        return out

    sheets = {}
    # ------------------------------------------------------------------ Data
    hg, hrows, hcols = grid(hwb, "Data")
    fg, _, _ = grid(fwb, "Data")
    labels = [_norm(r[1]) if len(r) > 1 else "" for r in hg]
    paths = ev._path_walk(labels, spec)
    hdr = hg[2]
    geo_col = {c: _norm(hdr[c - 1]) for c in range(3, 29) if c - 1 < len(hdr) and _norm(hdr[c - 1]) in geos}
    growth_col = {}
    for c in range(31, hcols + 1):
        g = _norm(hdr[c - 1]) if c - 1 < len(hdr) else ""
        if g in geos and g not in growth_col.values():
            growth_col[c] = g
    trend_col = {c: _norm(hdr[c - 1]) for c in range(58, hcols + 1) if c - 1 < len(hdr) and _norm(hdr[c - 1]) in geos}
    region_start = next((i for i, lab in enumerate(labels, start=1) if lab == "Region"), 10**9)
    dim_of_row, current = {}, None
    for i, lab in enumerate(labels, start=1):
        if lab in spec["dimensions"]:
            current = lab
        if i < region_start:
            dim_of_row[i] = current
    cells = []
    for r in range(1, hrows + 1):
        for c in range(1, hcols + 1):
            h = hg[r - 1][c - 1] if c - 1 < len(hg[r - 1]) else None
            v = fg[r - 1][c - 1] if r - 1 < len(fg) and c - 1 < len(fg[r - 1]) else None
            if layout_only and isinstance(h, (int, float)) and not isinstance(h, bool):
                h = None
            if h is None and v is None:
                continue
            cell = {"r": r, "c": c, "h": h if not isinstance(h, float) else round(h, 6),
                    "v": v if not isinstance(v, float) else round(v, 6), "k": "label"}
            p = paths[r - 1] if r - 1 < len(paths) else None
            if r > region_start:
                name = labels[r - 1]
                if c == 3:
                    cell["k"] = "geo_share"
                elif c == 4:
                    cell["k"] = "geo_cagr"
                elif c == 5:
                    cell["k"] = "geo_value"
                if "k" in cell:
                    cell["geo"] = name
                    cell["gen"] = (gen["region_share"].get(name) if c == 3 else gen["region_cagr"].get(name)
                                   if c == 4 else gen["region_value"].get(name))
                    if name in geos:
                        cell["drv_share_global"] = driver_country.get(name, 0) / dct
                        reg_name = region_of.get(name)
                        rtot = sum(driver_country.get(x, 0) for x in regions.get(reg_name, []))
                        cell["drv_share_region"] = driver_country.get(name, 0) / rtot if rtot else None
                        cell["parent"] = reg_name
                    elif name in regions:
                        cell["drv_share_global"] = sum(driver_country.get(x, 0) for x in regions[name]) / dct
            elif p and c in geo_col:
                g = geo_col[c]
                reg_name = region_of.get(g)
                parent_key = "__root__" if len(p[1]) == 1 else p[1][0]
                regional_share = ((((state.get("shares") or {}).get(p[0]) or {}).get(reg_name) or {})
                                  .get(parent_key) or {}).get(p[1][-1])
                cell.update(k="share", geo=g, dim=p[0], path=list(p[1]), gen=gen["shares"].get((p, g)),
                            hdisp=disp(hd["shares"], p), odisp=disp(gen["shares"], p), bias=seg_bias.get(p),
                            region=reg_name, regional=regional_share)
                if p[0] == dim and len(p[1]) == 1:
                    s = p[1][0]
                    cell.update(regulated=True, drv=driver_mix.get(g, {}).get(s),
                                our_global=our_global_mix.get(s), human_global=human_global_mix.get(s))
                elif p[0] in cons_dims and len(p[1]) == 1:
                    cell["consistency"] = True
            elif c in growth_col and (p or dim_of_row.get(r) or r == 1):
                cell.update(k="seg_cagr" if p else "geo_total_cagr", geo=growth_col[c],
                            dim=p[0] if p else dim_of_row.get(r), path=list(p[1]) if p else [])
            elif c in trend_col and p:
                cell.update(k="trend_code", geo=trend_col[c], dim=p[0], path=list(p[1]))
            else:
                cell["k"] = "label"
                if c == 2 and p:
                    cell["lvl"] = len(p[1])  # 1 = top-level segment, 2 = sub-segment of the row's parent
            cells.append(cell)
    sheets["Data"] = {"rows": hrows, "cols": hcols, "cells": cells}

    # ------------------------------------------------------------- Market Value
    hg, hrows, hcols = grid(hwb, "Market Value")
    fg, _, _ = grid(fwb, "Market Value")
    hdr_r = next(i for i, row in enumerate(hg, start=1) if any(_norm(x).lower() in ("regions", "region") for x in row))
    head = {}
    for j, x in enumerate(hg[hdr_r - 1]):  # first occurrence wins ('Share' repeats for the Volume block)
        if x and _norm(x).lower() not in head:
            head[_norm(x).lower()] = j + 1
    cells = []
    for r in range(1, hrows + 1):
        for c in range(1, hcols + 1):
            h = hg[r - 1][c - 1] if c - 1 < len(hg[r - 1]) else None
            v = fg[r - 1][c - 1] if r - 1 < len(fg) and c - 1 < len(fg[r - 1]) else None
            if layout_only and isinstance(h, (int, float)) and not isinstance(h, bool):
                h = None
            if h is None and v is None:
                continue
            cell = {"r": r, "c": c, "h": h if not isinstance(h, float) else round(h, 6),
                    "v": v if not isinstance(v, float) else round(v, 6), "k": "label"}
            name = _norm(hg[r - 1][1]) if len(hg[r - 1]) > 1 else ""
            if r > hdr_r and name:
                kind = {head.get("usd mn"): "mv_value", head.get("share"): "mv_share", head.get("cagr"): "mv_cagr"}.get(c)
                if kind:
                    gm = gen_mv.get(name) or {}
                    cell.update(k=kind, geo=name, gen={"mv_value": gm.get("usd mn"), "mv_share": gm.get("share"),
                                                       "mv_cagr": gm.get("cagr")}[kind])
                    if name in regions:
                        cell["drv_share_global"] = sum(driver_country.get(x, 0) for x in regions[name]) / dct
            cells.append(cell)
    sheets["Market Value"] = {"rows": hrows, "cols": hcols, "cells": cells}

    # --------------------------------------------------------------------- ASP
    hg, hrows, hcols = grid(hwb, "ASP")
    fg, _, _ = grid(fwb, "ASP")
    headers = [i for i, row in enumerate(hg, start=1) if len(row) > 2 and _norm(row[2]) == geos[0]]
    acol = {j + 1: _norm(x) for j, x in enumerate(hg[headers[0] - 1]) if _norm(x) in geos} if headers else {}
    block_of, total_row = {}, {}
    for bi, hr in enumerate(headers):
        rr = hr + 1
        while rr <= hrows and (_norm(hg[rr - 1][1]) if len(hg[rr - 1]) > 1 else ""):
            block_of[rr] = bi
            rr += 1
        total_row[hr - 1] = bi
    kinds = ["asp", "asp_share", "asp_value", "asp_volume"]
    cells = []
    for r in range(1, hrows + 1):
        for c in range(1, hcols + 1):
            h = hg[r - 1][c - 1] if c - 1 < len(hg[r - 1]) else None
            v = fg[r - 1][c - 1] if r - 1 < len(fg) and c - 1 < len(fg[r - 1]) else None
            if layout_only and isinstance(h, (int, float)) and not isinstance(h, bool):
                h = None
            if h is None and v is None:
                continue
            cell = {"r": r, "c": c, "h": h if not isinstance(h, float) else round(h, 6),
                    "v": v if not isinstance(v, float) else round(v, 6), "k": "label"}
            lab = _norm(hg[r - 1][1]) if len(hg[r - 1]) > 1 else ""
            if r in block_of and c in acol and lab:
                cell.update(k=kinds[block_of[r]], geo=acol[c], segment=lab,
                            gen=gen_asp.get((lab, acol[c])) if block_of[r] == 0 else None)
            elif r in total_row and c in acol:
                cell.update(k="asp_total", geo=acol[c], block=total_row[r])
            elif isinstance(h, str) and (h.startswith("ASP") and h[3:].isdigit()):
                cell["k"] = "trend_code"
            cells.append(cell)
    sheets["ASP"] = {"rows": hrows, "cols": hcols, "cells": cells}

    # ------------------------------------------------------------------ context
    inp = evaluation.get("input_sheet", {})
    bu = ((state.get("scope_reconciliation") or {}).get("cache") or {}).get("bottom_up") or {}
    context = {
        "market": market, "tag": tag,
        "global": inp.get("global_value_2025"), "global_cagr": inp.get("global_cagr"),
        "grades": (evaluation.get("verdict") or {}).get("grades"),
        "metrics": {k: inp.get(k) for k in ("regional_share", "country_share_in_region", "country_value",
                                            "segment_shares", "asp", "regional_cagr", "country_cagr")},
        "by_dimension": {d: x.get("shares") for d, x in (inp.get("by_dimension") or {}).items()},
        "partition_dim": dim, "segments": reg.get("segments"), "guard": guard,
        "consistency_notes": (reg.get("consistency") or {}).get("notes"), "consistency_llm": llm_cons,
        "events": ev_scan.get("events"), "events_applied": ev_applied, "event_changes": ev_changes,
        "events_for": {g: [e["event"] for e in events_for(g)] for g in list(regions) + geos},
        "bottom_up": [{k: c.get(k) for k in ("segment", "value_mn", "basis", "n_publishers", "spread_mn")}
                      for c in bu.get("components") or []],
        "regions": regions, "geographies": geos,
        "shares_trace": state.get("shares_trace") or {}, "geo_trace": state.get("geo_trace") or {},
        "market_trace": state.get("market_trace") or {},
        "usecases": None if not usecases else {
            "dimension": usecases.get("dimension"), "industries": usecases.get("industries"),
            "industry_value": usecases.get("industry_value"), "country_value": usecases.get("country_value"),
            "region_value": usecases.get("region_value"), "global_value": usecases.get("global_value"),
            "flags": usecases.get("flags"), "consistency": usecases.get("consistency"),
            "settings": usecases.get("settings"),
            "mix_profiles": (usecases.get("mix") or {}).get("profiles"),
            "mix_dims": (usecases.get("mix") or {}).get("dimensions") or [],
            "items": [{k: r.get(k) for k in ("industry", "use_case", "own", "spec", "driver_values", "units", "values_mn",
                                              "world_mn", "evidence", "published_context", "self_checks", "reexamined",
                                              "world_units_estimate", "world_units_formula", "world_units_circular",
                                              "price_per_quantity_unit", "driver_meta", "anchor",
                                              # the machine-model build
                                              "archetype", "world_inputs", "world_formula", "world_units",
                                              "driver_total", "driver_share", "bound_violations", "basis_counts",
                                              "blocked", "flags", "spread", "median_of_samples", "excluded_samples",
                                              "driver_votes", "samples", "published_gap", "attempts")}
                      for r in usecases.get("use_cases") or []]},
        "shares_final": {d: (v or {}).get("__global__") for d, v in (state.get("shares") or {}).items()},
        "asp_trace": state.get("asp_trace") or {},
        "human_global_mix": human_global_mix, "our_global_mix": our_global_mix,
        "narrative": (evaluation.get("narrative") or ""),
    }
    md = (hel / f"evaluation_{tag}.md").read_text(encoding="utf-8") if (hel / f"evaluation_{tag}.md").exists() else ""
    m = re.search(r"## Evaluator agent diagnosis\n\n(.*?)(\n## |\Z)", md, re.S)
    context["narrative"] = m.group(1).strip() if m else ""

    payload = json.dumps({"sheets": sheets, "context": context, "order": ["Data", "ASP", "Market Value"]},
                         default=str, ensure_ascii=False)
    html = TEMPLATE.read_text(encoding="utf-8").replace("/*__DATA__*/null", payload)
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / "index.html"
    out.write_text(html, encoding="utf-8")
    return out


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--market", default="Global High Energy Laser Market")
    p.add_argument("--human-data", default="src/newinfo/Data file-Global High Energy Laser Market.xlsx")
    p.add_argument("--tag", default="v10")
    p.add_argument("--gen-tag", help="tag of the generation run when --tag reused its Input Sheet")
    p.add_argument("--dir", default="output/hel", help="run output folder (e.g. output/gearbox)")
    p.add_argument("--layout-only", action="store_true",
                   help="--human-data is a blank template from an ME's labels (no human numbers to show)")
    p.add_argument("--serve", action="store_true")
    p.add_argument("--port", type=int, default=8765)
    args = p.parse_args()
    slug = re.sub(r"[^a-z0-9]+", "-", args.market.lower().replace("global ", "").replace(" market", "")).strip("-")
    out = build(args.market, ROOT / args.human_data, ROOT / "demo" / f"{slug}-{args.tag}", args.tag, args.gen_tag,
                run_dir=ROOT / args.dir, layout_only=args.layout_only)
    print(f"[demo] {out} ({out.stat().st_size / 1e6:.1f} MB)")
    if args.serve:
        import functools
        import http.server
        handler = functools.partial(http.server.SimpleHTTPRequestHandler, directory=str(out.parent))
        print(f"[demo] serving on http://localhost:{args.port}")
        http.server.ThreadingHTTPServer(("127.0.0.1", args.port), handler).serve_forever()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
