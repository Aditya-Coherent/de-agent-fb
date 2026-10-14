"""Structure agent, data writer, evaluator agent and scope reconciliation.

Uses the human High Energy Laser pair in src/newinfo as the fixture: the
human market evaluated against itself must score perfectly, and a market
scaled by a known factor must have exactly that factor attributed to
'level'."""
from __future__ import annotations

import math
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from me_engine.generation import data_writer, evaluator_agent as ev, scope, structure_agent as sa

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "src" / "newinfo" / "Data file-Global High Energy Laser Market.xlsx"
ME = ROOT / "src" / "newinfo" / "ME - Global High Energy Laser Market.xlsx"
MARKET = "Global High Energy Laser Market"
pytestmark = pytest.mark.skipif(not (DATA.exists() and ME.exists()), reason="High Energy Laser fixture not present")


@pytest.fixture(scope="module")
def spec():
    return sa.extract(DATA, ME, MARKET)


def test_volume_scale_from_band_title():
    assert sa.volume_scale("Market Volume (Units)") == 1e6
    assert sa.volume_scale("Market Volume (Th Liters)") == 1e3
    assert sa.volume_scale("Market Volume (Mn Units)") == 1.0


def test_structure_spec_is_complete_and_consistent(spec):
    assert len(spec["geographies"]) == 26 and len(spec["regions"]) == 6 and len(spec["dimensions"]) == 8
    assert spec["me_layout"]["volume_unit"] == "Units" and spec["me_layout"]["years"] == list(range(2021, 2034))
    assert set(spec["duplicate_labels"]) == {"Industrial Manufacturing", "Medical"}
    assert all(c["ok"] for c in spec["checks"])
    paths = [tuple(s["path"]) for d in spec["dimensions"].values() for s in d]
    assert len(paths) == len(set(paths))  # duplicate labels stay distinct by path


def test_blank_templates_keep_layout_and_drop_every_number(spec, tmp_path):
    sa.blank_template(ME, tmp_path / "me.xlsx")
    data_writer.blank_data_template(DATA, tmp_path / "data.xlsx")
    ws = load_workbook(tmp_path / "me.xlsx", read_only=True)["Global"]
    rows = list(ws.iter_rows(values_only=True))
    labelled = [r for r in rows if len(r) > 2 and r[2]]
    assert labelled and not any(isinstance(v, (int, float)) for r in labelled for v in r[3:])
    assert rows[3][3] == 2021  # year header kept
    d = load_workbook(tmp_path / "data.xlsx", read_only=True)["Data"]
    assert not any(isinstance(v, (int, float)) for r in d.iter_rows(values_only=True) for v in r)
    assert sa.conformance(spec, tmp_path / "data.xlsx", tmp_path / "me.xlsx")["score"] == 1.0  # layout intact


def test_human_against_itself_is_perfect(spec):
    rep = ev.evaluate(spec, DATA, ME, DATA, ME)
    assert rep["structure"]["score"] == 1.0
    assert math.isclose(rep["me"]["global_2025"]["ratio"], 1.0)
    for comp in rep["me"]["decomposition"].values():
        assert comp["typical_x"] == 1.0
    assert rep["input_sheet"]["segment_shares"]["mae_pts"] == 0.0
    assert rep["input_sheet"]["segment_shares"]["n"] == 97 * 26
    assert rep["me"]["coverage"]["filled"] == 1.0


def test_known_scaling_is_attributed_to_level(spec, tmp_path):
    scaled = tmp_path / "me_x2.xlsx"
    wb = load_workbook(ME)
    for ws in wb.worksheets:
        band = None
        for row in ws.iter_rows():
            marker = row[3].value if len(row) > 3 else None
            if isinstance(marker, str) and marker.strip():
                band = marker.strip()
                continue
            if band and band.startswith("Market Value") and len(row) > 2 and row[2].value:
                for c in row[3:16]:
                    if isinstance(c.value, (int, float)):
                        c.value *= 2.0
    wb.save(scaled)
    rep = ev.evaluate(spec, DATA, ME, DATA, scaled)
    d = rep["me"]["decomposition"]
    assert math.isclose(rep["me"]["global_2025"]["ratio"], 2.0, rel_tol=1e-9)
    assert d["level"]["typical_x"] == 2.0 and d["geo_mix"]["typical_x"] == 1.0
    assert d["seg_mix"]["typical_x"] == 1.0 and d["growth"]["typical_x"] == 1.0


def test_data_writer_round_trips_the_human_file(spec, tmp_path):
    data_writer.blank_data_template(DATA, tmp_path / "blank.xlsx")
    out = tmp_path / "rewritten.xlsx"
    stats = data_writer.fill(spec, tmp_path / "blank.xlsx", DATA, ME, out)
    assert stats["shares"] == 97 * 26 and stats["growth"] == 97 * 26
    rep = ev.evaluate(spec, DATA, ME, out, ME)
    assert rep["input_sheet"]["segment_shares"]["mae_pts"] == 0.0
    assert rep["input_sheet"]["segment_growth"]["mae_pts"] < 0.5   # CAGR block recomputed from the ME


def test_segment_sizing_uses_publisher_consensus_not_one_firm():
    snippets = {
        "Industrial": ["Firm A: the industrial laser market was USD 9.0 billion in 2025.",
                       "Firm B estimates industrial lasers at USD 7.0 billion in 2025.",
                       "Firm A also said USD 12.0 billion in another release in 2025."],
        "Medical": ["Firm C: the medical laser market was USD 5.0 billion in 2025."],
        "Sensing": ["nothing here"],
    }

    class LLM:
        def complete_json(self, system, user, **_):
            if "Which of these segments" in user:
                return {"covered": ["Defense"], "scope_of_cited_figure": "directed energy", "reason": "x"}
            if "Segment: Industrial" in user:
                return {"figures": [
                    {"publisher": "Firm A Research", "value_usd_mn_2025": 9000, "quote": snippets["Industrial"][0],
                     "snippet_index": 0, "adjustment": 1.0},
                    {"publisher": "Firm B Insights", "value_usd_mn_2025": 7000, "quote": snippets["Industrial"][1],
                     "snippet_index": 1, "adjustment": 1.0},
                    {"publisher": "Firm A", "value_usd_mn_2025": 12000, "quote": snippets["Industrial"][2],
                     "snippet_index": 2, "adjustment": 1.0}],  # same publisher twice -> counted once
                    "estimate_usd_mn_2025": 20000}
            if "Segment: Medical" in user:
                return {"figures": [{"publisher": "Firm C", "value_usd_mn_2025": 5000, "quote": snippets["Medical"][0],
                                     "snippet_index": 0, "adjustment": 0.8}], "estimate_usd_mn_2025": 2500}
            return {"figures": [{"publisher": "Made Up", "value_usd_mn_2025": 4000, "quote": "invented sentence",
                                 "snippet_index": 0}], "estimate_usd_mn_2025": 1000}

    nesting = {"By Application": {"__root__": ["Defense", "Industrial", "Medical", "Sensing"]}}

    def search(qs):
        return snippets[qs[0].split()[0]]

    rec = scope.reconcile(LLM(), "Global High Energy Laser Market", nesting, 3930.0, "USD 3.93 Billion in 2025", search)
    ind, med, sen = rec["components"]
    assert ind["basis"] == "consensus" and ind["grounded"] and ind["n_publishers"] == 2
    assert ind["value_mn"] == 8000                       # median of 9,000 and 7,000 -- not the model's 20,000
    assert med["basis"] == "single_source" and not med["grounded"]
    assert math.isclose(med["value_mn"], (5000 * 0.8 * 2500) ** 0.5)
    assert sen["basis"] == "estimate" and sen["value_mn"] == 1000 and sen["rejected"] == 1   # invented quote dropped
    assert math.isclose(rec["anchor_mn"], 3930 + 8000 + med["value_mn"] + 1000)
    assert "median of 2 publishers" in scope.brief(rec)
    assert scope.partition_dimension({"By Laser Type": {}, "By End User": {}}) == "By End User"


# --------------------------------------------------------------------------- #
# assumptor regulation (published figures are artifacts; real-world data decides)
# --------------------------------------------------------------------------- #
def test_geo_buckets_cover_named_groups_and_rests():
    from me_engine.generation import assumptor_regulator as ar
    geos = ["U.S.", "ASEAN", "GCC", "Rest of Europe", "Central Africa", "Rest of Asia Pacific", "North Africa"]
    assert ar.bucket_of("USA", {"region": "North America"}, geos) == "U.S."
    assert ar.bucket_of("VNM", {"region": "East Asia & Pacific"}, geos) == "ASEAN"
    assert ar.bucket_of("SAU", {"region": "Middle East & North Africa"}, geos) == "GCC"
    assert ar.bucket_of("POL", {"region": "Europe & Central Asia"}, geos) == "Rest of Europe"
    assert ar.bucket_of("KAZ", {"region": "Europe & Central Asia"}, geos) == "Rest of Asia Pacific"
    assert ar.bucket_of("NGA", {"region": "Sub-Saharan Africa"}, geos) == "Central Africa"
    assert ar.bucket_of("EGY", {"region": "Middle East & North Africa"}, geos) == "North Africa"
    drivers = {"USA": {"region": "North America", "gdp_usd": 100.0, "military_expenditure_usd": 10.0},
               "POL": {"region": "Europe & Central Asia", "gdp_usd": 50.0},           # missing -> region ratio
               "DEU": {"region": "Europe & Central Asia", "gdp_usd": 40.0, "military_expenditure_usd": 2.0}}
    by_geo = ar.driver_by_geo(drivers, ["U.S.", "Rest of Europe", "Germany"], "military_expenditure_usd")
    assert by_geo["U.S."] == 10.0 and math.isclose(by_geo["Rest of Europe"], 50.0 * 0.05)


def test_assumptor_formula_is_regulated_by_honest_artifacts():
    from me_engine.generation import assumptor_regulator as ar

    class LLM:
        def __init__(self, answers):
            self.answers = list(answers)

        def complete_json(self, system, user, **_):
            return self.answers.pop(0)

    comp = {"figures": [{"publisher": "A", "adjusted_mn": 5000, "independent": True},
                        {"publisher": "B", "adjusted_mn": 7000, "independent": True},
                        {"publisher": "unattributed", "adjusted_mn": 30000, "independent": False}]}
    # inside the independent range -> the assumptor's own number stands
    ok = ar.assume_segment(LLM([{"formula": "driver_total_usd_mn * k", "constants": {"k": {"value": 0.0006}}}]),
                           "M", "Medical", "", "health_expenditure_usd", 10_000e9, comp)
    assert math.isclose(ok["value_mn"], 6000) and ok["n_independent"] == 2 and not ok["flags"]
    # wildly implausible (unit error) -> repaired by a second formula
    rep = ar.assume_segment(LLM([{"formula": "driver_total_usd_mn * k", "constants": {"k": {"value": 1e-9}}},
                                 {"formula": "driver_total_usd_mn * k", "constants": {"k": {"value": 0.0005}}}]),
                            "M", "Medical", "", "health_expenditure_usd", 10_000e9, comp)
    assert rep["how"] == "formula (repaired)" and math.isclose(rep["value_mn"], 5000)
    # plausible but outside the 2+ independent publishers' range -> clipped and flagged
    hi = ar.assume_segment(LLM([{"formula": "driver_total_usd_mn * k", "constants": {"k": {"value": 0.0009}}}]),
                           "M", "Medical", "", "health_expenditure_usd", 10_000e9, comp)
    assert hi["value_mn"] == 7000 and "clipped" in hi["flags"][-1]


# --------------------------------------------------------------------------- #
# events agent: real-world events move the forecast, verified and bounded
# --------------------------------------------------------------------------- #
def test_events_are_verified_dated_bounded_and_roll_up(tmp_path):
    from datetime import datetime, timezone
    from me_engine.generation import events_agent as ea

    today = datetime(2026, 10, 13, tzinfo=timezone.utc)
    articles = [
        {"date": "2026-10-01", "source": "A", "url": "u0", "title": "War closes strait",
         "body": "Shipping through the strait has stopped after the attacks."},
        {"date": "2023-01-01", "source": "B", "url": "u1", "title": "Old news", "body": "Something happened long ago."},
        {"date": "2026-09-01", "source": "C", "url": "u2", "title": "Big US contract", "body": "The army ordered lasers."},
    ]

    class LLM:
        def complete_json(self, system, user, **_):
            return {"events": [
                {"event": "Strait closure", "article_index": 0, "quote": "Shipping through the strait has stopped",
                 "direction": "up", "impact_pp_per_year": 40, "duration_years": 8, "affects": ["Global"], "confidence": 1},
                {"event": "Stale", "article_index": 1, "quote": "Something happened long ago",
                 "impact_pp_per_year": 2, "duration_years": 2, "affects": ["Global"], "confidence": 1},
                {"event": "Invented", "article_index": 2, "quote": "this sentence is not in the article",
                 "impact_pp_per_year": 2, "duration_years": 2, "affects": ["U.S."], "confidence": 1},
                {"event": "US order", "article_index": 2, "quote": "The army ordered lasers",
                 "direction": "up", "impact_pp_per_year": 1.6, "duration_years": 4, "affects": ["U.S."], "confidence": 0.5}]}

    regions = {"North America": ["U.S.", "Canada"]}
    evs = ea.extract_events(LLM(), "M", articles, list(regions), ["U.S.", "Canada"], today, 540)
    assert [e["event"] for e in evs] == ["Strait closure", "US order"]          # stale + invented dropped
    assert evs[0]["effective_cagr_pp"] == ea.MAX_EVENT_PP                       # +40 pp/yr capped at 3 pp
    assert math.isclose(evs[1]["effective_cagr_pp"], 1.6 * 4 / 8 * 0.5)         # 0.4 pp
    deltas = ea.geo_deltas(evs, regions)
    assert math.isclose(deltas["U.S."], 3.4) and math.isclose(deltas["Canada"], 3.0)

    # apply to a copy of the human Data file: region = value-weighted mean of its countries
    import shutil
    f = tmp_path / "data.xlsx"
    shutil.copyfile(DATA, f)
    res = ea.apply(f, {"deltas_pp": deltas}, regions)
    us = next(c for c in res["changes"] if c["geo"] == "U.S.")
    na = next(c for c in res["changes"] if c["geo"] == "North America")
    assert math.isclose(us["delta_pp"], 3.4) and 3.0 < na["delta_pp"] < 3.4
    assert res["global_cagr_adjusted"] > res["global_cagr_baseline"]           # all-positive events never lower it
