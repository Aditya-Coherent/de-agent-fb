"""Write a generated market into the human Data file's exact structure.

The generator's own sheet writer has its own layout (segments grouped by
level, a different ASP table, no per-segment growth block). To compare -- and
to hand over a file that looks like the analyst's -- the generated numbers are
placed into a blank copy of the human Data file (structure_agent output),
keyed by full segment path so duplicate labels never collide:

    Data         segment shares (per country), the CAGR 2025-33 block (from the
                 generated ME), region / country shares
    Market Value regional value (US$ Mn), share, CAGR; Global row
    ASP          block 1 ASP per product per country, block 2 shares, block 3
                 value (US$ Mn), block 4 volume (units), and their total rows

Cells the generator has no number for stay EMPTY (e.g. the analyst's
CMI_Trend curve-template IDs) -- never copied from the human file.
"""
from __future__ import annotations

import math
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from . import evaluator_agent as ev
from .structure_agent import _norm

BASE, END = 2025, 2033


def blank_data_template(data_file: Path | str, out: Path | str) -> dict[str, int]:
    """Copy the human Data file and clear every number and every per-cell
    analyst code (CMI_Trend_.., ASP7 ..) -- labels and headers stay."""
    wb = load_workbook(data_file)
    cleared = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            for c in row:
                v = c.value
                is_number = isinstance(v, (int, float)) and not isinstance(v, bool)
                is_code = isinstance(v, str) and (v.startswith("CMI_Trend") or (v.startswith("ASP") and v[3:].isdigit()))
                if is_number or is_code:
                    c.value = None
                    cleared += 1
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return {"cells_cleared": cleared}


def _gen_asp(path: Path, geos: list[str]) -> dict[tuple[str, str], float]:
    """ASP per (product, country) from the generator's own ASP sheet (any layout:
    the header row is the one holding the first country name)."""
    ws = load_workbook(path, data_only=True, read_only=True)["ASP"]
    rows = list(ws.iter_rows(values_only=True))
    out = {}
    for i, r in enumerate(rows):
        cols = {_norm(v): j for j, v in enumerate(r) if _norm(v) in geos}
        if geos[0] not in cols:
            continue
        lab_col = min(cols.values()) - 1
        for r2 in rows[i + 1:]:
            lab = _norm(r2[lab_col]) if lab_col < len(r2) else ""
            if not lab:
                if out:
                    break
                continue
            for g, j in cols.items():
                v = r2[j] if j < len(r2) else None
                if isinstance(v, (int, float)) and v > 0:
                    out[(lab, g)] = float(v)
        break
    return out


def _market_value(path: Path) -> dict[str, dict[str, float | None]]:
    rows = list(load_workbook(path, data_only=True, read_only=True)["Market Value"].iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if any(_norm(v).lower() in ("regions", "region") for v in r))
    h = [_norm(v).lower() for v in rows[hdr_i]]
    ci = {k: next((i for i, v in enumerate(h) if v == k), None) for k in ("usd mn", "share", "cagr")}
    ci_name = next(i for i, v in enumerate(h) if v in ("regions", "region"))
    out = {}
    for r in rows[hdr_i + 1:]:
        name = _norm(r[ci_name]) if ci_name < len(r) else ""
        if name:
            out[name] = {k: (float(r[i]) if i is not None and i < len(r) and isinstance(r[i], (int, float)) else None)
                         for k, i in ci.items()}
    return out


def fill(spec: dict[str, Any], template: Path | str, gen_data: Path | str, gen_me: Path | str | None,
         out: Path | str) -> dict[str, Any]:
    gen_data = Path(gen_data)
    gd = ev.read_data_file(gen_data, spec)
    gm = ev.read_me(Path(gen_me), spec) if gen_me else {}
    geos = spec["geographies"]
    wb = load_workbook(template)
    stats = {"shares": 0, "growth": 0, "region_shares": 0, "market_value": 0, "asp_cells": 0}

    # ---- Data sheet ----
    ws = wb["Data"]
    labels = [_norm(ws.cell(r, 2).value) for r in range(1, ws.max_row + 1)]
    paths = ev._path_walk(labels, spec)
    geo_col = {_norm(ws.cell(3, c).value): c for c in range(3, 30) if _norm(ws.cell(3, c).value)}
    growth_col = {}
    for c in range(31, ws.max_column + 1):
        g = _norm(ws.cell(3, c).value)
        if g in geos and g not in growth_col:
            growth_col[g] = c
    dim_of_row = {}
    current = None
    for r, lab in enumerate(labels, start=1):
        if lab in spec["dimensions"]:
            current = lab
            dim_of_row[r] = lab
    for r, p in enumerate(paths, start=1):
        if not p:
            continue
        for g in geos:
            v = gd["shares"].get((p, g))
            if v is not None and g in geo_col:
                ws.cell(r, geo_col[g]).value = v
                stats["shares"] += 1
            vals = ((gm.get(g) or {}).get("value") or {}).get(p) or {}
            if g in growth_col and vals.get(BASE) and vals.get(END):
                ws.cell(r, growth_col[g]).value = (vals[END] / vals[BASE]) ** (1 / (END - BASE)) - 1
                stats["growth"] += 1
    for r, dim in dim_of_row.items():  # dimension title row + row 1: the country's total CAGR
        for g in geos:
            tot = ((gm.get(g) or {}).get("value") or {}).get(("__total__", dim)) or {}
            if g in growth_col and tot.get(BASE) and tot.get(END):
                cagr = (tot[END] / tot[BASE]) ** (1 / (END - BASE)) - 1
                ws.cell(r, growth_col[g]).value = cagr
                if r == min(dim_of_row):
                    ws.cell(1, growth_col[g]).value = cagr
    start = next((i for i, lab in enumerate(labels, start=1) if lab == "Region"), None)
    if start:
        for r in range(start + 1, ws.max_row + 1):
            lab = _norm(ws.cell(r, 2).value)
            if lab == "Global":
                ws.cell(r, 3).value = 1
            elif lab in gd["region_share"]:
                ws.cell(r, 3).value = gd["region_share"][lab]
                stats["region_shares"] += 1
            # D: CAGR, E: 2025 value (US$ Mn) -- what Stage B reads as each geography's anchor
            if lab in gd["region_cagr"] and lab != "Global":
                ws.cell(r, 4).value = gd["region_cagr"][lab]
            if lab in gd["region_value"]:
                ws.cell(r, 5).value = gd["region_value"][lab]

    # ---- Market Value sheet ----
    mv = _market_value(gen_data)
    wsm = wb["Market Value"]
    hdr = next(r for r in range(1, 10) if any(_norm(wsm.cell(r, c).value).lower() in ("regions", "region")
                                                for c in range(1, wsm.max_column + 1)))
    h = {_norm(wsm.cell(hdr, c).value).lower(): c for c in range(1, wsm.max_column + 1) if wsm.cell(hdr, c).value}
    name_c = h.get("regions") or h.get("region")
    share_c = next(c for c in range(1, wsm.max_column + 1) if _norm(wsm.cell(hdr, c).value).lower() == "share")
    for r in range(hdr + 1, wsm.max_row + 1):
        name = _norm(wsm.cell(r, name_c).value)
        g = mv.get(name)
        if not g:
            continue
        if h.get("usd mn") and g.get("usd mn") is not None:
            wsm.cell(r, h["usd mn"]).value = g["usd mn"]
        wsm.cell(r, share_c).value = 1.0 if name == "Global" else g.get("share")
        if h.get("cagr") and g.get("cagr") is not None:
            wsm.cell(r, h["cagr"]).value = g["cagr"]
        stats["market_value"] += 1

    # ---- ASP sheet: four blocks under header rows holding the first country ----
    wsa = wb["ASP"]
    asp = _gen_asp(gen_data, geos)
    priced = spec["priced_dimension"]
    headers = [r for r in range(1, wsa.max_row + 1) if _norm(wsa.cell(r, 3).value) == geos[0]]
    acol = {_norm(wsa.cell(headers[0], c).value): c for c in range(3, 30)} if headers else {}

    def block_rows(hr: int) -> list[tuple[int, Any]]:
        out, labs = [], []
        r = hr + 1
        while r <= wsa.max_row and (_norm(wsa.cell(r, 2).value) or not labs):
            labs.append((r, _norm(wsa.cell(r, 2).value)))
            r += 1
        pw = ev._path_walk([priced] + [lab for _, lab in labs if lab != priced], spec)[1:]
        it = iter(pw)
        for rr, lab in labs:
            out.append((rr, None if lab == priced or not lab else next(it, None)))
        return out

    for bi, hr in enumerate(headers[:4]):
        for rr, p in block_rows(hr):
            if not p:
                continue
            for g in geos:
                c = acol.get(g)
                if not c:
                    continue
                v = None
                gv = (gm.get(g) or {})
                if bi == 0:
                    leaf = p[1][-1]
                    v = asp.get((leaf, g)) if not any(s["label"] == leaf and s["has_children"]
                                                       for s in spec["dimensions"][priced]) else 0
                elif bi == 1:
                    v = gd["shares"].get((p, g))
                elif bi == 2:
                    v = ((gv.get("value") or {}).get(p) or {}).get(BASE)
                elif bi == 3:
                    v = ((gv.get("volume") or {}).get(p[1][-1]) or {}).get(BASE) \
                        or ((gv.get("volume") or {}).get(p) or {}).get(BASE)
                    v = round(v) if v else v
                if v is not None:
                    wsa.cell(rr, c).value = v
                    stats["asp_cells"] += 1
        # total row above the header: block 0 / 3 -> units per country, block 2 -> country value
        if bi in (0, 2, 3) and hr > 1:
            for g in geos:
                c = acol.get(g)
                gv = gm.get(g) or {}
                if bi == 2:
                    v = ((gv.get("value") or {}).get(("__total__", priced)) or {}).get(BASE)
                else:
                    v = ((gv.get("volume") or {}).get(("__total__", priced)) or {}).get(BASE)
                    v = round(v) if v else v
                if c and v is not None:
                    wsa.cell(hr - 1, c).value = v
            glob = gm.get("Global") or {}
            tot = ((glob.get("volume" if bi != 2 else "value") or {}).get(("__total__", priced)) or {}).get(BASE)
            if tot:
                wsa.cell(hr - 1, 2).value = round(tot) if bi != 2 else tot
                if bi == 0:
                    wsa.cell(1, 2).value = round(tot / 1000)
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return stats
