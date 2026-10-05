"""Load a company/distributor list from whatever format the source agent hands us.

Supported inputs:
- .csv / .xlsx — the company-finder agent's report format: a header row
  (#, Company, Brand, Parent_or_Independent, Website, Functionality,
  Geography, Is_Relevant, Summary) with section-break rows in between
  ("=== Manufacturers of X (38) ===" in the .csv, "Manufacturers of X  (38)"
  as a lone-cell row in the .xlsx). Only the "Manufacturers" section is kept
  by default — raw-material suppliers sell into a different upstream
  market, and distributor revenue includes resale markup and often spans
  multiple manufacturers/categories, so neither maps cleanly onto this
  market's size (see the demand-vs-supply discussion this exists for).
- .json — a flat list of names, or a list of objects with at least a
  "name" key and any number of extra fact keys (country, known_revenue,
  ...) threaded through to the estimator as context.
- .txt  — one company name per line.
"""
from __future__ import annotations

import csv
import json
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

_SECTION_RE = re.compile(r"^=*\s*(.+?)\s*\(\d+\)\s*=*$")


@dataclass(frozen=True, slots=True)
class Company:
    name: str
    facts: dict = field(default_factory=dict)


def _section_name(cell: str) -> str | None:
    """If `cell` is a section-break row ('=== Manufacturers of X (38) ==='
    or the bare '(38)'-suffixed variant), return the section label."""
    m = _SECTION_RE.match(cell.strip())
    return m.group(1).strip() if m else None


def _rows_to_companies(rows: list[list[str]], sections: tuple[str, ...]) -> list[Company]:
    header_idx = next((i for i, r in enumerate(rows) if "Company" in r), None)
    if header_idx is None:
        # No recognizable header — treat every non-empty first cell as a name.
        return [Company(name=str(r[0]).strip()) for r in rows if r and str(r[0]).strip()]

    header = [str(h).strip() for h in rows[header_idx]]
    col = {h: i for i, h in enumerate(header) if h}
    name_col = col.get("Company")
    if name_col is None:
        return []

    skip_cols = {"#", "Company", "Is_Relevant", "Summary", "Website"}
    fact_cols = [c for c in header if c and c not in skip_cols and c != "Company"]
    wanted = tuple(s.lower() for s in sections)

    companies: list[Company] = []
    current_section: str | None = None
    for row in rows[header_idx + 1:]:
        cells = [str(c).strip() if c is not None else "" for c in row]
        non_empty = [c for c in cells if c]
        if len(non_empty) == 1:
            section = _section_name(non_empty[0]) or non_empty[0]
            current_section = section
            continue
        if len(cells) <= name_col or not cells[name_col]:
            continue
        if wanted and not any(w in (current_section or "").lower() for w in wanted):
            continue
        is_relevant_col = col.get("Is_Relevant")
        if is_relevant_col is not None and len(cells) > is_relevant_col:
            if cells[is_relevant_col] and cells[is_relevant_col].lower() not in ("yes", ""):
                continue
        facts = {c: cells[col[c]] for c in fact_cols if len(cells) > col[c] and cells[col[c]]}
        if current_section:
            facts["category"] = current_section
        companies.append(Company(name=cells[name_col], facts=facts))
    return companies


def _load_csv(path: Path, sections: tuple[str, ...]) -> list[Company]:
    with path.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    return _rows_to_companies(rows, sections)


def _load_xlsx(path: Path, sections: tuple[str, ...]) -> list[Company]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    return _rows_to_companies(rows, sections)


def load_companies(path: Path | str, sections: tuple[str, ...] = ("Manufacturer",)) -> list[Company]:
    """`sections` filters to section-break labels containing any of these
    (case-insensitive) substrings; pass () to keep everything. Only applies
    to the .csv/.xlsx company-finder report format — .json/.txt have no
    section concept."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Company list not found: {path}")

    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _load_csv(path, sections)
    if suffix == ".xlsx":
        return _load_xlsx(path, sections)

    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        companies: list[Company] = []
        for item in data:
            if isinstance(item, str):
                companies.append(Company(name=item.strip()))
            elif isinstance(item, dict):
                name = str(item.get("name") or item.get("company") or "").strip()
                if not name:
                    continue
                facts = {k: v for k, v in item.items() if k not in ("name", "company") and v}
                companies.append(Company(name=name, facts=facts))
        return companies

    # Plain text: one company per line.
    lines = path.read_text(encoding="utf-8").splitlines()
    return [Company(name=line.strip()) for line in lines if line.strip()]


def save_companies_csv(companies: list[Company], path: Path | str) -> Path:
    """Persist a discovered/loaded list to the same `my_companies/` CSV shape
    (no section markers — a flat "Manufacturers" block) so it can be re-loaded
    by `load_companies` and sits alongside company-finder-sourced files."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fact_keys: list[str] = []
    for c in companies:
        for k in c.facts:
            if k not in fact_keys:
                fact_keys.append(k)

    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["#", "Company", *fact_keys])
        writer.writerow([f"=== Manufacturers ({len(companies)}) ==="])
        for i, c in enumerate(companies, start=1):
            writer.writerow([i, c.name, *[c.facts.get(k, "") for k in fact_keys]])
    return path


# --------------------------------------------------------------------------- #
# Screening before estimation: role filter + parent/brand dedupe
# --------------------------------------------------------------------------- #
# Both fix real double counting seen in recorded runs: distributors (Cardinal
# Health, McKesson, Henry Schein) re-sell the same gloves the manufacturers
# already reported (Disposable Medical Gloves came out ~4x the published
# size), and one corporate group listed several times under parent/brand/
# regional names (Deoleo + Bertolli/Carbonell/Carapelli, Sovena + Sovena
# Group + Andorinha) hid ~$4.4B of duplicates in Olive Oil.

# Trailing tokens that don't distinguish one business from another — legal
# forms, group/holding wording, and regional-arm qualifiers.
_CORP_SUFFIXES = frozenset({
    "inc", "incorporated", "corp", "corporation", "co", "company", "companies", "ltd", "limited",
    "llc", "lp", "llp", "plc", "sa", "sau", "sas", "sl", "slu", "spa", "srl", "ag", "se", "nv", "bv",
    "gmbh", "kg", "kgaa", "ab", "asa", "as", "oyj", "bhd", "berhad", "sdn", "pte", "pty", "pvt",
    "private", "public", "tbk", "pt", "kk", "cv", "de", "and", "the", "holding", "holdings", "group",
    "international", "intl", "industries", "industry", "global", "worldwide", "enterprises",
    "asia", "pacific", "europe", "portugal", "spain", "italy", "usa", "us", "america", "americas",
    "uk", "india", "china", "japan", "thailand", "malaysia",
})

_NO_PARENT_VALUES = frozenset({"", "independent", "none", "n/a", "na", "-", "self", "unknown", "private"})

EXCLUDED_ROLES = frozenset({"distributor", "retailer"})
_ROLE_KEYS = ("role", "Role", "category", "Category", "type", "Type", "company_type", "Company_Type")
_PARENT_KEYS = ("parent_company", "Parent_Company", "parent", "Parent", "Parent_or_Independent")


def name_tokens(name: str) -> list[str]:
    """Lowercased ASCII word tokens of a company name with parenthesised
    asides removed and trailing corporate/regional suffixes stripped
    ("Borges International Group" -> ["borges"], "Deoleo, S.A." -> ["deoleo"]).
    Never strips a name down to nothing."""
    text = unicodedata.normalize("NFKD", str(name)).encode("ascii", "ignore").decode().lower()
    text = re.sub(r"\([^)]*\)", " ", text).replace("&", " and ")
    text = re.sub(r"[.']", "", text)  # "S.p.A." -> "spa", "McKesson's" -> "mckessons"
    tokens = [t for t in re.split(r"[^a-z0-9]+", text) if t]
    stripped = list(tokens)
    while stripped and stripped[-1] in _CORP_SUFFIXES:
        stripped.pop()
    return stripped or tokens


def normalize_company_name(name: str) -> str:
    return " ".join(name_tokens(name))


def name_similarity(a: str, b: str) -> float:
    """Token Dice similarity (0-1) between two company names, ignoring
    corporate suffixes. Deliberately token-level, not character-level:
    "Borges" vs "Borgestad ASA" is 0.8 by difflib's character ratio but a
    different business entirely, while "TopGlove" vs "Top Glove Corporation
    Bhd" is the same one (handled by the joined-string check)."""
    ta, tb = name_tokens(a), name_tokens(b)
    if not ta or not tb:
        return 0.0
    if "".join(ta) == "".join(tb):
        return 1.0
    sa, sb = set(ta), set(tb)
    return 2 * len(sa & sb) / (len(sa) + len(sb))


def company_role(company: Company) -> str | None:
    """manufacturer | brand_owner | distributor | retailer | other, from an
    explicit role/category fact (discovery output, or a user file's column).
    None when the company carries no such fact — callers keep it as-is."""
    for key in _ROLE_KEYS:
        raw = company.facts.get(key)
        if not raw:
            continue
        text = str(raw).strip().lower()
        # Making the product wins over also distributing it ("manufacturer & distributor").
        if any(w in text for w in ("manufactur", "producer", "maker", "processor")):
            return "manufacturer"
        if any(w in text for w in ("distribut", "wholesal", "reseller", "dealer", "trader")):
            return "distributor"
        if "retail" in text or "supermarket" in text:
            return "retailer"
        if "brand" in text:
            return "brand_owner"
        return "other"
    return None


def parent_company_name(company: Company) -> str | None:
    for key in _PARENT_KEYS:
        raw = str(company.facts.get(key) or "").strip()
        if not raw:
            continue
        raw = re.sub(r"^(parent\s*[:\-]|subsidiary of|owned by|part of)\s*", "", raw, flags=re.IGNORECASE)
        if raw.strip().lower() in _NO_PARENT_VALUES:
            return None
        return raw.strip()
    return None


def screen_companies(companies: list[Company]) -> tuple[list[Company], list[dict], list[dict]]:
    """Deterministic pre-estimation screen. Returns (kept, excluded, merged):

    1. Distributors/retailers are dropped (re-sold product is already in a
       manufacturer's revenue) and reported in `excluded` with the reason.
    2. Companies resolving to the same normalized parent are collapsed to
       one entry — the parent itself when it's in the list (top of the
       parent chain), otherwise the first listed — and reported in `merged`.
    """
    kept_by_role: list[Company] = []
    excluded: list[dict] = []
    for c in companies:
        role = company_role(c)
        if role in EXCLUDED_ROLES:
            excluded.append({
                "name": c.name, "role": role,
                "reason": f"{role} -- re-sold product is already counted in the manufacturers' revenue",
            })
        else:
            kept_by_role.append(c)

    # Union-find over normalized names, linking each company to its parent so
    # chains (Bertolli -> Deoleo -> <private-equity owner>) end up in one group.
    root: dict[str, str] = {}

    def find(k: str) -> str:
        root.setdefault(k, k)
        while root[k] != k:
            root[k] = root[root[k]]
            k = root[k]
        return k

    own_keys = [normalize_company_name(c.name) for c in kept_by_role]
    parent_keys: list[str | None] = []
    for c, own in zip(kept_by_role, own_keys):
        parent = parent_company_name(c)
        pkey = normalize_company_name(parent) if parent else None
        parent_keys.append(pkey)
        find(own)
        if pkey:
            root[find(own)] = find(pkey)

    groups: dict[str, list[int]] = {}
    for i, own in enumerate(own_keys):
        groups.setdefault(find(own), []).append(i)

    kept: list[Company] = []
    merged: list[dict] = []
    for members in groups.values():
        member_keys = {own_keys[i] for i in members}
        named_parents = {parent_keys[i] for i in members if parent_keys[i]}
        # The parent at the top of the chain: its own parent isn't a member.
        # Among those, prefer one some other member names as its parent
        # (Deoleo over its brand Bertolli when both point at an outside owner).
        tops = [i for i in members if parent_keys[i] in (None, own_keys[i]) or parent_keys[i] not in member_keys]
        tops.sort(key=lambda i: (own_keys[i] not in named_parents, i))
        rep = tops[0] if tops else members[0]
        kept.append(kept_by_role[rep])
        others = [kept_by_role[i].name for i in members if i != rep]
        if others:
            merged.append({
                "kept": kept_by_role[rep].name,
                "merged": others,
                "reason": "same parent company / brand of a listed parent -- counted once",
            })
    order = {id(c): i for i, c in enumerate(kept_by_role)}
    kept.sort(key=lambda c: order[id(c)])
    return kept, excluded, merged
