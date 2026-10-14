"""Use-case demand builder: the market built bottom-up, the way an analyst
reasons -- industry -> use case -> physical world quantity -> spread over
countries -> price -> value. Works for any product: durable goods counted in
units and consumables measured in tonnes / litres / kg.

Per use case u (base year), the agent gives STRUCTURED inputs and code does
all the arithmetic (numbers written in reasons are never read):

    durable     world units/yr = (new_machines_per_year + installed_machines / product_life_years)
                                 x adoption_share x units_per_machine
                machine = the individual piece of equipment the product goes into
                (the buyer is the machine, never the plant); new builds AND
                replacements are both counted
    consumable  world qty/yr   = world_activity_per_year x adoption_share x dose_per_activity_unit
                (the activity may be the chosen driver's world total)

    quantity(c,u) = world quantity(u) x driver(c) / SUM driver     the driver (World Bank
                    or FAOSTAT) only spreads the world quantity; it is never turned into counts
    value(c,u)    = quantity(c,u) x price(u)                         [US$ / yr]
        price       per the SAME unit as the quantity, at the level the market is
                    valued (manufacturer / ex-factory unless the definition says
                    otherwise). A price quoted per another unit is converted by
                    code (kg <-> t, L <-> t via a stated density), never by the agent.
    value(c)      = SUM over use cases -- use cases are added, never multiplied.

Stability and sanity:
    - the agent designs each use case N_SAMPLES times independently; the MEDIAN
      of each structured input is used, and all samples and their spread are kept;
    - inputs outside generic physical bounds (BOUNDS) are fed back with code's
      arithmetic ("your inputs give N units/yr") and retried; if still out of
      bounds they are kept and flagged;
    - snippets and published figures from the blocklisted publisher(s) (default:
      the market report's own publisher) are dropped before the agent sees them;
    - published market sizes from research firms are context ("what others
      predict"): an industry more than 3x away from all of them is flagged, not
      changed (re-examination is opt-in);
    - when the market itself is an official statistic (a FAOSTAT item), each
      country's use-case quantities are scaled to its official consumption --
      official data sets the level, the use cases only split it (`anchor_to_official`).

Steps: use cases -> structured world inputs (agent, N samples) -> median ->
world quantity (code) -> driver share per country -> official anchor -> roll up
-> ledger. Every number, input and conversion goes into a ledger
(`build_ledger`); the build stays in the backend -- only final values are
written to the Data file.
"""
from __future__ import annotations

import json
import math
import re
import statistics
from pathlib import Path
from typing import Any, Callable

from . import assumptor_regulator as ar
from .scope import partition_dimension

BASE_YEAR = 2025
WB_CACHE = "wb_usecase_drivers_v4.json"  # v4: + machinery manufacturing (where machine builders are)
_WB_SOURCE = {"publisher": "World Bank, World Development Indicators", "url": "https://data.worldbank.org", "tier": 3}

# country-level drivers (World Bank), all available for (nearly) every country
_RAW = {"population": "SP.POP.TOTL", "gdp_usd": "NY.GDP.MKTP.CD", "military_expenditure_usd": "MS.MIL.XPND.CD",
        "manufacturing_value_added_usd": "NV.IND.MANF.CD", "industry_value_added_usd": "NV.IND.TOTL.CD",
        "medhigh_tech_share_of_mfg_pct": "NV.MNF.TECH.ZS.UN", "health_pc_usd": "SH.XPD.CHEX.PC.CD",
        "hospital_beds_per_1000": "SH.MED.BEDS.ZS", "physicians_per_1000": "SH.MED.PHYS.ZS",
        "rd_pct_gdp": "GB.XPD.RSDV.GD.ZS", "researchers_per_million": "SP.POP.SCIE.RD.P6",
        "high_tech_exports_usd": "TX.VAL.TECH.CD", "air_passengers": "IS.AIR.PSGR",
        "fixed_broadband_subscriptions": "IT.NET.BBND", "urban_population": "SP.URB.TOTL",
        "agriculture_value_added_usd": "NV.AGR.TOTL.CD", "arable_land_ha": "AG.LND.ARBL.HA",
        "cereal_production_t": "AG.PRD.CREL.MT", "oil_rents_pct_gdp": "NY.GDP.PETR.RT.ZS",
        "gas_rents_pct_gdp": "NY.GDP.NGAS.RT.ZS", "co2_power_mt": "EN.GHG.CO2.PI.MT.CE.AR5",
        "co2_industrial_combustion_mt": "EN.GHG.CO2.IC.MT.CE.AR5", "co2_industrial_processes_mt": "EN.GHG.CO2.IP.MT.CE.AR5",
        "machinery_share_of_mfg_pct": "NV.MNF.MTRN.ZS.UN"}
MENU = {
    "military_expenditure_usd": "military expenditure, US$ per year",
    "manufacturing_value_added_usd": "manufacturing value added, US$ per year",
    "medhigh_tech_manufacturing_usd": "medium- and high-tech manufacturing value added, US$ per year",
    "machinery_manufacturing_usd": "machinery and transport equipment manufacturing value added, US$ per year -- "
                                   "tracks where machine builders (OEMs) build machines",
    "industry_value_added_usd": "industry value added (incl. construction, energy, mining), US$ per year",
    "agriculture_value_added_usd": "agriculture, forestry and fishing value added, US$ per year",
    "health_expenditure_usd": "current health expenditure, US$ per year",
    "hospital_beds": "hospital beds (count)",
    "physicians": "physicians (count)",
    "rd_expenditure_usd": "R&D expenditure, US$ per year",
    "researchers": "researchers in R&D (count)",
    "high_tech_exports_usd": "high-technology exports, US$ per year",
    "air_passengers": "air passengers carried per year (count)",
    "fixed_broadband_subscriptions": "fixed broadband subscriptions (count)",
    "arable_land_ha": "arable land (hectares)",
    "cereal_production_t": "cereal production (tonnes per year)",
    "urban_population": "urban population (count)",
    "oil_rents_usd": "oil rents (value of crude oil produced less cost), US$ per year -- tracks upstream oil activity",
    "gas_rents_usd": "natural gas rents, US$ per year -- tracks upstream gas activity",
    "co2_power_mt": "CO2 from power generation, Mt per year -- tracks thermal power plant activity",
    "co2_industrial_combustion_mt": "CO2 from industrial fuel combustion, Mt per year -- tracks heavy-industry energy use (refining, chemicals, steel)",
    "co2_industrial_processes_mt": "CO2 from industrial processes, Mt per year -- tracks cement, steel, chemicals and fertilizer output",
    "gdp_usd": "GDP, US$ per year",
    "population": "population (count)",
}
_WB_CODE = {**_RAW, "medhigh_tech_manufacturing_usd": "NV.IND.MANF.CD x NV.MNF.TECH.ZS.UN",
            "health_expenditure_usd": "SH.XPD.CHEX.PC.CD x SP.POP.TOTL", "hospital_beds": "SH.MED.BEDS.ZS x SP.POP.TOTL",
            "physicians": "SH.MED.PHYS.ZS x SP.POP.TOTL", "rd_expenditure_usd": "GB.XPD.RSDV.GD.ZS x NY.GDP.MKTP.CD",
            "researchers": "SP.POP.SCIE.RD.P6 x SP.POP.TOTL", "oil_rents_usd": "NY.GDP.PETR.RT.ZS x NY.GDP.MKTP.CD",
            "gas_rents_usd": "NY.GDP.NGAS.RT.ZS x NY.GDP.MKTP.CD",
            "machinery_manufacturing_usd": "NV.IND.MANF.CD x NV.MNF.MTRN.ZS.UN"}
FAO_PREFIX = "fao:"
# FAOSTAT's "China" row is mainland + Hong Kong + Macao + Taiwan; the parts are listed separately
_FAO_AGGREGATES = {"world", "china"}


def fetch_drivers(existing: dict[str, Any] | None = None) -> dict[str, Any]:
    """World Bank drivers per country. With `existing` (an older cache), only the
    indicators it lacks are downloaded -- the API can take a minute per indicator."""
    if existing:
        out = {iso: dict(c) for iso, c in existing.items()}
    else:
        countries = ar._get(f"{ar._WB}/country", {"format": "json", "per_page": 400})[1]
        out = {c["id"]: {"name": c["name"], "region": str((c.get("region") or {}).get("value") or "").strip()}
               for c in countries if (c.get("region") or {}).get("value") not in (None, "Aggregates")}
    have = {k for c in out.values() for k in c}
    for key, code in _RAW.items():
        if key in have:
            continue
        rows = ar._get(f"{ar._WB}/country/all/indicator/{code}", {"format": "json", "per_page": 20000,
                                                                  "date": "2015:2024"})[1]
        latest: dict[str, tuple[int, float]] = {}
        for r in rows or []:
            iso, val = r.get("countryiso3code"), r.get("value")
            if iso in out and val is not None and (iso not in latest or int(r["date"]) > latest[iso][0]):
                latest[iso] = (int(r["date"]), float(val))
        for iso, (_, val) in latest.items():
            out[iso][key] = val
    for c in out.values():
        pop, gdp = c.get("population"), c.get("gdp_usd")
        if pop:
            if c.get("health_pc_usd"):
                c["health_expenditure_usd"] = c["health_pc_usd"] * pop
            if c.get("hospital_beds_per_1000"):
                c["hospital_beds"] = c["hospital_beds_per_1000"] * pop / 1000
            if c.get("physicians_per_1000"):
                c["physicians"] = c["physicians_per_1000"] * pop / 1000
            if c.get("researchers_per_million"):
                c["researchers"] = c["researchers_per_million"] * pop / 1e6
        if gdp and c.get("rd_pct_gdp"):
            c["rd_expenditure_usd"] = c["rd_pct_gdp"] / 100 * gdp
        if gdp and c.get("oil_rents_pct_gdp") is not None:
            c["oil_rents_usd"] = c["oil_rents_pct_gdp"] / 100 * gdp
        if gdp and c.get("gas_rents_pct_gdp") is not None:
            c["gas_rents_usd"] = c["gas_rents_pct_gdp"] / 100 * gdp
        if c.get("manufacturing_value_added_usd") and c.get("machinery_share_of_mfg_pct"):
            c["machinery_manufacturing_usd"] = c["manufacturing_value_added_usd"] * c["machinery_share_of_mfg_pct"] / 100
        if c.get("manufacturing_value_added_usd") and c.get("medhigh_tech_share_of_mfg_pct"):
            c["medhigh_tech_manufacturing_usd"] = c["manufacturing_value_added_usd"] * c["medhigh_tech_share_of_mfg_pct"] / 100
    return out


def load_drivers(cache_dir: Path) -> dict[str, Any]:
    """The current driver cache; when it is missing, the newest older version is
    extended with only the indicators it lacks (not re-downloaded from scratch)."""
    path = cache_dir / WB_CACHE
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    older = sorted(cache_dir.glob("wb_usecase_drivers*.json"), key=lambda q: q.stat().st_mtime, reverse=True)
    existing = json.loads(older[0].read_text(encoding="utf-8")) if older else None
    data = fetch_drivers(existing)
    cache_dir.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data), encoding="utf-8")
    return data


# ---- drivers per geography ----------------------------------------------------
def _load_fao_cache() -> dict[str, Any]:
    try:
        from me_engine.demandCountry_approach2A.topdown import official_data as od
        return od.load_cache()
    except Exception:  # noqa: BLE001 -- no FAOSTAT data: World Bank drivers only
        return {}


def _iso_resolver(drivers: dict[str, Any]) -> Callable[[str], str | None]:
    """FAOSTAT area name -> ISO3 (World Bank names + the assumptor's alias table)."""
    from me_engine.demandCountry_approach2A.assumptor.covariates import _ALIASES, _norm as cnorm

    names: dict[str, str] = {}
    for iso, c in drivers.items():
        names.setdefault(cnorm(c["name"]), iso)
    for iso, c in drivers.items():
        names.setdefault(cnorm(c["name"].split(",")[0]), iso)

    def iso3(area: str) -> str | None:
        k = cnorm(area)
        return _ALIASES.get(k) or names.get(k) or names.get(cnorm(re.sub(r"\(.*?\)", "", k)))
    return iso3


def fao_by_geo(dataset: str, item: str, drivers: dict[str, Any], geos: list[str], fao_cache: dict[str, Any],
               base_year: int = BASE_YEAR) -> tuple[dict[str, float], dict[str, Any]]:
    """FAOSTAT tonnes per geography bucket at the base year (each country's last
    3 data years moved with its own trend), plus what was covered."""
    from me_engine.demandCountry_approach2A.topdown import official_data as od

    areas = (fao_cache.get(dataset) or {}).get(item) or {}
    iso3 = _iso_resolver(drivers)
    out = {g: 0.0 for g in geos}
    latest, unplaced = [], []
    for area, series in areas.items():
        if area.strip().lower() in _FAO_AGGREGATES:
            continue
        iso = iso3(area)
        g = ar.bucket_of(iso, drivers[iso], geos) if iso in drivers else None
        if not g:
            unplaced.append(area)
            continue
        value, year, _ = od._project(series, base_year)
        if value > 0:
            out[g] += value
            latest.append(year)
    meta = {"source": f"FAOSTAT {od.DATASETS[dataset]['label']} '{item}', {od.DATASETS[dataset]['element']}",
            "url": od.DATASETS[dataset]["page"], "latest_year": max(latest) if latest else None,
            "oldest_year": min(latest) if latest else None, "unplaced_areas": unplaced[:40]}
    return out, meta


class DriverTable:
    """Driver values per geography (World Bank or FAOSTAT), cached, with where
    each came from -- the ledger needs to know which numbers were imputed."""

    def __init__(self, wb: dict[str, Any], geos: list[str], fao_cache: dict[str, Any] | None = None) -> None:
        self.wb, self.geos, self.fao = wb, geos, fao_cache or {}
        self._vals: dict[str, dict[str, float]] = {}
        self._meta: dict[str, dict[str, Any]] = {}

    def fao_items(self) -> list[tuple[str, str]]:
        return [(ds, item) for ds in ("fbs", "rfb") for item in sorted(self.fao.get(ds) or {})]

    def valid(self, key: str) -> bool:
        if key in MENU:
            return True
        if key.startswith(FAO_PREFIX):
            parts = key.split(":", 2)
            return len(parts) == 3 and parts[2] in (self.fao.get(parts[1]) or {})
        return False

    def by_geo(self, key: str) -> dict[str, float]:
        if key not in self._vals:
            if key.startswith(FAO_PREFIX):
                _, ds, item = key.split(":", 2)
                vals, meta = fao_by_geo(ds, item, self.wb, self.geos, self.fao)
                meta.update(unit="t per year", label="derived", imputed_share={g: 0.0 for g in self.geos})
            else:
                vals, imputed = self._wb(key)
                meta = {"source": f"{_WB_SOURCE['publisher']} {_WB_CODE.get(key, key)}", "url": _WB_SOURCE["url"],
                        "unit": MENU.get(key, key), "imputed_share": imputed}
            self._vals[key], self._meta[key] = vals, meta
        return self._vals[key]

    def meta(self, key: str) -> dict[str, Any]:
        self.by_geo(key)
        return self._meta[key]

    def _wb(self, key: str) -> tuple[dict[str, float], dict[str, float]]:
        """Indicator summed over each bucket; a country missing it contributes GDP x
        (its region's median indicator/GDP ratio, else the world median) so a gap
        never zeroes a bucket. Returns the imputed share too."""
        ratio: dict[str, list[float]] = {}
        for c in self.wb.values():
            if c.get(key) and c.get("gdp_usd"):
                ratio.setdefault(c["region"], []).append(c[key] / c["gdp_usd"])
        med = {r: statistics.median(v) for r, v in ratio.items()}
        world_med = statistics.median(x for v in ratio.values() for x in v) if ratio else None
        out = {g: 0.0 for g in self.geos}
        imputed = {g: 0.0 for g in self.geos}
        for iso, c in self.wb.items():
            g = ar.bucket_of(iso, c, self.geos)
            if not g:
                continue
            v = c.get(key)
            r = med.get(c.get("region"), world_med)
            if v is None and c.get("gdp_usd") and r is not None:
                v = c["gdp_usd"] * r
                imputed[g] += v
            out[g] += float(v or 0.0)
        return out, {g: (imputed[g] / out[g] if out[g] else 0.0) for g in self.geos}


# ---- use cases ----------------------------------------------------------------
def use_cases(nesting: dict[str, Any]) -> tuple[str, list[dict[str, str]]]:
    dim = partition_dimension(nesting)
    tree = nesting.get(dim) or {}
    out = []
    for industry in tree.get("__root__") or []:
        kids = tree.get(industry) or []
        for uc in kids or [industry]:
            out.append({"industry": industry, "use_case": uc, "own": not kids})
    return dim, out


def _fmt(v: float) -> str:
    a = abs(v)
    return f"{v / 1e12:.2f}tn" if a >= 1e12 else f"{v / 1e9:.2f}bn" if a >= 1e9 else f"{v / 1e6:.2f}mn" if a >= 1e6 else f"{v:,.0f}"


def _example_geos(geos: list[str]) -> list[str]:
    preferred = [g for g in ("U.S.", "China", "Germany", "India") if g in geos]
    return (preferred + [g for g in geos if g not in preferred])[:4]


# ---- the agent's world quantity: structured inputs, multiplied by code ------------------
ARCHETYPES = ("durable", "consumable")
DURABLE_INPUTS = ("new_machines_per_year", "installed_machines", "product_life_years", "adoption_share",
                  "units_per_machine")
CONSUMABLE_INPUTS = ("world_activity_per_year", "adoption_share", "dose_per_activity_unit")
# generic physical bounds (inclusive); every other input is a count / amount and must be > 0
BOUNDS = {"product_life_years": (5.0, 40.0), "adoption_share": (0.0, 1.0), "units_per_machine": (1.0, 4.0)}
N_SAMPLES = 3            # independent designs per use case; the median of each input is used
ATTEMPTS_PER_SAMPLE = 3  # first answer + 2 retries with code's feedback
DRIVER_ACTIVITY = "driver"  # consumable: the chosen driver itself is the world activity
# the publisher of the market report being built is never evidence or context for it
DEFAULT_BLOCKLIST = ("Coherent Market Insights", "coherentmarketinsights")


def archetype(spec: dict[str, Any]) -> str:
    a = str(spec.get("archetype") or "").strip().lower()
    return a if a in ARCHETYPES else ""


FIRST_FIT_INPUTS = ("new_machines_per_year", "adoption_share", "units_per_machine")


def counts_replacements(spec: dict[str, Any]) -> bool:
    """False when the market counts first-fit units on new machines only (an OEM
    market): the replacement term is then left out of the world quantity."""
    return spec.get("count_replacements", True) is not False


def input_names(spec: dict[str, Any]) -> tuple[str, ...]:
    if archetype(spec) == "durable":
        return DURABLE_INPUTS if counts_replacements(spec) else FIRST_FIT_INPUTS
    return CONSUMABLE_INPUTS


def _entry_value(entry: Any) -> float | str | None:
    v = entry.get("value") if isinstance(entry, dict) else entry
    if isinstance(v, str) and v.strip().lower() == DRIVER_ACTIVITY:
        return DRIVER_ACTIVITY
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f if math.isfinite(f) else None


def activity_from_driver(spec: dict[str, Any]) -> bool:
    return (archetype(spec) == "consumable"
            and _entry_value((spec.get("world") or {}).get("world_activity_per_year")) == DRIVER_ACTIVITY)


def world_inputs(spec: dict[str, Any]) -> dict[str, float]:
    """The structured world inputs as numbers (numbers in reason text are never read)."""
    if not archetype(spec):
        raise ValueError(f"archetype must be one of {ARCHETYPES}, got '{spec.get('archetype')}'")
    w = spec.get("world") if isinstance(spec.get("world"), dict) else {}
    out, missing = {}, []
    for k in input_names(spec):
        v = _entry_value(w.get(k))
        if k == "world_activity_per_year" and v == DRIVER_ACTIVITY:
            continue
        if v is None or v == DRIVER_ACTIVITY:
            missing.append(k)
        else:
            out[k] = float(v)
    if missing:
        raise ValueError(f"missing or non-numeric world inputs: {', '.join(missing)}")
    return out


def world_units(spec: dict[str, Any], driver_total: float | None = None) -> float:
    """World quantity bought per year, computed by code:
    durable    (new machines/yr + installed machines / product life) x adoption x units per machine
    consumable world activity x adoption x dose (activity = the driver's world total when stated so)."""
    x = world_inputs(spec)
    if archetype(spec) == "durable":
        if not counts_replacements(spec):
            return x["new_machines_per_year"] * x["adoption_share"] * x["units_per_machine"]
        if x["product_life_years"] <= 0:
            raise ValueError("product_life_years must be > 0")
        return ((x["new_machines_per_year"] + x["installed_machines"] / x["product_life_years"])
                * x["adoption_share"] * x["units_per_machine"])
    activity = driver_total if activity_from_driver(spec) else x["world_activity_per_year"]
    if activity is None:
        raise ValueError("world activity is the driver, but no driver total was given")
    return activity * x["adoption_share"] * x["dose_per_activity_unit"]


def world_formula(spec: dict[str, Any]) -> str:
    if archetype(spec) == "durable" and not counts_replacements(spec):
        return "new_machines_per_year x adoption_share x units_per_machine (first fit on new machines only)"
    if archetype(spec) == "durable":
        return "(new_machines_per_year + installed_machines / product_life_years) x adoption_share x units_per_machine"
    return "world_activity_per_year x adoption_share x dose_per_activity_unit"


def bound_violations(spec: dict[str, Any]) -> list[str]:
    out = []
    for k, v in world_inputs(spec).items():
        if k in BOUNDS:
            lo, hi = BOUNDS[k]
            if v < lo or v > hi or (lo == 0 and v <= 0):
                out.append(f"{k} = {v:g} is outside {lo:g}-{hi:g}")
        elif v <= 0:
            out.append(f"{k} = {v:g} must be > 0")
    return out


def computed_text(spec: dict[str, Any], n: float, driver_total: float | None = None) -> str:
    """Code's arithmetic on the agent's inputs, in words the agent can check."""
    x = world_inputs(spec)
    unit = spec.get("quantity_unit") or "units"
    if archetype(spec) == "durable" and not counts_replacements(spec):
        return (f"your inputs give {x['new_machines_per_year']:,.0f} new machines/yr x {x['adoption_share']:g} adoption x "
                f"{x['units_per_machine']:g} per machine = {n:,.0f} {unit}/yr (first fit only)")
    if archetype(spec) == "durable":
        a, u = x["adoption_share"], x["units_per_machine"]
        new, rep = x["new_machines_per_year"] * a * u, x["installed_machines"] / x["product_life_years"] * a * u
        return (f"your inputs give ({x['new_machines_per_year']:,.0f} new machines/yr + {x['installed_machines']:,.0f} "
                f"installed / {x['product_life_years']:g} yr) x {a:g} adoption x {u:g} per machine = {n:,.0f} {unit}/yr "
                f"({new:,.0f} for new builds + {rep:,.0f} replacements)")
    act = driver_total if activity_from_driver(spec) else x["world_activity_per_year"]
    return (f"your inputs give {act:,.4g} activity x {x['adoption_share']:g} adoption x "
            f"{x['dose_per_activity_unit']:g} dose = {n:,.0f} {unit}/yr")


CIRCULAR = re.compile(r"(market|revenue|sales|spend)[^.]{0,80}(\$|usd|billion|bn\b|million|mn\b)[^.]{0,120}"
                      r"(divid|/|÷|per unit|implies|implied|\bat (an? )?(average |typical )?(\w+ ){0,2}price)", re.I)


def circular_reason(spec: dict[str, Any]) -> bool:
    """True if a world input was backed out of a published market size."""
    w = spec.get("world") if isinstance(spec.get("world"), dict) else {}
    reasons = [str(e.get("reason") or "") for e in w.values() if isinstance(e, dict)]
    return any(CIRCULAR.search(t) for t in reasons)


# ---- publisher blocklist --------------------------------------------------------------
def _compact(text: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def blocked(text: str, blocklist: tuple[str, ...] | list[str] = DEFAULT_BLOCKLIST) -> bool:
    """True when the text names a blocked publisher. An all-caps term is an
    acronym (whole word, case-sensitive); any other term matches ignoring case,
    spaces and punctuation (so a name also matches its domain)."""
    for term in blocklist:
        t = str(term or "").strip()
        if not t:
            continue
        if t.isupper() and " " not in t:
            if re.search(rf"(?<![A-Za-z0-9]){re.escape(t)}(?![A-Za-z0-9])", str(text or "")):
                return True
        elif _compact(t) and _compact(t) in _compact(text):
            return True
    return False


def _snippet_text(s: Any) -> str:
    if isinstance(s, dict):
        return " ".join(str(s.get(k) or "") for k in ("text", "title", "url", "domain", "source"))
    return str(s)


def filter_snippets(snippets: list[Any], blocklist: tuple[str, ...] | list[str] = DEFAULT_BLOCKLIST
                    ) -> tuple[list[Any], int]:
    kept = [s for s in snippets if not blocked(_snippet_text(s), blocklist)]
    return kept, len(snippets) - len(kept)


def filter_published(figures: list[Any], blocklist: tuple[str, ...] | list[str] = DEFAULT_BLOCKLIST
                     ) -> tuple[list[Any], int]:
    kept = [f for f in figures
            if not blocked(" ".join(str(v) for v in f.values()) if isinstance(f, dict) else str(f), blocklist)]
    return kept, len(figures) - len(kept)


def _evidence(snippets: list[Any]) -> list[dict[str, str]]:
    out = []
    for i, s in enumerate(snippets[:10], start=1):
        if isinstance(s, dict):
            out.append({"id": f"E{i}", "source": str(s.get("domain") or s.get("source") or "web search"),
                        "text": str(s.get("text") or ""), "url": str(s.get("url") or "")})
        else:
            out.append({"id": f"E{i}", "source": "web search", "text": str(s)})
    return out


# ---- units and price ------------------------------------------------------------
_COUNT_WORDS = {"unit", "units", "each", "piece", "pieces", "pc", "pcs", "item", "items", "number", "no"}


def _unit_words(text: str) -> set[str]:
    return {w.rstrip("s") for w in re.findall(r"[a-z]+", str(text or "").lower()) if w not in ("of", "per", "one", "a")}


_CURRENCY = re.compile(r"\b(usd|us\$|dollars?)\b|\$", re.I)
_MONEY_SCALE = {"thousand": 1e3, "k": 1e3, "million": 1e6, "mn": 1e6, "mio": 1e6, "m": 1e6, "billion": 1e9,
                "bn": 1e9}


def _currency_scale(per: str) -> tuple[float, str]:
    """('million USD' -> 1e6, ''), ('thousand US$ per tonne' -> 1e3, 'tonne'): a `per`
    that names a currency is a scale on the price, not the unit it is quoted per."""
    if not _CURRENCY.search(per):
        return 1.0, per
    words = re.findall(r"[a-z]+", _CURRENCY.sub(" ", per.lower()))
    scale = math.prod(_MONEY_SCALE[w] for w in words if w in _MONEY_SCALE) or 1.0
    rest = " ".join(w for w in words if w not in _MONEY_SCALE and w not in ("per", "a", "each", "one"))
    return float(scale), rest


def price_conversion(spec: dict[str, Any]) -> float:
    """Multiplier turning the quoted price (US$ per `per`) into US$ per
    `quantity_unit`. Raises ValueError when the two units can't be reconciled."""
    from me_engine.demandCountry_approach2A.engine.units import _base_unit, unit_scale

    qu = str(spec.get("quantity_unit") or "").strip()
    p = spec.get("price_usd_per_unit")
    scale, per = _currency_scale(str((p.get("per") if isinstance(p, dict) else "") or "").strip())
    if scale != 1.0:
        return scale * price_conversion({**spec, "price_usd_per_unit": {**p, "per": per}})
    if not qu or not per:
        return 1.0  # price stated per quantity unit
    a, b = _base_unit(qu), _base_unit(per)
    if a and b:
        density = spec.get("density_kg_per_l")
        if a[0] != b[0] and not (isinstance(density, (int, float)) and density > 0):
            raise ValueError(f"price is per '{per}' but quantity is in '{qu}': give density_kg_per_l to convert")
        return unit_scale(qu, per, float(density) if density else None) or 1.0
    if a or b:
        raise ValueError(f"price is per '{per}' but quantity is in '{qu}': one is a mass/volume unit, the other is "
                         "not -- quote the price per quantity_unit")
    wq, wp = _unit_words(qu), _unit_words(per)
    if wq & wp or wp <= _COUNT_WORDS or wq <= _COUNT_WORDS:
        return 1.0
    raise ValueError(f"price is per '{per}' but quantity is in '{qu}': quote the price per quantity_unit")


def price_of(spec: dict[str, Any]) -> float:
    """US$ per quantity_unit (converted by code); 0 when no usable price."""
    p = spec.get("price_usd_per_unit")
    try:
        quoted = float(p.get("value") if isinstance(p, dict) else p)
    except (TypeError, ValueError):
        return 0.0
    return quoted * price_conversion(spec)


# ---- the agent's design ---------------------------------------------------------------
_DESIGN_SYSTEM = (
    "You are a market analyst sizing demand bottom-up, one use case at a time, for ANY product. Published "
    "market-research figures are CONTEXT (what other firms predict): read them to understand the market, but do NOT "
    "derive your number from them; build it from physical demand.\n"
    "FIRST decide the product's ARCHETYPE:\n"
    "  'durable'    : a unit of equipment or a component that goes into a MACHINE (the individual piece of "
    "equipment it is built into or serves) and is bought when machines are built and when the one in service wears "
    "out. Counted in units.\n"
    "  'consumable' : used up in proportion to an activity (tonnes, litres, kg of product per unit of activity).\n"
    "DURABLE -- give five numbers; CODE multiplies them, you do not:\n"
    "  units per year = (new_machines_per_year + installed_machines / product_life_years) x adoption_share x "
    "units_per_machine\n"
    "  machine              : ONE individual machine of this use case (one turbine, one compressor, one pump, one "
    "vehicle). The buyer is the MACHINE, never the plant, site, company or project: a plant running six machines is "
    "six machines.\n"
    "  new_machines_per_year: machines of this use case built / shipped / ordered worldwide per year -- every size "
    "class and every end-user industry of this use case (utility and industrial, large and small), whole world.\n"
    "  installed_machines   : machines of this use case in operation worldwide (the fleet).\n"
    "  product_life_years   : years one unit of the PRODUCT serves in a machine before it is replaced (if it lasts "
    "as long as the machine, the machine's life). Physical range 5-40 years.\n"
    "  adoption_share       : share of those machines that use this product at all (0-1).\n"
    "  units_per_machine    : product units in one machine (1-4).\n"
    "  New builds AND replacements are both demand: never leave either out.\n"
    "CONSUMABLE -- give three numbers; code multiplies them: world_activity_per_year (in `activity_unit`; or the "
    "string \"driver\" when the driver you chose IS that activity, e.g. 'fao:fbs:Vegetables' tonnes) x "
    "adoption_share (0-1) x dose_per_activity_unit (quantity_unit per one activity_unit).\n"
    "DRIVER: pick ONE country-level driver from the menu that tracks where these machines (or this activity) are. "
    "Code uses it ONLY to split your world quantity across countries in proportion (a country's share = its driver "
    "value / the world total). Do NOT turn the driver into machine or plant counts, and do not compute with it.\n"
    "ARITHMETIC: give every input as a plain number; code does all arithmetic. Numbers written in reasons are "
    "ignored. Never get a count by dividing a market's US$ size by a price.\n"
    "Every input needs a unit, a reason (where the number comes from) and a basis: an evidence id (e.g. 'E3') if "
    "evidence states or directly supports it, else 'assumption'. Prefer evidence on machines: units shipped or "
    "ordered per year, installed base, service life.\n"
    "PRICE: the price of one `quantity_unit` (or say in `per` which unit it is quoted per -- code converts kg/t/L; "
    "for litres<->tonnes give `density_kg_per_l`). Use the price level the market is valued at: what the "
    "product's makers receive (manufacturer / ex-factory), not the retail price of a finished good containing it. "
    "A price written in thousands or millions of US$ must say so in `per` (e.g. 'million USD').\n"
    "SCOPE: the product is everything the taxonomy lists; the market's name is only a label -- do not multiply in "
    "narrowing factors such as 'high-energy share'. Count only the product itself, not the larger product it goes "
    "into, and count each physical use once. Return json."
)

_INPUT_JSON = '{"value": <number>, "unit": "<unit>", "basis": "<E-id or assumption>", "reason": "<why>"}'


def design(llm: Any, market: str, uc: dict[str, str], definition: str, evidence: list[dict[str, str]],
           published: list[dict[str, Any]], driver_examples: dict[str, dict[str, float]],
           feedback: str = "", fao_items: list[tuple[str, str]] | None = None, *, sample: int = 1,
           n_samples: int = 1, salt: str = "") -> dict[str, Any]:
    ev_text = "\n".join(f"[{e['id']}] ({e['source']}) {e['text']}" for e in evidence) or "(no use-case evidence found)"
    examples = {k: {c: _fmt(v) for c, v in vals.items()} for k, vals in driver_examples.items()}
    fao_text = ""
    if fao_items:
        by_ds: dict[str, list[str]] = {}
        for ds, item in fao_items:
            by_ds.setdefault(ds, []).append(item)
        fao_text = ("\nFAOSTAT drivers, tonnes per year per country (key 'fao:<dataset>:<item>'; fbs = food domestic "
                    "supply, rfb = fertilizer apparent consumption):\n"
                    + "\n".join(f"  {ds}: " + "; ".join(items) for ds, items in by_ds.items()))
    # the sample index varies the prompt so the LLM cache (keyed by prompt) returns an independent answer
    sample_line = (f"\nThis is independent estimate #{sample} of {n_samples}"
                   + (f" (run {salt})" if salt else "") + ": reason it through from scratch.\n")
    return llm.complete_json(
        _DESIGN_SYSTEM,
        f"Market: {market}\nIndustry: {uc['industry']}\nUse case: {uc['use_case']}\n"
        f"How the market's taxonomy defines the product: {definition}\n\n"
        f"World Bank driver menu (key: meaning): {json.dumps(MENU)}{fao_text}\n"
        f"World Bank driver values for orientation: {json.dumps(examples)}\n\n"
        f"Evidence about this use case:\n{ev_text}\n\n"
        f"What research firms predict for the industry (context only, never the answer): "
        f"{json.dumps(published, ensure_ascii=False)[:1500]}\n"
        + sample_line
        + (f"\nREVIEW OF YOUR PREVIOUS ATTEMPT: {feedback}\n" if feedback else "")
        + '\nReturn JSON: {"archetype": "<durable | consumable>", '
          '"driver": "<menu key or fao:<dataset>:<item> -- used only to split the world quantity across countries>", '
          '"buyers": "<who buys it, in plain words>", "machine": "<durable: what ONE machine is>", '
          '"quantity_unit": "<units for durables; t | kg | L for consumables>", '
          '"units_label": "<what one quantity_unit of the product is>", '
          '"activity_unit": "<consumable: unit of world_activity_per_year>", '
          '"world": {"new_machines_per_year": ' + _INPUT_JSON + ', "installed_machines": ' + _INPUT_JSON
        + ', "product_life_years": ' + _INPUT_JSON + ', "adoption_share": ' + _INPUT_JSON
        + ', "units_per_machine": ' + _INPUT_JSON + '}  (consumable instead: {"world_activity_per_year": '
        + _INPUT_JSON + ' or value "driver", "adoption_share": {...}, "dose_per_activity_unit": {...}}), '
          '"price_usd_per_unit": {"value": <number>, "per": "<unit the price is quoted per>", "basis": "<E-id or '
          'assumption>", "reason": "<why, and which price level>"}, "density_kg_per_l": <number or null>, '
          '"reasoning": "<3-4 sentences: the demand logic in plain words>"}',
    )


def _evaluate(cand: dict[str, Any], table: "DriverTable", geos: list[str]) -> dict[str, Any]:
    """Code's evaluation of one design; ValueError when it cannot be used at all."""
    key = str(cand.get("driver") or "")
    if not table.valid(key):
        raise ValueError(f"unknown driver '{key}' -- pick a key from the menu")
    dv = table.by_geo(key)
    total = math.fsum(max(0.0, dv.get(g, 0.0)) for g in geos)
    if total <= 0:
        raise ValueError(f"driver '{key}' is zero in every geography -- pick another")
    n = world_units(cand, total)
    if not (n > 0 and math.isfinite(n)):
        raise ValueError(f"the inputs give {n} units per year -- no demand")
    price = price_of(cand)
    if price <= 0:
        raise ValueError("no usable price")
    return {"world_units": n, "driver_total": total, "price": price, "violations": bound_violations(cand),
            "circular": circular_reason(cand)}


def _sample(ask: Callable[[str, int], dict[str, Any]], k: int, table: "DriverTable", geos: list[str]
            ) -> dict[str, Any] | None:
    """One independent design, retried with code's feedback; out-of-bounds inputs
    that survive every retry are kept and flagged."""
    feedback, log, fallback = "", [], None
    for attempt in range(ATTEMPTS_PER_SAMPLE):
        try:
            cand = ask(feedback, k)
            ev = _evaluate(cand, table, geos)
        except Exception as exc:  # noqa: BLE001
            feedback = (f"Your answer could not be used ({exc}). Give the archetype, a driver key from the menu, "
                        "every world input as a plain positive number, and a positive price with the unit it is "
                        "quoted per.")
            log.append(feedback)
            continue
        rec = {"sample": k, "attempts": attempt + 1, "spec": cand, **ev, "feedback": list(log),
               "kept_out_of_bounds": False}
        problems = list(ev["violations"])
        if ev["circular"]:
            problems.append("a count was obtained by dividing a published market size by a price -- rebuild it "
                            "from machines (or activity)")
        if not problems:
            return rec
        fallback = {**rec, "kept_out_of_bounds": True}
        feedback = (f"Code computed your world quantity: {computed_text(cand, ev['world_units'], ev['driver_total'])}. "
                    "Problems: " + "; ".join(problems) + ". Re-check these inputs -- count individual machines, not "
                    "plants; include new builds and replacements; life is the product's service life in years -- and "
                    "correct them.")
        log.append(feedback)
    if fallback:
        fallback["feedback"] = log
    return fallback


def _unit_key(text: Any) -> str:
    return re.sub(r"s$", "", str(text or "").strip().lower())


def _sample_group(s: dict[str, Any]) -> tuple[str, str, str]:
    """Samples are comparable input-by-input only when they count the same thing."""
    sp = s["spec"]
    if archetype(sp) == "durable":
        return ("durable", _unit_key(sp.get("quantity_unit")), "")
    act = (f"driver:{sp.get('driver')}" if activity_from_driver(sp)
           else _unit_key(((sp.get("world") or {}).get("world_activity_per_year") or {}).get("unit")
                          if isinstance((sp.get("world") or {}).get("world_activity_per_year"), dict)
                          else sp.get("activity_unit")))
    return ("consumable", _unit_key(sp.get("quantity_unit")), act)


def _median_entry(entries: list[Any]) -> dict[str, Any]:
    """Median of one input across samples, keeping the basis and reason of the
    sample it came from (with two samples: their mean, cited only if both cite)."""
    pairs = sorted(((float(_entry_value(e)), e if isinstance(e, dict) else {"value": e}) for e in entries),
                   key=lambda p: p[0])
    n = len(pairs)
    if n % 2:
        v, e = pairs[n // 2]
        return {**e, "value": v}
    (a, ea), (b, eb) = pairs[n // 2 - 1], pairs[n // 2]
    ia, ib = re.findall(r"E\d+", str(ea.get("basis") or "")), re.findall(r"E\d+", str(eb.get("basis") or ""))
    return {"value": (a + b) / 2, "unit": ea.get("unit") or eb.get("unit"),
            "basis": " ".join(dict.fromkeys(ia + ib)) if ia and ib else "assumption",
            "reason": f"mean of the two middle estimates {a:g} ({ea.get('reason') or ''}) and {b:g} "
                      f"({eb.get('reason') or ''})"}


def _spread(values: list[float]) -> dict[str, Any]:
    lo, hi = min(values), max(values)
    return {"values": values, "min": lo, "median": statistics.median(values), "max": hi,
            "max_over_min": (hi / lo) if lo > 0 else None}


def combine(samples: list[dict[str, Any]]) -> tuple[dict[str, Any], dict[str, Any]]:
    """The median of each structured input over the comparable samples (not the
    median of their final values). The price, units and wording come from the
    first sample of the largest comparable group; the driver is the most chosen."""
    groups: dict[tuple[str, str, str], list[dict[str, Any]]] = {}
    for s in samples:
        groups.setdefault(_sample_group(s), []).append(s)
    best = max(groups.values(), key=len)
    base = best[0]["spec"]
    drivers = [str(s["spec"].get("driver")) for s in best]
    driver = max(dict.fromkeys(drivers), key=lambda d: (drivers.count(d), -drivers.index(d)))
    world: dict[str, Any] = {}
    spread: dict[str, Any] = {}
    for k in input_names(base):
        if k == "world_activity_per_year" and activity_from_driver(base):
            world[k] = base["world"][k]
            continue
        entries = [s["spec"]["world"][k] for s in best]
        world[k] = _median_entry(entries)
        spread[k] = _spread([float(_entry_value(e)) for e in entries])
    spread["world_units_per_sample"] = _spread([s["world_units"] for s in samples])
    used = [s["sample"] for s in best]
    return {**base, "driver": driver, "world": world}, {
        "used_samples": used, "excluded_samples": [s["sample"] for s in samples if s["sample"] not in used],
        "driver_votes": drivers, "spread": spread}


def _search_queries(uc: dict[str, str], term: str) -> list[str]:
    """Evidence on the machine fleet, not on market size."""
    u = uc["use_case"]
    return [f"{u} units shipped orders per year worldwide",
            f"{u} installed base number of units in operation worldwide",
            f"{term} {u} service life replacement interval years",
            f"{u} {uc['industry']} {term} consumption volume per year"]


def basis_counts(spec: dict[str, Any], evidence: list[dict[str, str]]) -> dict[str, int]:
    """How many world inputs (and the price) cite an evidence id vs are assumptions."""
    entries = [e for k, e in (spec.get("world") or {}).items()
               if not (k == "world_activity_per_year" and activity_from_driver(spec))]
    labels = [_label(e, evidence)[0] for e in entries + [spec.get("price_usd_per_unit")]]
    return {"cited": sum(lab != "assumption" for lab in labels), "assumption": labels.count("assumption")}


# ---- the build --------------------------------------------------------------------
def build(llm: Any, spec: dict[str, Any], nesting: dict[str, Any], market: str, cache_dir: Path,
          search: Callable[[list[str]], list[Any]], published_by_industry: dict[str, list[dict[str, Any]]],
          say: Callable[[str], None] = print, fao_cache: dict[str, Any] | None = None, *,
          n_samples: int = N_SAMPLES, salt: str = "", blocklist: tuple[str, ...] | list[str] = DEFAULT_BLOCKLIST,
          reexamine_published: bool = False) -> dict[str, Any]:
    """`reexamine_published`: False (default) only FLAGS an industry more than 3x
    away from every published figure; True redesigns it once and keeps the
    redesign if it moves toward the publishers' midpoint."""
    from .scope import segment_definition

    dim, ucs = use_cases(nesting)
    geos = spec["geographies"]
    table = DriverTable(load_drivers(cache_dir), geos,
                        _load_fao_cache() if fao_cache is None else fao_cache)
    examples = {}
    for key in MENU:
        vals = table.by_geo(key)
        examples[key] = {g: vals.get(g, 0.0) for g in _example_geos(geos)}
    term = market.replace("Global", "").replace("Market", "").strip()
    pubs_by_ind = {ind: filter_published(figs or [], blocklist) for ind, figs in published_by_industry.items()}
    results = []
    for uc in ucs:
        snippets, n_snip = filter_snippets(search(_search_queries(uc, term)), blocklist)
        evidence = _evidence(snippets)
        published, n_pub = pubs_by_ind.get(uc["industry"], ([], 0))
        say(f"[use case] {uc['use_case']}: blocklist removed {n_snip} search snippet(s), {n_pub} published figure(s)")
        definition = segment_definition(nesting, dim, uc["industry"])

        def ask(feedback: str, k: int, uc=uc, definition=definition, evidence=evidence, published=published):
            return design(llm, market, uc, definition, evidence, published, examples, feedback, table.fao_items(),
                          sample=k, n_samples=n_samples, salt=salt)

        samples = [s for k in range(1, n_samples + 1) if (s := _sample(ask, k, table, geos))]
        if not samples:
            say(f"[use case] {uc['use_case']}: no usable design in {n_samples} samples")
            continue
        cspec, info = combine(samples)
        try:
            ev = _evaluate(cspec, table, geos)
        except Exception as exc:  # noqa: BLE001
            say(f"[use case] {uc['use_case']}: median design unusable ({exc})")
            continue
        flags = [f"sample {s['sample']} kept out of bounds after {s['attempts']} attempts: "
                 + "; ".join(s["violations"] + (["circular count"] if s["circular"] else []))
                 for s in samples if s["kept_out_of_bounds"]]
        if ev["violations"]:
            flags.append("median inputs out of bounds: " + "; ".join(ev["violations"]))
        if info["excluded_samples"]:
            flags.append(f"samples {info['excluded_samples']} count a different unit and were left out of the median")
        results.append(_item(uc, cspec, table, geos, ev, evidence=evidence, published=published, samples=samples,
                             info=info, flags=flags, blocked={"snippets": n_snip, "published_figures": n_pub}))
        r = results[-1]
        say(f"[use case] {uc['industry']} / {uc['use_case']}: {r['world_units']:,.0f} {cspec.get('quantity_unit') or 'units'}"
            f"/yr x ${r['price_per_quantity_unit']:,.4g} = ${r['world_mn']:,.0f}M via {cspec['driver']}; "
            f"units spread over samples x{(info['spread']['world_units_per_sample']['max_over_min'] or 0):.2f}"
            + (f"; FLAGS: {' | '.join(flags)}" if flags else ""))

    # sanity: an industry far from every published figure is flagged (or, if switched on, re-examined once)
    flags = []
    for ind in sorted({r["industry"] for r in results}):
        pubs = sorted(p["usd_mn"] for p in pubs_by_ind.get(ind, ([], 0))[0] if p.get("usd_mn"))
        tot = sum(r["world_mn"] for r in results if r["industry"] == ind)
        if not pubs or pubs[0] / 3 <= tot <= pubs[-1] * 3:
            continue
        gap = {"industry_total_mn": tot, "published_min_mn": pubs[0], "published_max_mn": pubs[-1],
               "ratio_to_midpoint": tot / math.sqrt(pubs[0] * pubs[-1])}
        msg = (f"{ind}: bottom-up ${tot:,.0f}M is more than 3x away from every published figure "
               f"(${pubs[0]:,.0f}-{pubs[-1]:,.0f}M)")
        if reexamine_published:
            tot2 = _reexamine(llm, results, ind, pubs, msg, market, nesting, dim, examples, table, geos, n_samples,
                              salt)
            msg += f" -- re-examined -> ${tot2:,.0f}M"
        else:
            msg += " -- flagged only, number unchanged"
        for r in results:
            if r["industry"] == ind:
                r["flags"].append(msg)
                r["published_gap"] = gap
        flags.append(msg)
        say(f"[use case] {msg}")

    out = {"dimension": dim, "use_cases": results, "flags": flags, "base_year": BASE_YEAR,
           "settings": {"n_samples": n_samples, "salt": salt, "blocklist": list(blocklist),
                        "reexamine_published": reexamine_published, "bounds": BOUNDS},
           "drivers_used": {r["spec"]["driver"]: table.meta(r["spec"]["driver"]) for r in results}}
    return rollup(out, spec)


def _reexamine(llm: Any, results: list[dict[str, Any]], ind: str, pubs: list[float], msg: str, market: str,
               nesting: dict[str, Any], dim: str, examples: dict[str, Any], table: "DriverTable", geos: list[str],
               n_samples: int, salt: str) -> float:
    """Opt-in: redesign each use case of the industry once; keep a redesign only
    if it moves the industry toward the publishers' midpoint."""
    from .scope import segment_definition

    mid = math.sqrt(pubs[0] * pubs[-1])
    fb = (msg + ". That does not make them right, but re-examine your machine counts, product life, adoption and "
          "units per machine for unit errors or scope mistakes.")
    for i, r in enumerate(results):
        if r["industry"] != ind:
            continue
        try:
            s2 = design(llm, market, r, segment_definition(nesting, dim, ind), r["evidence"], r["published_context"],
                        examples, fb, table.fao_items(), sample=1, n_samples=1, salt=f"{salt}-reexamine")
            ev2 = _evaluate(s2, table, geos)
            ind_old = sum(x["world_mn"] for x in results if x["industry"] == ind)
            ind_new = ind_old - r["world_mn"] + ev2["world_units"] * ev2["price"] / 1e6
            if abs(math.log(max(ind_new, 1e-9) / mid)) < abs(math.log(max(ind_old, 1e-9) / mid)):
                results[i] = {**_item(r, s2, table, geos, ev2, evidence=r["evidence"], published=r["published_context"],
                                      samples=[], info={"spread": r["spread"], "used_samples": [],
                                                        "excluded_samples": [], "driver_votes": []},
                                      flags=r["flags"], blocked=r["blocked"]),
                              "samples": r["samples"], "attempts": r["attempts"] + 1, "reexamined": fb}
        except Exception:  # noqa: BLE001
            continue
    return sum(r["world_mn"] for r in results if r["industry"] == ind)


def _item(uc: dict[str, Any], spec_uc: dict[str, Any], table: DriverTable, geos: list[str], ev: dict[str, Any], *,
          evidence: list[dict[str, str]], published: list[dict[str, Any]], samples: list[dict[str, Any]],
          info: dict[str, Any], flags: list[str], blocked: dict[str, int]) -> dict[str, Any]:
    """World quantity (code's arithmetic on the inputs) spread over countries by
    the driver's share only: units(c) = world units x driver(c) / sum(driver)."""
    dvals = {g: max(0.0, table.by_geo(spec_uc["driver"]).get(g, 0.0)) for g in geos}
    total, n, price = ev["driver_total"], ev["world_units"], ev["price"]
    share = {g: dvals[g] / total for g in geos}
    units = {g: n * share[g] for g in geos}
    values = {g: units[g] * price / 1e6 for g in geos}
    return {"industry": uc["industry"], "use_case": uc["use_case"], "own": uc.get("own"), "evidence": evidence,
            "published_context": published, "spec": spec_uc, "archetype": archetype(spec_uc),
            "world_inputs": world_inputs(spec_uc), "world_formula": world_formula(spec_uc),
            "world_units": n, "world_units_estimate": n, "world_units_circular": ev["circular"],
            "driver_values": dvals, "driver_total": total, "driver_share": share, "driver_meta": table.meta(spec_uc["driver"]),
            "units": units, "price_per_quantity_unit": price, "values_mn": values, "world_mn": sum(values.values()),
            "attempts": sum(s["attempts"] for s in samples), "bound_violations": ev["violations"],
            "basis_counts": basis_counts(spec_uc, evidence), "blocked": blocked, "flags": flags,
            "spread": info["spread"], "median_of_samples": info["used_samples"],
            "excluded_samples": info["excluded_samples"], "driver_votes": info["driver_votes"],
            "samples": [{"sample": s["sample"], "attempts": s["attempts"], "archetype": archetype(s["spec"]),
                         "driver": s["spec"].get("driver"), "machine": s["spec"].get("machine"),
                         "quantity_unit": s["spec"].get("quantity_unit"), "world": s["spec"].get("world"),
                         "world_units": s["world_units"], "price": s["price"],
                         "price_spec": s["spec"].get("price_usd_per_unit"), "violations": s["violations"],
                         "circular": s["circular"], "kept_out_of_bounds": s["kept_out_of_bounds"],
                         "reasoning": s["spec"].get("reasoning"), "feedback": s["feedback"]} for s in samples]}


def rollup(res: dict[str, Any], spec: dict[str, Any]) -> dict[str, Any]:
    """use case -> industry -> country -> region -> world (recomputed from use-case values)."""
    results = res["use_cases"]
    geos = spec["geographies"]
    industries = sorted({r["industry"] for r in results})
    country_value = {g: math.fsum(r["values_mn"].get(g, 0.0) for r in results) for g in geos}
    industry_value = {g: {ind: math.fsum(r["values_mn"].get(g, 0.0) for r in results if r["industry"] == ind)
                          for ind in industries} for g in geos}
    regions = spec["regions"]
    region_value = {rg: math.fsum(country_value[c] for c in cs if c in country_value) for rg, cs in regions.items()}
    total = math.fsum(region_value.values())
    return {**res, "industries": industries, "country_value": country_value, "industry_value": industry_value,
            "region_value": region_value, "global_value": total,
            "region_share": {rg: v / total for rg, v in region_value.items()} if total else {},
            "country_share_in_region": {c: country_value[c] / region_value[rg] for rg, cs in regions.items()
                                        for c in cs if region_value.get(rg)}}


# ---- official anchor -----------------------------------------------------------------
def official_by_geo(market: str, spec: dict[str, Any], cache_dir: Path, llm: Any | None = None,
                    fao_cache: dict[str, Any] | None = None) -> dict[str, Any] | None:
    """Official consumption of the market's own product per geography (tonnes),
    when the market IS one FAOSTAT item; else None."""
    from me_engine.demandCountry_approach2A.topdown import official_data as od

    cache = _load_fao_cache() if fao_cache is None else fao_cache
    if not cache:
        return None
    term = market.replace("Global", "").replace("Market", "").strip()
    hit = od.match_item(term, cache, llm)
    if not hit:
        return None
    dataset, item = hit
    wb = load_drivers(cache_dir)
    vals, meta = fao_by_geo(dataset, item, wb, spec["geographies"], cache)
    return {"dataset": dataset, "item": item, "by_geo": vals, **meta}


def anchor_to_official(res: dict[str, Any], official: dict[str, Any], spec: dict[str, Any]) -> dict[str, Any]:
    """Scale each country's use-case quantities so they add up to its official
    consumption of the product. Official statistics set the level; the use cases
    only decide the split. A country without official data keeps the bottom-up
    level and is flagged."""
    from me_engine.demandCountry_approach2A.engine.units import unit_scale

    flags = list(res.get("flags") or [])
    to_t: dict[int, float] = {}
    for i, r in enumerate(res["use_cases"]):
        s = r["spec"]
        f = unit_scale(str(s.get("quantity_unit") or ""), "tonnes", s.get("density_kg_per_l"))
        if f is None:
            flags.append(f"official anchor skipped: '{r['use_case']}' is not in mass/volume units "
                         f"({s.get('quantity_unit')})")
            return {**res, "flags": flags}
        to_t[i] = f
    by_geo: dict[str, dict[str, Any]] = {}
    for g in spec["geographies"]:
        bu = math.fsum(r["units"].get(g, 0.0) * to_t[i] for i, r in enumerate(res["use_cases"]))
        off = float(official["by_geo"].get(g) or 0.0)
        if bu <= 0 or off <= 0:
            by_geo[g] = {"official_t": off, "bottom_up_t": bu, "factor": None}
            flags.append(f"{g}: no official consumption figure -- bottom-up level kept" if off <= 0 else
                         f"{g}: official consumption {off:,.0f} t but no use case reaches it -- not distributed")
            continue
        f = off / bu
        for r in res["use_cases"]:
            r["units"][g] *= f
            r["values_mn"][g] *= f
            r.setdefault("anchor", {})[g] = f
        by_geo[g] = {"official_t": off, "bottom_up_t": bu, "factor": f}
    for r in res["use_cases"]:
        r["world_mn"] = math.fsum(r["values_mn"].values())
    anchored = {**res, "flags": flags,
                "official_anchor": {k: v for k, v in official.items() if k != "by_geo"} | {"by_geo": by_geo}}
    return rollup(anchored, spec)


# ---- ledger -----------------------------------------------------------------------------
def _label(entry: Any, evidence: list[dict[str, str]]) -> tuple[str, dict[str, Any] | None, str]:
    """(label, source, note) for an agent-supplied number: a web snippet it cites
    is tier 5 -> source_found_unverified; anything else is an assumption."""
    e = entry if isinstance(entry, dict) else {}
    reason = str(e.get("reason") or "").strip()
    ids = re.findall(r"E\d+", str(e.get("basis") or ""))
    by_id = {x["id"]: x for x in evidence}
    cited = [by_id[i] for i in ids if i in by_id]
    if cited:
        return ("source_found_unverified",
                {"tier": 5, "publisher": "web search", "evidence_ids": [x["id"] for x in cited],
                 "quote": " | ".join(str(x["text"])[:300] for x in cited)}, reason)
    return "assumption", None, reason or "agent assumption (no reason given)"


def build_ledger(res: dict[str, Any], spec: dict[str, Any], run_id: str = "R000") -> Any:
    """Every number of the build in the 2A ledger: drivers, constants, prices,
    conversions, world inputs, driver shares, official anchor, quantities, values and roll-ups.
    Derived numbers are recomputed here and checked against the build (`mismatches`)."""
    from me_engine.demandCountry_approach2A.ledger import Ledger

    L = Ledger(re.sub(r"[^A-Za-z0-9]", "", run_id) or "R000", "UC")
    year = res.get("base_year", BASE_YEAR)
    geos = spec["geographies"]
    mod = "usecase_builder"
    anchor_geo = (res.get("official_anchor") or {}).get("by_geo") or {}
    official = {}
    for g, a in anchor_geo.items():
        if a.get("factor"):
            meta = res["official_anchor"]
            official[g] = L.fact(f"Official consumption, {g}", a["official_t"], "t", g, "QTY", "derived",
                                 year=year, weakest_input_label="verified_observed", module=mod,
                                 source={"publisher": meta.get("source"), "url": meta.get("url"), "tier": 2},
                                 note=f"FAOSTAT, latest data year {meta.get('latest_year')}, moved to {year} with each "
                                      "country's own trend; summed over the geography's member countries")
    value_rows: dict[str, list[Any]] = {g: [] for g in geos}
    industry_rows: dict[tuple[str, str], list[Any]] = {}
    qty_t: dict[str, list[tuple[Any, float]]] = {g: [] for g in geos}
    for r in res["use_cases"]:
        s, ev = r["spec"], r.get("evidence") or []
        tag = f"{r['industry']} / {r['use_case']}"
        qu = str(s.get("quantity_unit") or s.get("units_label") or "units")
        p = s.get("price_usd_per_unit") or {}
        lab, src, note = _label(p, ev)
        per = str(p.get("per") or qu) if isinstance(p, dict) else qu
        quoted = L.fact(f"{tag}: price", float(p.get("value") if isinstance(p, dict) else p), f"US$ per {per}",
                        "Global", "PRC", lab, year=year, source=src, note=note, module=mod)
        conv = price_conversion(s)
        price = quoted if conv == 1.0 else L.derive(
            f"{tag}: price per {qu}", quoted.value * conv, f"US$ per {qu}", "Global", "PRC",
            f"{quoted.number_id} × {conv:g}", [quoted], year=year, note=f"unit conversion by code ({per} -> {qu})")
        meta = r.get("driver_meta") or {}
        drivers = {}
        for g in geos:
            imp = (meta.get("imputed_share") or {}).get(g, 0.0)
            dlabel = meta.get("label") or ("estimated_from_facts" if imp > 0 else "verified_observed")
            dnote = (f"{imp:.0%} of this geography's value imputed from GDP x its region's median ratio"
                     if imp > 0 else "summed over the geography's member countries")
            drivers[g] = L.fact(f"{tag}: driver {s['driver']}, {g}", r["driver_values"].get(g, 0.0),
                                str(meta.get("unit") or ""), g, "QTY", dlabel, year=year, module=mod,
                                weakest_input_label="verified_observed" if dlabel == "derived" else None,
                                source={"publisher": meta.get("source"), "url": meta.get("url")}, note=dnote)
        dtotal = L.total(f"{tag}: driver {s['driver']}, world total", list(drivers.values()),
                         str(meta.get("unit") or ""), "Global", "QTY", year=year)
        # world inputs (median over the agent's samples) -> world quantity, by code
        spread = r.get("spread") or {}
        inputs = {}
        for k in input_names(s):
            if k == "world_activity_per_year" and activity_from_driver(s):
                inputs[k] = dtotal
                continue
            c = s["world"][k]
            lab, src, note = _label(c, ev)
            sp = spread.get(k)
            if sp:
                note = (f"{note}; median of {len(sp['values'])} samples "
                        f"({', '.join(f'{v:g}' for v in sp['values'])})").lstrip("; ")
            inputs[k] = L.fact(f"{tag}: {k}", float(_entry_value(c)), str(c.get("unit") or "") if isinstance(c, dict)
                               else "", "Global", "PAR", lab, year=year, source=src, note=note, module=mod)
        i = {k: n.number_id for k, n in inputs.items()}
        formula = (f"{i['new_machines_per_year']} × {i['adoption_share']} × {i['units_per_machine']}"
                   if archetype(s) == "durable" and not counts_replacements(s) else
                   f"({i['new_machines_per_year']} + {i['installed_machines']} / {i['product_life_years']}) × "
                   f"{i['adoption_share']} × {i['units_per_machine']}" if archetype(s) == "durable" else
                   f"{i['world_activity_per_year']} × {i['adoption_share']} × {i['dose_per_activity_unit']}")
        world = L.derive(f"{tag}: world quantity per year", world_units(s, dtotal.value), qu, "Global", "QTY", formula,
                         list(inputs.values()), year=year, module=mod, note=world_formula(s))
        L.check(world, r["world_units"], f"{tag} world quantity", rel_tol=1e-6)
        for g in geos:
            share = L.derive(f"{tag}: driver share, {g}", drivers[g].value / dtotal.value, "share", g, "SHR",
                             f"{drivers[g].number_id} / {dtotal.number_id}", [drivers[g], dtotal], year=year,
                             module=mod, note="the driver only spreads the world quantity")
            q = L.mul(f"{tag}: quantity, {g}", world, share, qu, g, "QTY", year=year)
            if g in official and r.get("anchor", {}).get(g):
                qty_t[g].append((q, _to_tonnes(s)))
            r.setdefault("_q", {})[g] = q
        r["_price"] = price
    anchor_factor = {}
    for g, off in official.items():
        parts = qty_t[g]
        bu = L.derive(f"Bottom-up quantity before anchor, {g}", math.fsum(q.value * t for q, t in parts), "t", g, "QTY",
                      " + ".join(f"{q.number_id} × {t:g}" for q, t in parts), [q for q, _ in parts], year=year)
        anchor_factor[g] = L.derive(f"Official anchor factor, {g}", off.value / bu.value, "ratio", g, "FAC",
                                    f"{off.number_id} / {bu.number_id}", [off, bu], year=year,
                                    note="official consumption sets the level; use cases set the split")
    for r in res["use_cases"]:
        tag = f"{r['industry']} / {r['use_case']}"
        qu = str(r["spec"].get("quantity_unit") or r["spec"].get("units_label") or "units")
        for g in geos:
            q = r["_q"][g]
            if g in anchor_factor and r.get("anchor", {}).get(g):
                q = L.mul(f"{tag}: quantity anchored, {g}", q, anchor_factor[g], qu, g, "QTY", year=year)
            v = L.mul(f"{tag}: value, {g}", q, r["_price"], "US$", g, "VAL", year=year)
            L.check(v, r["values_mn"].get(g, 0.0) * 1e6, f"{tag} value {g}", rel_tol=1e-6)
            value_rows[g].append(v)
            industry_rows.setdefault((r["industry"], g), []).append(v)
        r.pop("_q", None)
        r.pop("_price", None)
    country = {}
    for g in geos:
        for ind in res.get("industries") or []:
            parts = industry_rows.get((ind, g))
            if parts:
                n = L.total(f"{ind}, {g}", parts, "US$", g, year=year)
                L.check(n, res["industry_value"][g][ind] * 1e6, f"industry {ind} {g}", rel_tol=1e-6)
        if value_rows[g]:
            country[g] = L.total(f"Market value, {g}", value_rows[g], "US$", g, year=year)
            L.check(country[g], res["country_value"][g] * 1e6, f"country {g}", rel_tol=1e-6)
    regions = []
    for rg, cs in spec["regions"].items():
        parts = [country[c] for c in cs if c in country]
        if parts:
            regions.append(L.total(f"Market value, {rg}", parts, "US$", rg, year=year))
            L.check(regions[-1], res["region_value"][rg] * 1e6, f"region {rg}", rel_tol=1e-6)
    if regions:
        world = L.total("Market value, world", regions, "US$", "Global", year=year)
        L.check(world, res["global_value"] * 1e6, "world", rel_tol=1e-6)
    return L


def _to_tonnes(spec: dict[str, Any]) -> float:
    from me_engine.demandCountry_approach2A.engine.units import unit_scale
    return unit_scale(str(spec.get("quantity_unit") or ""), "tonnes", spec.get("density_kg_per_l")) or 0.0


def write(spec: dict[str, Any], data_file: Path, ub: dict[str, Any]) -> dict[str, int]:
    """Write the bottom-up result into the human-structured Data file: country
    and region values and shares, the industry mix per country and the use-case
    mix within each industry per country."""
    from openpyxl import load_workbook

    from . import evaluator_agent as ev
    from .structure_agent import _norm

    reg_like = {"region_value": ub["region_value"], "region_share": ub["region_share"],
                "country_value": ub["country_value"], "country_share_in_region": ub["country_share_in_region"],
                "global_value": ub["global_value"], "dimension": ub["dimension"],
                "country_mix": {g: {ind: (v / ub["country_value"][g] if ub["country_value"][g] else 0.0)
                                    for ind, v in mix.items()} for g, mix in ub["industry_value"].items()}}
    stats = ar.write(spec, data_file, reg_like)
    wb = load_workbook(data_file)
    ws = wb["Data"]
    labels = [_norm(ws.cell(r, 2).value) for r in range(1, ws.max_row + 1)]
    paths = ev._path_walk(labels, spec)
    geo_col = {_norm(ws.cell(3, c).value): c for c in range(3, 30) if _norm(ws.cell(3, c).value)}
    by_uc = {(r["industry"], r["use_case"]): r for r in ub["use_cases"]}
    n = 0
    for row, p in enumerate(paths, start=1):
        if p and p[0] == ub["dimension"] and len(p[1]) == 2:
            ind, uc = p[1]
            r = by_uc.get((ind, uc))
            if not r:
                continue
            for g, c in geo_col.items():
                tot = ub["industry_value"].get(g, {}).get(ind, 0.0)
                if tot > 0:
                    ws.cell(row, c).value = r["values_mn"].get(g, 0.0) / tot
                    n += 1
    stats["asp_cells"] = _write_asp(wb, spec, ub, _norm)
    wb.save(data_file)
    stats["use_case_cells"] = n
    return stats


def _write_asp(wb: Any, spec: dict[str, Any], ub: dict[str, Any], norm: Callable[[Any], str]) -> int:
    """The bottom-up price of each use case into the ASP sheet's first block (one
    price per use case, every country), so the ME's volume = value / price matches
    the units the build counted. Only use cases counted in units of the product."""
    if "ASP" not in wb.sheetnames or not spec.get("geographies"):
        return 0
    wsa = wb["ASP"]
    first = spec["geographies"][0]
    header = next((r for r in range(1, wsa.max_row + 1) if norm(wsa.cell(r, 3).value) == first), None)
    if header is None:
        return 0
    acol = {norm(wsa.cell(header, c).value): c for c in range(3, 30) if norm(wsa.cell(header, c).value)}
    price = {norm(r["use_case"]).lower(): r["price_per_quantity_unit"] for r in ub["use_cases"]
             if _unit_words(r["spec"].get("quantity_unit") or "units") <= _COUNT_WORDS | {"unit"}}
    n, r = 0, header + 1
    while r <= wsa.max_row and (norm(wsa.cell(r, 2).value) or r == header + 1):
        p = price.get(norm(wsa.cell(r, 2).value).lower())
        if p:
            for g in spec["geographies"]:
                if g in acol:
                    wsa.cell(r, acol[g]).value = p
                    n += 1
        r += 1
    return n
