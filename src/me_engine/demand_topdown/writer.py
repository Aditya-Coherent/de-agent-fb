"""Final market report writer: JSON + a presentation-ready XLSX.

Sheets:
    Summary            global -> region -> country totals
    <one per dimension> countries grouped by region, region subtotals, Global row
    Method & Evidence  how each dimension was built, share evidence status, checks
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_REGION_FILL = PatternFill("solid", fgColor="D9E1F2")
_GLOBAL_FILL = PatternFill("solid", fgColor="FCE4D6")
_BOLD = Font(bold=True)
_NUM = "#,##0.0"


def write_report_json(report: dict[str, Any], path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    return path


def _sheet_title(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]\*\?/\\:]", "", name)[:28] or "Dimension"
    title, i = base, 2
    while title.lower() in used:
        title = f"{base[:25]} {i}"
        i += 1
    used.add(title.lower())
    return title


def _header(ws, row: int, values: list[str]) -> None:
    for col, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=v)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize(ws, first_width: int = 30, width: int = 16) -> None:
    ws.column_dimensions["A"].width = first_width
    for col in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _write_summary(ws, report: dict[str, Any]) -> None:
    geo = report["geography"]
    ws.append([f"{report['market_name']} -- Approach {report.get('approach')}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([report.get("methodology", "")])
    ws.append([f"Base year: {report.get('base_year')}   Unit: {report.get('unit')}"])
    ws.append([])
    _header(ws, 5, ["Geography", "Level", f"Market value ({report.get('unit')})", "Share of global"])
    total = geo["global"] or 0.0
    ws.append(["Global", "Global", geo["global"], 1.0 if total else 0.0])
    for c in ws[ws.max_row]:
        c.fill = _GLOBAL_FILL
        c.font = _BOLD
    for region, rv in geo["regions"].items():
        ws.append([region, "Region", rv, rv / total if total else 0.0])
        for c in ws[ws.max_row]:
            c.fill = _REGION_FILL
            c.font = _BOLD
        for country, cv in geo["countries"].items():
            if geo["country_region"].get(country) == region:
                ws.append([f"   {country}", "Country", cv, cv / total if total else 0.0])
    for row in ws.iter_rows(min_row=6, min_col=3, max_col=4):
        row[0].number_format = _NUM
        row[1].number_format = "0.0%"
    _autosize(ws, 34, 22)


def _write_dimension(ws, table: dict[str, Any], geo: dict[str, Any]) -> None:
    segments = table["segments"]
    method = "BOTTOM-UP (built per country)" if table["method"] == "bottom_up" else "TOP-DOWN (country total x share)"
    ws.append([f"{table['dimension']} -- {method}"])
    ws["A1"].font = Font(bold=True, size=13)
    if table.get("evidence_status"):
        ws.append([f"Evidence status: {table['evidence_status']}"])
    else:
        ws.append([])
    ws.append([])
    header_row = ws.max_row + 1
    _header(ws, header_row, ["Geography", *segments, "Total"])

    def add(label: str, row: dict[str, float], fill=None) -> None:
        vals = [row.get(s, 0.0) for s in segments]
        ws.append([label, *vals, round(sum(vals), 4)])
        r = ws.max_row
        for col in range(2, len(segments) + 3):
            ws.cell(row=r, column=col).number_format = _NUM
        if fill is not None:
            for c in ws[r]:
                c.fill = fill
                c.font = _BOLD

    add("Global", table["global_values"], _GLOBAL_FILL)
    for region, rv in table["region_values"].items():
        add(region, rv, _REGION_FILL)
        for country, cv in table["country_values"].items():
            if geo["country_region"].get(country) == region:
                add(f"   {country}", cv)

    if table.get("shares"):
        ws.append([])
        ws.append(["Shares used (fraction of each country's total)"])
        ws.cell(row=ws.max_row, column=1).font = _BOLD
        _header(ws, ws.max_row + 1, ["Country", *segments])
        for country, sh in table["shares"].items():
            ws.append([country, *[sh.get(s, 0.0) for s in segments]])
            for col in range(2, len(segments) + 2):
                ws.cell(row=ws.max_row, column=col).number_format = "0.0%"
    _autosize(ws)


def _write_method(ws, report: dict[str, Any]) -> None:
    ws.append(["Method & Evidence"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([report.get("methodology", "")])
    ws.append([])
    _header(ws, 4, ["Dimension", "Method", "Taxonomy source", "Share evidence", "Verified quotes", "Reasoning / flags"])
    for t in report["segmentation"]:
        quotes = [e["quote"] for e in t.get("evidence", []) if e.get("verified_in_snippets")]
        notes = "; ".join([t.get("reasoning", ""), *t.get("flags", [])]).strip("; ")
        ws.append([
            t["dimension"], t["method"], t.get("taxonomy_source", "usage taxonomy (approved)"),
            t.get("evidence_status", ""), " | ".join(quotes), notes,
        ])
        for c in ws[ws.max_row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.append([])
    ws.append(["Consistency checks (every dimension must add back to the same country and global totals)"])
    ws.cell(row=ws.max_row, column=1).font = _BOLD
    _header(ws, ws.max_row + 1, ["Dimension", "Max country gap (USD Mn)", "Global gap (USD Mn)", "Passed"])
    for chk in report.get("consistency_checks", []):
        ws.append([chk["dimension"], chk["max_country_gap_usd_mn"], chk["global_gap_usd_mn"], "YES" if chk["passed"] else "NO"])
    _autosize(ws, 28, 24)
    ws.column_dimensions["F"].width = 70
    ws.column_dimensions["E"].width = 50


def format_terminal_summary(report: dict[str, Any]) -> str:
    geo = report["geography"]
    lines = [f"=== {report['market_name']} -- Approach {report.get('approach')} (USD Mn, {report.get('base_year')}) ==="]
    lines.append(f"Global: ${geo['global']:,.1f}M")
    for region, rv in geo["regions"].items():
        lines.append(f"  {region}: ${rv:,.1f}M")
        for country, cv in geo["countries"].items():
            if geo["country_region"].get(country) == region:
                lines.append(f"      {country}: ${cv:,.1f}M")
    for t in report["segmentation"]:
        tag = "bottom-up" if t["method"] == "bottom_up" else f"top-down, {t.get('evidence_status', '')}"
        lines.append(f"\n{t['dimension']} ({tag}):")
        for seg, v in t["global_values"].items():
            lines.append(f"  {seg}: ${v:,.1f}M")
    failed = [c["dimension"] for c in report.get("consistency_checks", []) if not c["passed"]]
    lines.append("\nConsistency checks: " + ("all passed" if not failed else f"FAILED for {failed}"))
    return "\n".join(lines)


def write_report_workbook(report: dict[str, Any], path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    _write_summary(ws, report)
    used = {"summary", "method & evidence"}
    for table in report["segmentation"]:
        _write_dimension(wb.create_sheet(_sheet_title(table["dimension"], used)), table, report["geography"])
    _write_method(wb.create_sheet("Method & Evidence"), report)
    wb.save(path)
    return path
