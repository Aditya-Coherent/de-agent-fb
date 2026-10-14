"""Bottom-up shares for every non-use-case dimension.

A human analyst does not guess a country's laser-type split; they ask what each
use case buys (Welding buys continuous-wave fiber lasers of 1-20 kW, Counter-UAS
buys beam-combined fiber arrays above 100 kW) and weight by how much of each use
case a country has. Here:

    profile:  for each use case, the agent gives the mix of every other
              dimension's segments it buys (top level sums to 1, sub-segments
              within each parent sum to 1), each with a reason and basis
    aggregate: share(country, dim, segment)
                 = sum_uc value(uc, country) * mix(uc, dim, segment) / sum_uc value(uc, country)
               sub-segment within parent p
                 = sum_uc value * mix(p) * mix(child | p) / sum_uc value * mix(p)

Country differences come only from each country's use-case composition, which in
turn comes from its real-world drivers -- no country number is invented.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any


def _norm(d: dict[str, Any], keys: list[str]) -> dict[str, float]:
    vals = {}
    for k in keys:
        try:
            vals[k] = max(0.0, float(d.get(k) or 0.0))
        except (TypeError, ValueError):
            vals[k] = 0.0
    s = sum(vals.values())
    return {k: v / s for k, v in vals.items()} if s > 0 else {k: 1.0 / len(keys) for k in keys}


def profile(llm: Any, market: str, item: dict[str, Any], trees: dict[str, dict[str, list[str]]]) -> dict[str, Any]:
    """One call per use case: what it buys, in every other dimension."""
    s = item.get("spec") or {}
    ev = "\n".join(f"[{e['id']}] {str(e['text'])[:300]}" for e in (item.get("evidence") or [])[:8]) or "(none)"
    data = llm.complete_json(
        "You are a market analyst. For ONE use case of a product, state what it buys in each segmentation "
        "dimension: the share of this use case's spending that goes to each segment. Think physically: which "
        "technology, power class, operating mode, deployment, component content and buyer type this use case "
        "actually needs. Top-level shares of a dimension sum to 1; sub-segments within each parent sum to 1. "
        "Use 0 where the use case buys none of a segment. Give one short reason per dimension, citing evidence "
        "ids where they support it and saying 'assumption' otherwise. Return json.",
        f"Market: {market}\nIndustry: {item['industry']}\nUse case: {item['use_case']}\n"
        f"Who buys: {s.get('buyers', '')}\nOne unit: {s.get('units_label', '')}; price ${s.get('price_usd_per_unit', {}).get('value', '')}\n"
        f"Evidence:\n{ev}\n\nDimensions (top level, and sub-segments by parent): {json.dumps(trees, ensure_ascii=False)}\n"
        'Return JSON: {"<dimension>": {"top": {"<segment>": <share>}, "sub": {"<parent>": {"<child>": <share>}}, '
        '"reason": "<1-2 sentences>", "basis": "<E-ids or assumption>"}}',
    )
    out = {}
    for d, tree in trees.items():
        got = data.get(d) if isinstance(data.get(d), dict) else {}
        top = _norm(got.get("top") or {}, tree["__root__"])
        sub = {p: _norm(((got.get("sub") or {}).get(p) or {}), kids) for p, kids in tree.items() if p != "__root__"}
        out[d] = {"top": top, "sub": sub, "reason": str(got.get("reason") or "")[:600],
                  "basis": str(got.get("basis") or "assumption")[:80], "answered": bool(got.get("top"))}
    return out


def build(llm: Any, ub: dict[str, Any], nesting: dict[str, Any], market: str,
          say: Any = print) -> dict[str, Any]:
    trees = {d: t for d, t in nesting.items() if d != ub["dimension"] and (t or {}).get("__root__")}
    items = ub["use_cases"]
    profiles = {}
    for it in items:
        try:
            profiles[f"{it['industry']}|{it['use_case']}"] = profile(llm, market, it, trees)
        except Exception as exc:  # noqa: BLE001
            say(f"[mix] {it['use_case']}: profile failed ({exc})")
    geos = list(ub["country_value"])
    shares: dict[str, dict[str, dict[str, float]]] = {}  # dim -> geo -> "seg" or "parent|child" -> share
    contrib: dict[str, dict[str, dict[str, dict[str, float]]]] = {}  # dim -> geo -> key -> uc -> value_mn
    for d, tree in trees.items():
        shares[d], contrib[d] = {}, {}
        for g in geos:
            c: dict[str, dict[str, float]] = {}
            for it in items:
                key = f"{it['industry']}|{it['use_case']}"
                pr = profiles.get(key)
                v = it["values_mn"].get(g, 0.0)
                if not pr or v <= 0:
                    continue
                for seg, m in pr[d]["top"].items():
                    c.setdefault(seg, {})[key] = v * m
                    for child, mc in (pr[d]["sub"].get(seg) or {}).items():
                        c.setdefault(f"{seg}|{child}", {})[key] = v * m * mc
            tot = sum(sum(c.get(s, {}).values()) for s in tree["__root__"])
            sh = {s: (sum(c.get(s, {}).values()) / tot if tot else 0.0) for s in tree["__root__"]}
            for p, kids in tree.items():
                if p == "__root__":
                    continue
                pt = sum(c.get(p, {}).values())
                for k in kids:
                    sh[f"{p}|{k}"] = sum(c.get(f"{p}|{k}", {}).values()) / pt if pt else 1.0 / len(kids)
            shares[d][g], contrib[d][g] = sh, c
    say(f"[mix] {len(profiles)} use-case profiles over {len(trees)} dimensions")
    return {"profiles": profiles, "shares": shares, "contrib": contrib, "dimensions": list(trees)}


def write(spec: dict[str, Any], data_file: Path, mix: dict[str, Any]) -> int:
    from openpyxl import load_workbook

    from . import evaluator_agent as ev
    from .structure_agent import _norm as norm

    wb = load_workbook(data_file)
    ws = wb["Data"]
    labels = [norm(ws.cell(r, 2).value) for r in range(1, ws.max_row + 1)]
    paths = ev._path_walk(labels, spec)
    geo_col = {norm(ws.cell(3, c).value): c for c in range(3, 30) if norm(ws.cell(3, c).value)}
    n = 0
    for r, p in enumerate(paths, start=1):
        if not p or p[0] not in mix["shares"]:
            continue
        key = "|".join(p[1])
        for g, c in geo_col.items():
            v = mix["shares"][p[0]].get(g, {}).get(key)
            if v is not None:
                ws.cell(r, c).value = v
                n += 1
    wb.save(data_file)
    return n
