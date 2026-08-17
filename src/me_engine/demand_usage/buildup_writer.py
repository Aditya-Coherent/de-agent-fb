"""Write the transparent build-up artifact: Geography -> Usage -> Driver ->
Quantity -> Price -> Value, line by line. This is the artifact Desk Research
fundamentally cannot produce, since it never breaks the total down this way.

Two output forms of the same data:
    - JSON: full machine-readable state (for programmatic reuse/comparison)
    - .xlsx: a human-readable workbook (Build-Up, Summary, Validation sheets)
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Mapping

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

_PROV_FILL = {
    "observed": PatternFill("solid", fgColor="C6EFCE"),
    "derived": PatternFill("solid", fgColor="FFEB9C"),
    "assumption": PatternFill("solid", fgColor="FFC7CE"),
}


def build_report_dict(state: Mapping[str, Any]) -> dict[str, Any]:
    """Flatten the graph state into the auditable build-up chain plus metadata."""
    buckets = state.get("usage_buckets", []) or []
    countries = state.get("countries", []) or []
    drivers = state.get("drivers", {}) or {}
    quantities = state.get("quantities", {}) or {}
    prices = state.get("prices", {}) or {}
    usage_values = state.get("usage_values", {}) or {}
    driver_facts = state.get("driver_facts", {}) or {}
    price_facts = state.get("price_facts", {}) or {}
    unit_contract = state.get("unit_contract", {}) or {}

    rows: list[dict[str, Any]] = []
    for b in buckets:
        name = b["name"]
        driver = drivers.get(name, {})
        price = prices.get(name, {})
        price_fact = price_facts.get(name, {})
        contract = unit_contract.get(name, {})
        for c in countries:
            qty = quantities.get(name, {}).get(c, {})
            facts = driver_facts.get(name, {}).get(c, {})
            driver_raw = facts.get("driver_raw", {})
            intensity = facts.get("intensity", {})
            rows.append({
                "country": c,
                "usage_bucket": name,
                "bucket_category": b.get("bucket_category", "physical_consumption"),
                "usage_description": b.get("description", ""),
                "driver_name": driver.get("driver_name", ""),
                "driver_unit": driver.get("unit", ""),
                "final_demand_unit": driver.get("final_demand_unit", ""),
                "conversion_formula": facts.get("formula") or driver.get("conversion_formula", ""),
                "library_status": facts.get("library_status") or driver.get("library_status", ""),
                "quantity": qty.get("value"),
                "provenance": qty.get("provenance"),
                "provenance_note": qty.get("note", ""),
                "driver_evidence_type": driver_raw.get("evidence_type", ""),
                "driver_source": driver_raw.get("source", ""),
                "driver_citation": driver_raw.get("citation", ""),
                "driver_year": driver_raw.get("year", ""),
                "intensity_evidence_type": intensity.get("evidence_type", ""),
                "asp_usd_per_unit": price.get("asp_usd_per_unit"),
                "asp_evidence_type": price.get("evidence_type", ""),
                "asp_source": price_fact.get("source", ""),
                "asp_citation": price_fact.get("citation", ""),
                "value_usd_mn": usage_values.get(name, {}).get(c),
                "llm_call_failed": bool(qty.get("call_failed")) or bool(price.get("call_failed")),
                "unit_contract_status": contract.get("status", "unknown"),
                "normalized_asp_per_base_unit": contract.get("normalized_asp_per_base_unit"),
            })

    return {
        "market_name": state.get("market_name"),
        "methodology": "demand_side_usage",
        # Approved Phase 1 Market Setup -- carried through unchanged, never
        # re-derived by Phase 3 (spec section 18: "Market Setup" section).
        "market_setup": {
            "scope": state.get("scope"),
            "geography": countries,
            "base_year": state.get("base_year"),
            "market_unit": state.get("market_unit"),
            "physical_unit": state.get("physical_unit"),
        },
        "countries": countries,
        "usage_buckets": buckets,
        # Non-physical (financial/inventory) buckets excluded before pricing --
        # kept here for audit, never silently dropped from the record.
        "excluded_usage_buckets": state.get("excluded_usage_buckets", []),
        "usage_classifications": state.get("usage_classifications", {}),
        # Candidate major physical applications the bucket set may be missing --
        # flag only, evaluated before calculation, never merged into usage_buckets.
        "missing_usage_flags": state.get("missing_usage_flags", []),
        # Deterministic unit-contract verdict per bucket (units.py, no LLM) --
        # any bucket with status != "ok" was excluded from calculation entirely.
        "unit_contract": unit_contract,
        "unit_rejected_buckets": state.get("unit_rejected_buckets", []),
        "drivers": drivers,
        "build_up_rows": rows,
        "country_totals_usd_mn": state.get("country_totals"),
        "total_market_value_usd_mn": state.get("total_market_value"),  # Included Demand (physical consumption only)
        "related_demand_country_totals_usd_mn": state.get("related_demand_country_totals"),
        "related_demand_value_usd_mn": state.get("related_demand_value"),  # inventory/reserve, kept separate
        "validation": state.get("validation"),
        "evidence_quality": (state.get("validation") or {}).get("evidence_quality"),
        # Flag-only driver-magnitude cross-check + informational external
        # benchmark sweep (spec sections 11/16) -- never used to adjust any
        # quantity or the final total; see magnitude.py / triangulation.py.
        "driver_magnitude_flags": state.get("driver_magnitude_flags", []),
        "triangulation": state.get("triangulation", []),
        "triangulation_large_gap_flags": state.get("triangulation_large_gap_flags", []),
        "quality_flags": state.get("quality_flags"),
        "log": state.get("log"),
    }


def write_buildup_json(state: Mapping[str, Any], out_path: Path) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(build_report_dict(state), indent=2, ensure_ascii=False), encoding="utf-8")
    return out_path


def write_buildup_workbook(state: Mapping[str, Any], out_path: Path) -> Path:
    report = build_report_dict(state)
    wb = Workbook()
    bold = Font(bold=True)

    setup_ws = wb.active
    setup_ws.title = "Market Setup"
    setup = report.get("market_setup") or {}
    setup_rows = [
        ("Market Name", report.get("market_name")),
        ("Scope", setup.get("scope")),
        ("Geography", ", ".join(setup.get("geography") or [])),
        ("Base Year", setup.get("base_year")),
        ("Market Unit", setup.get("market_unit")),
        ("Physical Unit", setup.get("physical_unit")),
    ]
    for i, (label, value) in enumerate(setup_rows, start=1):
        setup_ws.cell(i, 1, label).font = bold
        setup_ws.cell(i, 2, value)

    ws = wb.create_sheet("Build-Up", 0)
    ws.title = "Build-Up"
    ws.cell(1, 1, "Market Name").font = bold
    ws.cell(1, 2, report["market_name"])
    ws.cell(2, 1, "Methodology").font = bold
    ws.cell(2, 2, "Demand-Side (Usage/Application, bottom-up, evidence-first, human-approved setup+taxonomy V3)")
    headers = [
        "Country", "Usage Bucket", "Bucket Category", "Driver", "Driver Unit", "Final Demand Unit",
        "Conversion Formula", "Library Status", "Quantity", "Provenance", "Provenance Note",
        "Driver Evidence Type", "Driver Source", "Driver Citation", "Driver Year",
        "Intensity Evidence Type", "ASP (US$/unit)", "ASP Evidence Type", "ASP Source",
        "ASP Citation", "Value (US$ Mn)", "LLM Call Failed",
        "Unit Contract Status", "Normalized ASP (per base unit)",
    ]
    for ci, h in enumerate(headers, start=1):
        ws.cell(4, ci, h).font = bold
    r = 5
    for row in report["build_up_rows"]:
        ws.cell(r, 1, row["country"])
        ws.cell(r, 2, row["usage_bucket"])
        cat_cell = ws.cell(r, 3, row.get("bucket_category", "physical_consumption"))
        if row.get("bucket_category") == "inventory_reserve":
            cat_cell.fill = PatternFill("solid", fgColor="BDD7EE")  # distinct blue -- Related, not Included
        ws.cell(r, 4, row["driver_name"])
        ws.cell(r, 5, row["driver_unit"])
        ws.cell(r, 6, row["final_demand_unit"])
        ws.cell(r, 7, row["conversion_formula"])
        ws.cell(r, 8, row["library_status"])
        ws.cell(r, 9, row["quantity"])
        prov_cell = ws.cell(r, 10, row["provenance"])
        if row["provenance"] in _PROV_FILL:
            prov_cell.fill = _PROV_FILL[row["provenance"]]
        ws.cell(r, 11, row["provenance_note"])
        driver_et_cell = ws.cell(r, 12, row["driver_evidence_type"])
        if row["driver_evidence_type"] in _PROV_FILL:
            driver_et_cell.fill = _PROV_FILL[row["driver_evidence_type"]]
        ws.cell(r, 13, row["driver_source"])
        ws.cell(r, 14, row["driver_citation"])
        ws.cell(r, 15, row["driver_year"])
        ws.cell(r, 16, row["intensity_evidence_type"])
        ws.cell(r, 17, row["asp_usd_per_unit"])
        asp_et_cell = ws.cell(r, 18, row["asp_evidence_type"])
        if row["asp_evidence_type"] in _PROV_FILL:
            asp_et_cell.fill = _PROV_FILL[row["asp_evidence_type"]]
        ws.cell(r, 19, row["asp_source"])
        ws.cell(r, 20, row["asp_citation"])
        ws.cell(r, 21, row["value_usd_mn"])
        fail_cell = ws.cell(r, 22, "YES -- NOT a genuine estimate" if row["llm_call_failed"] else "")
        if row["llm_call_failed"]:
            fail_cell.fill = PatternFill("solid", fgColor="FF0000")
        uc_status = row.get("unit_contract_status", "unknown")
        uc_cell = ws.cell(r, 23, uc_status)
        if uc_status != "ok":
            uc_cell.fill = PatternFill("solid", fgColor="FF0000")
        ws.cell(r, 24, row.get("normalized_asp_per_base_unit"))
        r += 1

    excluded_ws = wb.create_sheet("Excluded Usages")
    excluded_ws.cell(1, 1, "Bucket").font = bold
    excluded_ws.cell(1, 2, "Description").font = bold
    excluded_ws.cell(1, 3, "Status").font = bold
    excluded_ws.cell(1, 4, "Classification").font = bold
    excluded_ws.cell(1, 5, "Reason").font = bold
    excluded_ws.cell(1, 6, "Confidence").font = bold
    er = 2
    for exc in report.get("excluded_usage_buckets") or []:
        excluded_ws.cell(er, 1, exc.get("name", ""))
        excluded_ws.cell(er, 2, exc.get("description", ""))
        status_cell = excluded_ws.cell(er, 3, exc.get("status", "excluded_non_physical_usage"))
        status_cell.fill = PatternFill("solid", fgColor="FFC7CE")
        excluded_ws.cell(er, 4, exc.get("classification", ""))
        excluded_ws.cell(er, 5, exc.get("reason", ""))
        excluded_ws.cell(er, 6, exc.get("confidence", ""))
        er += 1
    if er == 2:
        excluded_ws.cell(2, 1, "(none -- every discovered usage bucket was classified as a genuine physical end-use)")

    coverage_ws = wb.create_sheet("Coverage Review")
    coverage_ws.cell(1, 1, "Candidate Missing Application").font = bold
    coverage_ws.cell(1, 2, "Reason").font = bold
    coverage_ws.cell(1, 3, "Evidence Type").font = bold
    coverage_ws.cell(1, 4, "Evidence Snippet").font = bold
    coverage_ws.cell(1, 5, "Confidence").font = bold
    coverage_ws.cell(1, 6, "Max Similarity to Existing Bucket").font = bold
    cr = 2
    for m in report.get("missing_usage_flags") or []:
        coverage_ws.cell(cr, 1, m.get("name", ""))
        coverage_ws.cell(cr, 2, m.get("reason", ""))
        et_cell = coverage_ws.cell(cr, 3, m.get("evidence_type", ""))
        et_cell.fill = PatternFill(
            "solid", fgColor="C6EFCE" if m.get("evidence_type") == "evidence_backed" else "FFEB9C"
        )
        coverage_ws.cell(cr, 4, m.get("evidence_snippet", ""))
        coverage_ws.cell(cr, 5, m.get("confidence", ""))
        coverage_ws.cell(cr, 6, m.get("max_similarity_to_existing", ""))
        cr += 1
    if cr == 2:
        coverage_ws.cell(2, 1, "(none -- no evidence-supported coverage gaps were found; NOT proof of full coverage)")

    magnitude_ws = wb.create_sheet("Driver Magnitude Checks")
    magnitude_ws.cell(1, 1, "Bucket").font = bold
    magnitude_ws.cell(1, 2, "Used Value").font = bold
    magnitude_ws.cell(1, 3, "Used Unit").font = bold
    magnitude_ws.cell(1, 4, "Independent Reference Value").font = bold
    magnitude_ws.cell(1, 5, "Reference Source").font = bold
    magnitude_ws.cell(1, 6, "Ratio").font = bold
    magnitude_ws.cell(1, 7, "Note").font = bold
    mr = 2
    for f in report.get("driver_magnitude_flags") or []:
        magnitude_ws.cell(mr, 1, f.get("bucket", ""))
        magnitude_ws.cell(mr, 2, f.get("used_value", ""))
        magnitude_ws.cell(mr, 3, f.get("used_unit", ""))
        magnitude_ws.cell(mr, 4, f.get("reference_value", ""))
        magnitude_ws.cell(mr, 5, f.get("reference_source", ""))
        ratio_cell = magnitude_ws.cell(mr, 6, f.get("ratio", ""))
        ratio_cell.fill = PatternFill("solid", fgColor="FFC7CE")
        magnitude_ws.cell(mr, 7, f.get("note", ""))
        mr += 1
    if mr == 2:
        magnitude_ws.cell(2, 1, "(none -- no bucket's driver total diverged 3x+ from an independent reference)")

    tri_ws = wb.create_sheet("Triangulation")
    tri_ws.cell(1, 1, "Benchmark Category").font = bold
    tri_ws.cell(1, 2, "Value").font = bold
    tri_ws.cell(1, 3, "Unit").font = bold
    tri_ws.cell(1, 4, "Year").font = bold
    tri_ws.cell(1, 5, "Source").font = bold
    tri_ws.cell(1, 6, "Citation").font = bold
    tri_ws.cell(1, 7, "Evidence Type").font = bold
    tri_ws.cell(1, 8, "Demand-Side Estimate (US$ Mn)").font = bold
    tri_ws.cell(1, 9, "Gap %").font = bold
    tr = 2
    for b in report.get("triangulation") or []:
        tri_ws.cell(tr, 1, b.get("category", ""))
        tri_ws.cell(tr, 2, b.get("value", ""))
        tri_ws.cell(tr, 3, b.get("unit", ""))
        tri_ws.cell(tr, 4, b.get("year", ""))
        tri_ws.cell(tr, 5, b.get("source", ""))
        tri_ws.cell(tr, 6, b.get("citation", ""))
        et_cell = tri_ws.cell(tr, 7, b.get("evidence_type", ""))
        if b.get("evidence_type") in _PROV_FILL:
            et_cell.fill = _PROV_FILL[b["evidence_type"]]
        tri_ws.cell(tr, 8, b.get("demand_side_estimate_usd_mn", ""))
        gap_cell = tri_ws.cell(tr, 9, b.get("gap_pct", ""))
        if isinstance(b.get("gap_pct"), (int, float)) and abs(b["gap_pct"]) >= 40.0:
            gap_cell.fill = PatternFill("solid", fgColor="FFEB9C")
        tr += 1
    if tr == 2:
        tri_ws.cell(2, 1, "(no independent benchmark evidence found -- informational sheet only, "
                          "never used to adjust the bottom-up total)")

    uc_ws = wb.create_sheet("Unit Contract")
    uc_ws.cell(1, 1, "Bucket").font = bold
    uc_ws.cell(1, 2, "Status").font = bold
    uc_ws.cell(1, 3, "Final Demand Unit (raw)").font = bold
    uc_ws.cell(1, 4, "Parsed Scale").font = bold
    uc_ws.cell(1, 5, "Parsed Base Unit").font = bold
    uc_ws.cell(1, 6, "Parsed Period").font = bold
    uc_ws.cell(1, 7, "Normalized ASP (per base unit)").font = bold
    uc_ws.cell(1, 8, "Reason (if rejected)").font = bold
    ucr = 2
    for name, c in (report.get("unit_contract") or {}).items():
        uc_ws.cell(ucr, 1, name)
        status_cell = uc_ws.cell(ucr, 2, c.get("status", "unknown"))
        status_cell.fill = PatternFill("solid", fgColor="C6EFCE" if c.get("status") == "ok" else "FFC7CE")
        uc_ws.cell(ucr, 3, c.get("final_unit_raw", ""))
        uc_ws.cell(ucr, 4, c.get("parsed_scale"))
        uc_ws.cell(ucr, 5, c.get("parsed_base_unit"))
        uc_ws.cell(ucr, 6, c.get("parsed_period"))
        uc_ws.cell(ucr, 7, c.get("normalized_asp_per_base_unit"))
        uc_ws.cell(ucr, 8, c.get("reason", ""))
        ucr += 1
    if ucr == 2:
        uc_ws.cell(2, 1, "(no buckets processed)")

    summary = wb.create_sheet("Summary")
    validation_for_summary = report.get("validation") or {}
    summary.cell(1, 1, "Reported Market View Breakdown").font = bold
    summary.cell(2, 1, "Included Demand (physical consumption)")
    summary.cell(2, 2, report["total_market_value_usd_mn"])
    summary.cell(3, 1, "+ Residual / Uncovered (implied by coverage gap -- NOT evidence-built, NOT in any bucket)")
    summary.cell(3, 2, validation_for_summary.get("residual_implied_usd_mn"))
    summary.cell(4, 1, "+ Related Demand (inventory/reserve, reported separately)")
    summary.cell(4, 2, report.get("related_demand_value_usd_mn"))
    summary.cell(5, 1, "= Reported Market View").font = bold
    summary.cell(5, 2, validation_for_summary.get("reported_market_view_usd_mn")).font = bold
    summary.cell(6, 1, "Estimated Coverage % (bounded pipeline estimate -- NOT verified market knowledge)")
    cov = validation_for_summary.get("coverage_pct")
    summary.cell(6, 2, f"{cov:.1f}%" if isinstance(cov, (int, float)) else "n/a")
    summary.cell(7, 1, "Coverage threshold (configurable)")
    thr = validation_for_summary.get("coverage_threshold_pct")
    summary.cell(7, 2, f"{thr:.0f}%" if isinstance(thr, (int, float)) else "n/a")

    summary.cell(9, 1, "Evidence Quality").font = bold
    eq = report.get("evidence_quality") or {}
    summary.cell(10, 1, "Observed inputs")
    summary.cell(10, 2, f"{eq.get('observed_pct', 0):.1f}%")
    summary.cell(11, 1, "Derived inputs")
    summary.cell(11, 2, f"{eq.get('derived_pct', 0):.1f}%")
    summary.cell(12, 1, "Assumption inputs")
    summary.cell(12, 2, f"{eq.get('assumption_pct', 0):.1f}%")

    summary.cell(14, 1, "Country").font = bold
    summary.cell(14, 2, "Included Demand (US$ Mn)").font = bold
    summary.cell(14, 3, "Related Demand (US$ Mn)").font = bold
    sr = 15
    related_by_country = report.get("related_demand_country_totals_usd_mn") or {}
    for country, total in (report["country_totals_usd_mn"] or {}).items():
        summary.cell(sr, 1, country)
        summary.cell(sr, 2, total)
        summary.cell(sr, 3, related_by_country.get(country, 0.0))
        sr += 1
    summary.cell(sr + 1, 1, "Global Total").font = bold
    summary.cell(sr + 1, 2, report["total_market_value_usd_mn"])
    summary.cell(sr + 1, 3, report.get("related_demand_value_usd_mn"))

    val = wb.create_sheet("Validation")
    validation = report.get("validation") or {}
    val.cell(1, 1, "Metric").font = bold
    val.cell(1, 2, "Value").font = bold
    vr = 2
    for key in (
        "deterministic_score", "plausibility_score", "arithmetic_ok",
        "coverage_pct", "coverage_threshold_pct", "included_demand_usd_mn",
        "related_demand_value_usd_mn", "residual_implied_usd_mn", "reported_market_view_usd_mn",
        "desk_research_reference_value", "reference_delta_pct",
    ):
        val.cell(vr, 1, key)
        val.cell(vr, 2, str(validation.get(key)))
        vr += 1
    for label, key in (
        ("Overlap flags", "overlap_flags"),
        ("Missing category flags", "missing_category_flags"),
        ("Dominant-bucket (unit mismatch risk) flags", "dominant_bucket_flags"),
        ("Country-cell outlier flags", "outlier_flags"),
        ("Intensity-fraction-clamped buckets", "intensity_clamped_buckets"),
        ("Unit contract violations (excluded from calculation)", "unit_contract_violations"),
        ("LLM CALL FAILURES (infrastructure, not genuine estimates)", "llm_call_failures"),
        ("Price plausibility flags", "price_plausibility_flags"),
        ("Unit mismatch flags", "unit_mismatch_flags"),
        ("Stale evidence source flags", "stale_source_flags"),
        ("Driver magnitude implausible flags", "driver_magnitude_flags"),
        ("Triangulation large-gap flags (informational)", "triangulation_large_gap_flags"),
        ("Other problems", "problems"),
    ):
        vr += 1
        val.cell(vr, 1, label).font = bold
        for item in validation.get(key, []) or []:
            vr += 1
            val.cell(vr, 2, str(item))

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(out_path)
        return out_path
    except PermissionError:
        # Most common cause: the file is currently open in Excel (a live
        # file-lock, not a real disk/permission problem) -- never let that
        # crash a run whose JSON output already saved successfully. Fall
        # back to a sibling filename instead of losing the run's output.
        alt_path = out_path.with_stem(out_path.stem + "_new")
        i = 2
        while alt_path.exists():
            alt_path = out_path.with_stem(f"{out_path.stem}_new{i}")
            i += 1
        wb.save(alt_path)
        return alt_path
