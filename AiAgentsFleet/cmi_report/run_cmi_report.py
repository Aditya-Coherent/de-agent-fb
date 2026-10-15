"""One command: CMI-template whitepaper for a market the fleet has run.

    python AiAgentsFleet/cmi_report/run_cmi_report.py --market-slug agricultural-biostimulants-market
    python AiAgentsFleet/cmi_report/run_cmi_report.py --market "Global Drone Market"   # runs the fleet first if needed

Steps: fleet fields (+ consistency pass) -> CMI agents (run_data.py, skipped if
cmi_fields.json exists unless --refresh-data) -> numeric model -> prose ->
charts -> .docx -> Word pass (real TOC page numbers + PDF) -> QA audit.

Output: AiAgentsFleet/output/<slug>/<slug>-cmi-whitepaper.docx (+ .pdf, charts/, qa_report.json)

Numbers are never jittered on this path: the run stops if AIFLEET_RANDOMIZE is
set, unless --allow-randomize is passed (see randomizer.py).
"""
from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
FLEET = HERE.parent
ROOT = FLEET.parent
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))

from me_engine.generation.config import GenConfig  # noqa: E402
from me_engine.generation.llm import GenLLM  # noqa: E402

from AiAgentsFleet import randomizer  # noqa: E402
from AiAgentsFleet.cmi_report import build as builder, charts as ch, fmt, model, prose, run_data, verify_facts  # noqa: E402
from AiAgentsFleet.cmi_report.guard import NumberGuard  # noqa: E402

TEMPLATE = ROOT / "templates" / "cmi_whitepaper_template.docx"
_EVIDENCE_QUERIES = {
    "ch1": ["{market} outlook 2033"], "ch2": ["{market} trends Middle East Africa"],
    "ch3": ["{market} North America Europe Asia Pacific Latin America"],
    "ch4": ["{market} Saudi Arabia UAE Qatar Kuwait"], "ch5": ["{market} South Africa Egypt Nigeria Morocco"],
    "ch6": ["{market} regulation Middle East Africa"], "ch7": ["{market} companies partnerships acquisitions"],
}


def _slug(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", name.strip().lower()).strip("-") or "market"


def ensure_fleet(market: str | None, slug: str) -> Path:
    d = FLEET / "output" / slug
    if not (d / "whitepaper_fields.json").exists():
        if not market:
            raise SystemExit(f"{d / 'whitepaper_fields.json'} missing -- pass --market to run the fleet first")
        subprocess.run([sys.executable, str(FLEET / "run_fleet.py"), "--market", market], check=True, cwd=ROOT)
    data = json.loads((d / "whitepaper_fields.json").read_text(encoding="utf-8"))
    if "fields_before_consistency" not in data:      # consistency pass not yet run
        sys.path.insert(0, str(FLEET / "report"))
        from run_report import run_consistency_pass  # noqa: E402
        run_consistency_pass(d, data["market_name"])
    return d


def make_charts(m: dict, ag: dict, cdir: Path) -> dict[str, Path]:
    C = {}
    C["cover"] = ch.cover(m, cdir / "cover.png")
    C["kpi"] = ch.kpi_panel(m, cdir / "kpi.png")
    C["cagr_share"] = ch.cagr_and_share(m, cdir / "cagr_share.png")
    dyn = ag.get("dynamics") or {}
    C["trends"] = ch.trend_cards(dyn.get("trends") or [], cdir / "trends.png")
    C["drivers"] = ch.drivers_restraints(dyn.get("mea_drivers") or [], dyn.get("mea_restraints") or [],
                                         cdir / "drivers.png")
    C["regions_stacked"] = ch.stacked_years({r: m["regions"][r]["series"] for r in ch.REGION_COLORS},
                                            ch.REGION_COLORS, cdir / "regions_stacked.png")
    C["scenarios"] = ch.scenarios(m, cdir / "scenarios.png")
    geos = ["Global", "Middle East", "Africa"] + m["profiled"]["Middle East"] + m["profiled"]["Africa"]
    sec = m["sections"]
    for gname in geos:
        safe = prose.slug_safe(gname)
        if sec["type"]["show"]:
            C[f"type_{gname}"] = ch.type_chart(m["segments"][gname]["type"], m["dimensions"]["type"]["singular"],
                                               cdir / f"type_{safe}.png")
        if sec["demand"]["show"]:
            C[f"demand_{gname}"] = ch.demand_chart(m["segments"][gname]["demand"], cdir / f"demand_{safe}.png")
    C["region_donuts"] = ch.region_donuts(m, cdir / "region_donuts.png")
    C["region_opportunity"] = ch.opportunity_bars(list(m["regions"].items()), ch.REGION_COLORS,
                                                  f"Absolute dollar opportunity {model.Y0}-{model.Y1} (incremental revenue)",
                                                  cdir / "region_opportunity.png")
    for sub, key in (("Middle East", "me"), ("Africa", "af")):
        units = list(m["countries"][sub])
        colors = {u: ch.CATEGORICAL[k % len(ch.CATEGORICAL)] for k, u in enumerate(units)}
        C[f"{key}_stacked"] = ch.stacked_years(m["country_series"][sub], colors, cdir / f"{key}_stacked.png")
        C[f"{key}_bubble"] = ch.bubble(m["countries"][sub], colors, cdir / f"{key}_bubble.png")
    reg = ag.get("regulation") or {}
    if sec["regulation"]["show"] and sec["milestones"]["show"]:
        C["timeline"] = ch.timeline(reg.get("milestones") or [], cdir / "timeline.png")
    # scorecard countries = regulator-table countries
    from AiAgentsFleet.cmi_report.groups import canonical
    keep_c = model.regulatory_countries(ag)
    sc = {canonical(k): v for k, v in (reg.get("scorecard") or {}).items() if canonical(k) in keep_c}
    if sec["regulation"]["show"] and sc:
        C["reg_scorecard"] = ch.heatmap(list(sc), reg.get("scorecard_dimensions") or [], list(sc.values()),
                                        "Regulatory enablement score (1 = restrictive, 5 = enabling) - CMI assessment",
                                        cdir / "reg_scorecard.png")
    ven = ag.get("vendors") or {}
    C["value_chain"] = ch.value_chain(m["value_chain"], cdir / "value_chain.png")
    C["vendor_global"] = ch.vendor_donut(m["vendor_chart_global"], "Global", cdir / "vendor_global.png")
    if sec["vendor_mea"]["show"]:
        C["vendor_mea"] = ch.vendor_donut(m["vendor_chart_mea"], "MEA", cdir / "vendor_mea.png")
    # one name per corporate group (same names as the share charts); a group listed twice keeps its first row
    caps = {}
    for n, row in (ven.get("capabilities") or {}).items():
        disp = model.display_name(m, n)
        caps.setdefault(disp, row)
    if sec["capabilities"]["show"]:
        C["capabilities"] = ch.heatmap(list(caps), ven.get("capability_dimensions") or [], list(caps.values()),
                                       "Vendor capability mapping (1 = limited, 5 = leading) - CMI assessment",
                                       cdir / "capabilities.png")
    C["phases"] = ch.phases((ag.get("strategy") or {}).get("phases") or [], cdir / "phases.png")
    C["process"] = ch.process(cdir / "process.png")
    pub = m["published_2025"]          # same-scope figures only (model.py)
    if pub:
        C["published"] = ch.published_vs_cmi(pub, m["global"]["v25"], cdir / "published.png")
    return C


_WORD_PS = r"""
$ErrorActionPreference = 'Stop'
$before = @(Get-Process WINWORD -ErrorAction SilentlyContinue | ForEach-Object Id)
$w = New-Object -ComObject Word.Application
$w.Visible = $false
$mine = @(Get-Process WINWORD -ErrorAction SilentlyContinue | Where-Object { $before -notcontains $_.Id } | ForEach-Object Id)
try {
  $d = $w.Documents.Open('%DOCX%', $false, $true)
  $d.Repaginate()
  $pages = @{}
  foreach ($p in $d.Paragraphs) {
    if ($p.OutlineLevel -eq 1) {
      $t = $p.Range.Text.Trim()
      if ($t -match '^(\d+)\s') { $n = [int]$matches[1]; if (-not $pages.ContainsKey("$n")) { $pages["$n"] = $p.Range.Information(3) } }
    }
  }
  if ('%PDF%' -ne '') { $d.SaveAs2('%PDF%', 17) }
  $d.Close($false)
  $pages | ConvertTo-Json -Compress
} finally {
  try { $w.Quit() } catch {}
  [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($w)
  # only the hidden Word instance this script started -- never a user's own Word window
  foreach ($id in $mine) { try { Stop-Process -Id $id -Force -ErrorAction Stop } catch {} }
}
exit 0
"""


def word_pass(docx_path: Path, pdf_path: Path | None) -> dict[int, int]:
    """Opens the .docx in MS Word (if installed): returns {chapter: page} and
    optionally exports a PDF. Returns {} when Word is unavailable."""
    script = _WORD_PS.replace("%DOCX%", str(docx_path)).replace("%PDF%", str(pdf_path or ""))

    def word_pids() -> set[int]:
        r = subprocess.run(["powershell", "-NoProfile", "-Command",
                            "Get-Process WINWORD -ErrorAction SilentlyContinue | ForEach-Object Id"],
                           capture_output=True, text=True)
        return {int(x) for x in r.stdout.split() if x.strip().isdigit()}
    before = word_pids()
    try:
        res = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
                             capture_output=True, text=True, timeout=900)
    except subprocess.TimeoutExpired:
        # the script's own cleanup never ran: stop only the Word instance(s) this step started
        for pid in word_pids() - before:
            subprocess.run(["powershell", "-NoProfile", "-Command", f"Stop-Process -Id {pid} -Force"],
                           capture_output=True)
        print("  [word] timed out -- TOC page numbers / PDF not refreshed")
        return {}
    try:
        # the page map is the JSON line Word printed; cleanup noise must not discard it
        line = next((x for x in reversed(res.stdout.strip().splitlines()) if x.startswith("{")), "")
        if not line:
            print(f"  [word] no page numbers returned: {res.stderr.strip()[-300:]}")
            return {}
        return {int(k): int(v) for k, v in json.loads(line).items()}
    except Exception as exc:  # noqa: BLE001
        print(f"  [word] skipped: {exc}")
        return {}


# standard sentences the builder itself prints in several places (not repetition)
BOILERPLATE = ("Every country in this table",)


def qa(docx_path: Path, m: dict, ag: dict, root: str) -> dict:
    """Audits the finished document: every number must exist in the model or
    the agent data it was printed from; no template leftovers."""
    import docx as _docx
    from AiAgentsFleet.cmi_report.register import FactGuard, build_register, count_numbers
    d = _docx.Document(str(docx_path))
    texts = [p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells]
    # Number consistency sweep: every number in every paragraph and table row must be the
    # register's display value of an entity named in that sentence / row (header row = context).
    ng = FactGuard(build_register(m, ag))
    # paragraphs are swept up to the appendix; the appendix's log paragraphs describe the checking process
    # (precision rule examples, counts) -- its tables (master data, register) are swept with the rest
    units, in_appendix = [], False
    for p in d.paragraphs:
        if p.style.name == "Heading 1" and "Appendix" in p.text:
            in_appendix = True
        if not in_appendix:
            units.append(p.text)
    for t in d.tables:
        header = " | ".join(c.text for c in t.rows[0].cells)
        units += [header + " || " + " | ".join(c.text for c in r.cells) for r in t.rows[1:]]
    issues, checked = [], 0
    for t in units:
        # section numbers ('6.1 ...'), figure/table numbers ('Figure 41:') and contents-page
        # page numbers ('... Way Forward<tab>31') are not data
        t = re.sub(r"\t\s*\d+\s*$", "", t)
        t = re.sub(r"^\s*(?:(?:Figure|Table)\s+\d+:|\d+(?:\.\d+)*\s)", "", t)
        checked += count_numbers(t)
        bad = ng.bad(t)
        if bad:
            issues.append({"numbers_not_in_model": [b.split(" ")[0] for b in bad], "why": bad, "text": t[:200]})
    leftovers = [t[:120] for t in texts if re.search(r"\[insert|draft template|publication template|disclaimer",
                                                     t, re.I)]
    market_words = set(re.findall(r"[a-z]+", m["market"].lower()))
    drone_words = [t[:120] for t in texts if re.search(r"\b(drones?|UAS|UAV|BVLOS)\b", t, re.I)] \
        if not market_words & {"drone", "drones", "uav", "uas"} else []
    # repeated sentences anywhere in the finished document (same words, no new figure)
    from AiAgentsFleet.cmi_report.guard import sentences
    seen, repeats = [], []
    # body prose only: paragraphs and single-cell callout boxes. Left out by design: data-table cells
    # (the same use case in two heat maps), the executive summary (it restates the body), "Finding:"
    # lines (recommendations cite the body) and the appendix (it logs the body).
    prose_texts, zone = [], "front"
    for p in d.paragraphs:
        if p.style.name == "Heading 1":
            zone = ("summary" if re.match(r"\s*0?1\s", p.text) else
                    "appendix" if "Appendix" in p.text else "body")
        if zone == "body" and not p.text.strip().startswith("Finding:"):
            prose_texts.append(p.text)
    prose_texts += [t.rows[0].cells[0].text for t in d.tables if len(t.rows) == 1 and len(t.columns) == 1
                    and "Key conclusions" not in t.rows[0].cells[0].text]
    for t in prose_texts:
        for s in sentences(t):
            w = {x for x in re.findall(r"[a-z0-9$.%+]+", s.lower()) if len(x) > 2}
            n = set(re.findall(r"\d[\d,.]*", s))
            # headings and table cells are parallel by design ('X Market by Country'); only prose counts.
            # "Finding:" lines quote the body on purpose (recommendations trace to findings).
            if len(s.split()) < 10 or s.lstrip().startswith(("Finding:",) + BOILERPLATE):
                continue
            if any(len(w & x) / len(w | x) >= 0.7 and n <= xn for x, xn in seen):
                repeats.append(s[:120])
            else:
                seen.append((w, n))
    removed = [f"{k}: {v['why']}" for k, v in m["sections"].items() if not v["show"]]
    # one fact = one number: every repeat of the same fact (entity + kind + year) must be identical
    from AiAgentsFleet.cmi_report import factcheck
    body = []
    for p in d.paragraphs:
        if p.style.name == "Heading 1" and "Appendix" in p.text:
            break
        body.append(p.text)
    fc = factcheck.check(build_register(m, ag), body,
                         [[[c.text.strip() for c in r.cells] for r in t.rows] for t in d.tables])
    return {"same_fact_inconsistencies": fc["inconsistent_repeated_facts"],
            "facts_stated_more_than_once": fc["facts_stated_more_than_once"],
            "fact_identity_flags": fc["mismatches"],
            "numbers_not_in_model": issues, "numbers_checked": checked, "template_leftovers": leftovers,
            "drone_wording": drone_words, "repeated_sentences": repeats, "removed_sections": removed}


def scrub(pr: dict, issues: list[dict]) -> int:
    """Removes from the prose every sentence the final QA flagged (a number
    that is not the model's), so the rebuilt document no longer contains it."""
    from AiAgentsFleet.cmi_report.guard import sentences
    flagged = [(set(i["numbers_not_in_model"]), i["text"]) for i in issues]
    removed = 0

    def fix(v):
        nonlocal removed
        if isinstance(v, str):
            keep = []
            for s in sentences(v):
                # the flagged token is in this sentence and this sentence is in the flagged text
                if any(any(tok and tok in s for tok in bad) and s[:40] in txt for bad, txt in flagged):
                    removed += 1
                else:
                    keep.append(s)
            return " ".join(keep)
        if isinstance(v, list):
            return [fix(x) for x in v]
        if isinstance(v, dict):
            return {k: fix(x) for k, x in v.items()}
        return v

    pr["slots"] = fix(pr["slots"])
    return removed


def main() -> int:
    p = argparse.ArgumentParser()
    g = p.add_mutually_exclusive_group(required=True)
    g.add_argument("--market-slug")
    g.add_argument("--market")
    p.add_argument("--refresh-data", action="store_true", help="re-run the CMI agents even if cached")
    p.add_argument("--no-word", action="store_true", help="skip the MS Word pass (TOC page numbers, PDF)")
    p.add_argument("--refresh-prose", action="store_true", help="rewrite the text even if the numbers are unchanged")
    p.add_argument("--allow-randomize", action="store_true",
                   help="permit AIFLEET_RANDOMIZE=1 (jittered numbers; never for client reports)")
    a = p.parse_args()
    randomizer.require_off(a.allow_randomize)
    slug = a.market_slug or _slug(a.market)

    mdir = ensure_fleet(a.market, slug)
    fleet = json.loads((mdir / "whitepaper_fields.json").read_text(encoding="utf-8"))
    cdir = mdir      # everything for a market lives in output/<slug>/
    if a.refresh_data or not (cdir / "cmi_fields.json").exists():
        run_data.run(slug)
    cmi = json.loads((cdir / "cmi_fields.json").read_text(encoding="utf-8"))
    # every regulator, requirement, milestone and vendor move checked claim by claim (PLAN-06); the
    # results are stored with the agent data and in fact_checks.json, and unconfirmed items are left out
    if verify_facts.ensure(cdir, fleet["market_name"], cmi["fields"],
                           verify_facts.judge_llm(GenConfig.load(ROOT), ROOT / "output" / ".gen_cache.db")):
        (cdir / "cmi_fields.json").write_text(json.dumps(cmi, indent=2, ensure_ascii=False), encoding="utf-8")
    ag = verify_facts.apply(cmi["fields"])
    m = model.build(fleet, cmi)
    definition = m["definition"]
    ag["_definition"] = definition
    model.write(mdir)
    print(f"Model: global {fmt.usd(m['global']['v26'])} ({model.Y0}) -> {fmt.usd(m['global']['v33'])} ({model.Y1}), "
          f"MEA {fmt.usd(m['regions']['Middle East & Africa']['v26'])}")

    llm = GenLLM(GenConfig.load(ROOT), cache_path=ROOT / "output" / ".gen_cache.db")
    market = fleet["market_name"]

    def evidence(chap: str) -> list[str]:
        qs = [q.format(market=market) for q in _EVIDENCE_QUERIES.get(chap, [])]
        return [s["text"] for s in run_data.search(market, qs, chap, [], min_hits=2)] if qs else []

    # Text is reused while the numbers and sections are unchanged (layout fixes then
    # cost no LLM calls); --refresh-prose writes it anew.
    import hashlib
    mhash = hashlib.sha256(json.dumps(m, sort_keys=True, default=str).encode()).hexdigest()[:16]
    ppath = cdir / "cmi_prose.json"
    old = json.loads(ppath.read_text(encoding="utf-8")) if ppath.exists() else {}
    if (not a.refresh_prose and old.get("model_hash") == mhash and old.get("version") == prose.PROSE_VERSION):
        pr = old
        prose.set_template(TEMPLATE)
        n = prose.reguard(pr, m, ag)
        print("  [prose] numbers unchanged -- reusing cmi_prose.json"
              + (f" ({n} sentence(s) removed by the current guards)" if n else ""))
    else:
        pr = prose.write(mdir, m, ag, llm, TEMPLATE, evidence)
        pr["model_hash"] = mhash
        ppath.write_text(json.dumps(pr, indent=2, ensure_ascii=False), encoding="utf-8")
    charts = make_charts(m, ag, cdir / "charts")

    out = cdir / f"{slug}-cmi-whitepaper.docx"
    root = builder._root(market)
    builder.build(TEMPLATE, out, m, pr, ag, charts, definition)
    # Number consistency sweep: every number vs the register; flagged sentences are removed and
    # the sweep repeated until clean (max 3 rounds) -- the counts go into the verification log
    q = qa(out, m, ag, root)
    verification = {"checked": q["numbers_checked"], "first_sweep": len(q["numbers_not_in_model"]), "fixed": 0}
    for _ in range(3):
        if not q["numbers_not_in_model"]:
            break
        n = scrub(pr, q["numbers_not_in_model"])
        verification["fixed"] += n
        if not n:
            break
        builder.build(TEMPLATE, out, m, pr, ag, charts, definition)
        q = qa(out, m, ag, root)
    verification["final_sweep"] = len(q["numbers_not_in_model"])
    log = builder.build(TEMPLATE, out, m, pr, ag, charts, definition, verification=verification)
    if not a.no_word:
        pages = word_pass(out, None)
        if pages:
            log = builder.build(TEMPLATE, out, m, pr, ag, charts, definition, toc_pages=pages,
                                verification=verification)
            word_pass(out, out.with_suffix(".pdf"))
    # the complete Number Register beside the report (the appendix carries its derived facts)
    import csv
    from AiAgentsFleet.cmi_report.register import build_register
    with (cdir / f"{slug}-number-register.csv").open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["ID", "Fact", "Display value", "Metric", "Entities"])
        for e in build_register(m, ag):
            w.writerow([e["id"], e["fact"], e["display"], e["metric"], "; ".join(x for x in e["entities"] if x)])
    report = {"build_log": log, "prose_dropped": pr.get("dropped"), "verification": verification,
              "headline_size_check": {k: m["size_check"].get(k) for k in ("status", "notes", "source")},
              "fact_review": {k: m["fact_review"].get(k) for k in ("status", "claims", "supported", "corrected",
                                                                   "omitted", "blockers", "warnings")},
              **qa(out, m, ag, root)}
    (cdir / "qa_report.json").write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"QA: {len(report['numbers_not_in_model'])} text block(s) with numbers not in the model, "
          f"{len(report['template_leftovers'])} template leftover(s), {len(report['drone_wording'])} drone-wording hit(s), "
          f"{len(report['repeated_sentences'])} repeated sentence(s); {report['numbers_checked']} numbers checked; "
          f"{len(report['same_fact_inconsistencies'])} inconsistent repeated fact(s) across "
          f"{report['facts_stated_more_than_once']} facts stated more than once")
    for r in report["removed_sections"]:
        print(f"  section left out -- {r}")
    print(f"Done: {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
