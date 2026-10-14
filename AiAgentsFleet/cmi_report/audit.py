"""Audits a finished CMI whitepaper against the report standard:
internal consistency, factual-accuracy labelling, and relevance.

Two kinds of checks, reported separately:
  * computed -- exact checks on the model and the .docx (sums, CAGRs, share
    claims vs growth, ties, blank cells, announced counts, profiles, labels);
  * ai_review -- one LLM review per market for judgement calls (missing major
    countries, parent/subsidiary double counting, segments vs definition,
    regulator currency, recommendations without a finding). These are
    assessments, not proofs.

    python AiAgentsFleet/cmi_report/audit.py --market-slug <slug> [--no-ai]
Writes output/<slug>/audit_report.json.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
FLEET = HERE.parent
ROOT = FLEET.parent
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))

import docx  # noqa: E402

from AiAgentsFleet.cmi_report import fmt, model  # noqa: E402
from AiAgentsFleet.cmi_report.guard import sentences  # noqa: E402

_NUMWORD = {"two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7, "eight": 8, "nine": 9, "ten": 10}


def _doc_texts(path: Path) -> tuple[list[str], list[list[list[str]]]]:
    d = docx.Document(str(path))
    paras = [p.text for p in d.paragraphs if p.text.strip()]
    tables = [[[c.text.strip() for c in r.cells] for r in t.rows] for t in d.tables]
    return paras, tables


# --------------------------------------------------------------------------- #
# computed checks
# --------------------------------------------------------------------------- #
def check_math(m: dict) -> list[str]:
    out = []
    tol = lambda a, b: abs(a - b) <= max(0.5, abs(b) * 0.002)  # noqa: E731
    g = m["global"]
    for y in m["years"]:
        s = sum(r["series"][y] for r in m["regions"].values())
        if not tol(s, g["series"][y]):
            out.append(f"regions sum to {s:.1f} in {y}, global is {g['series'][y]:.1f}")
    mea = m["regions"]["Middle East & Africa"]
    for k in ("v26", "v33"):
        s = sum(r[k] for r in m["subregions"].values())
        if not tol(s, mea[k]):
            out.append(f"Middle East + Africa = {s:.1f} ({k}), MEA = {mea[k]:.1f}")
    for reg, d in m["countries"].items():
        parent = m["regions"].get(reg) or m["subregions"].get(reg)
        for k in ("v26", "v33"):
            s = sum(c[k] for c in d.values())
            if not tol(s, parent[k]):
                out.append(f"{reg} countries sum to {s:.1f} ({k}), region is {parent[k]:.1f}")
    for geo, dims in m["segments"].items():
        parent = (g if geo == "Global" else m["subregions"].get(geo)
                  or next((d[geo] for d in m["countries"].values() if geo in d), None))
        for dim, segs in dims.items():
            if not segs or not m["sections"][dim]["show"]:
                continue
            for k in ("v26", "v33"):
                s = sum(x[k] for x in segs.values())
                if parent and not tol(s, parent[k]):
                    out.append(f"{geo} {dim} segments sum to {s:.1f} ({k}), total is {parent[k]:.1f}")
    # an aggregate's CAGR lies within the range of its components' CAGRs
    def within(name, agg, parts):
        cs = [p["cagr"] for p in parts]
        if cs and not (min(cs) - 0.05 <= agg["cagr"] <= max(cs) + 0.05):
            out.append(f"{name}: CAGR {agg['cagr']:.2f}% is outside its components' range {min(cs):.2f}-{max(cs):.2f}%")
    within("Global (regions)", g, list(m["regions"].values()))
    within("Middle East & Africa (sub-regions)", mea, list(m["subregions"].values()))
    for reg, d in m["countries"].items():
        within(f"{reg} (countries)", m["regions"].get(reg) or m["subregions"].get(reg), list(d.values()))
    # CAGR and increment arithmetic on every node
    nodes = [("global", g), *m["regions"].items(), *m["subregions"].items(),
             *[(c, v) for d in m["countries"].values() for c, v in d.items()]]
    for name, n in nodes:
        if abs(model.cagr(n["v26"], n["v33"]) - n["cagr"]) > 0.05:
            out.append(f"{name}: CAGR {n['cagr']:.2f}% does not match its 2026/2033 values")
        if abs((n["v33"] - n["v26"]) - n["incr"]) > 0.05:
            out.append(f"{name}: incremental value is not 2033 minus 2026")
    return out


_SHARE = re.compile(r"\b(gain(?:s|ing)?|los(?:es|ing|e)|increas\w*|ris\w*|expand\w*|declin\w*|fall\w*|shrink\w*)"
                    r"\b[^.]{0,30}\bshare\b|\bshare\b[^.]{0,30}\b(ris\w*|increas\w*|gain\w*|fall\w*|declin\w*|"
                    r"shrink\w*|drop\w*)", re.I)


def check_share_claims(m: dict, paras: list[str]) -> list[str]:
    """'X gains share' must mean X grows faster than its parent, and vice versa."""
    out = []
    ents = {**{r: (v, m["global"]) for r, v in m["regions"].items()},
            **{r: (v, m["regions"]["Middle East & Africa"]) for r, v in m["subregions"].items()}}
    for d_reg, d in m["countries"].items():
        parent = m["regions"].get(d_reg) or m["subregions"].get(d_reg)
        ents.update({c: (v, parent) for c, v in d.items()})
    ents["MEA"] = ents["Middle East & Africa"]
    for p in paras:
        for sent in sentences(p):
            # judge each clause on its own: in "X is the largest, while Y's share falls" the claim is about Y
            for clause in re.split(r",\s*(?:while|whereas|but)\b|;", sent):
                mm = _SHARE.search(clause)
                if not mm:
                    continue
                up = bool(re.search(r"gain|increas|ris|expand", mm.group(0), re.I))
                # the entity named closest before the share wording is the claim's subject
                before = clause[: mm.start()] + " " + mm.group(0)
                hits = [(before.rfind(n), n) for n in ents
                        if re.search(rf"(?<![A-Za-z]){re.escape(n)}(?![A-Za-z])", before)]
                if not hits:
                    continue
                name = max(hits, key=lambda h: (h[0], len(h[1])))[1]
                v, parent = ents[name]
                if up and v["cagr"] < parent["cagr"] or (not up and v["cagr"] > parent["cagr"]):
                    out.append(f"'{clause.strip()[:110]}' -- {name} grows {fmt.pct(v['cagr'])} vs parent "
                               f"{fmt.pct(parent['cagr'])}")
    return out


def check_ties(m: dict) -> list[str]:
    """Rankings whose top two are equal at printed precision."""
    out = []
    groups = {"regions": m["regions"], "subregions": m["subregions"], **m["countries"]}
    for geo, dims in m["segments"].items():
        for dim, segs in dims.items():
            if segs and m["sections"][dim]["show"]:
                groups[f"{geo} {dim}"] = segs
    for gname, d in groups.items():
        for key, label in (("cagr", "fastest"), ("v33", "largest")):
            vals = sorted(((fmt.pct(v[key]) if key == "cagr" else fmt.usd(v[key])), n) for n, v in d.items())
            top = sorted(d, key=lambda n: -d[n][key])[:2]
            if len(top) == 2:
                a, b = top
                fa = fmt.pct(d[a][key]) if key == "cagr" else fmt.usd(d[a][key])
                fb = fmt.pct(d[b][key]) if key == "cagr" else fmt.usd(d[b][key])
                if fa == fb:
                    out.append(f"{gname}: '{label}' is a tie at {fa} between {a} and {b}")
    return out


def check_tables(tables) -> list[str]:
    out = []
    for ti, t in enumerate(tables):
        for r in t[1:]:
            for c in r:
                if c in ("", "-", "–", "—", "n/a"):
                    out.append(f"table {ti + 1} ('{t[0][0]}'), row '{r[0][:40]}': empty or '-' cell")
    return out


def check_counts(paras: list[str], tables, ag: dict) -> list[str]:
    """'Six trends', 'Five priority actions', 'Three-phase' must be delivered in full."""
    out = []
    delivered = {
        "trend": len((ag.get("dynamics") or {}).get("trends") or []),
        "phase": len((ag.get("strategy") or {}).get("phases") or []),
    }
    for p in paras + [t[0][0] for t in tables if t and t[0]]:
        for mm in re.finditer(r"\b(two|three|four|five|six|seven|eight|nine|ten)[\s-]+(?:\w+\s+){0,2}?"
                              r"(trend|phase|priority action|action|finding|driver|scenario)s?\b", p, re.I):
            n, what = _NUMWORD[mm.group(1).lower()], mm.group(2).lower()
            if what in delivered and delivered[what] and delivered[what] != n:
                out.append(f"'{mm.group(0)}' announced, {delivered[what]} delivered")
            if what in ("priority action", "action"):
                box = next((t for t in tables if t and t[0] and mm.group(0).lower() in t[0][0].lower()), None)
                if box:
                    items = len(re.findall(r"(?:^|\n)\s*\d+\.", box[0][0]))
                    if items and items != n:
                        out.append(f"'{mm.group(0)}' announced, {items} listed")
    return sorted(set(out))


def check_profiles(m: dict, paras: list[str]) -> list[str]:
    """Every country named in a table needs at least a short profile ('Country: ...'
    line, or a Chapter 4/5 profile heading)."""
    profiled = {u for v in m["profiled"].values() for u in v}
    starts = {p.split(":")[0].strip() for p in paras if ":" in p[:40]}
    table_countries = [c for d in m["countries"].values() for c in d if not c.startswith("Rest of")]
    missing = [c for c in table_countries if c not in profiled and c not in starts]
    return [f"{len(missing)} of {len(table_countries)} countries in the tables have no profile: " + ", ".join(missing)] \
        if missing else []


def check_exec_vs_body(paras: list[str]) -> list[str]:
    """Numbers in the executive summary must also appear in the body."""
    try:
        start = next(i for i, p in enumerate(paras) if re.match(r"^\s*0?1\s+Executive Summary", p))
        end = next(i for i, p in enumerate(paras) if i > start and re.match(r"^\s*0?2\s+", p))
    except StopIteration:
        return ["could not locate the executive summary"]
    body = " ".join(paras[end:])
    out = []
    for p in paras[start:end]:
        for tok in re.findall(r"[+$]?\$?\d[\d,.]*[BM%]?", p):
            t = tok.strip(".,")
            if re.search(r"\d", t) and len(t) > 2 and t not in body and not re.fullmatch(r"20\d\d[EF]?", t):
                out.append(f"'{t}' in the executive summary does not appear in the body")
    return sorted(set(out))


_EST = re.compile(r"\b(estimat\w*|forecast\w*|project\w*|CMI|expected|assess\w*|indicative|scenario)\b", re.I)


def check_labels(paras: list[str]) -> dict:
    """Share of number-bearing PARAGRAPHS (not headings, not the appendix) with no
    estimate/forecast/CMI label, and count of [to verify] / [unverified] markers."""
    body = []
    for p in paras:
        if "\t" in p:                       # contents-page line
            continue
        if re.match(r"\s*1?1\s+Appendix", p):
            break
        if re.match(r"^\s*\d+(\.\d+)?\s+\S", p) and len(p) < 90:      # a numbered heading
            continue
        body.append(p)
    with_num = [p for p in body if re.search(r"\$\d|\d%", p) and len(p) > 60]
    unlabelled = [s for s in with_num if not _EST.search(s)]
    markers = sum(len(re.findall(r"\[(?:to verify|unverified)\]", p, re.I)) for p in paras)
    return {"number_paragraphs": len(with_num), "without_estimate_label": len(unlabelled),
            "examples": [s[:110] for s in unlabelled[:5]], "to_verify_markers": markers}


# --------------------------------------------------------------------------- #
# AI review (judgement)
# --------------------------------------------------------------------------- #
_AI_SYSTEM = (
    "You are a senior market-research editor auditing a market whitepaper for relevance and factual risk. Be "
    "concrete and conservative: only flag what you are confident about, and say why. Return ONLY JSON."
)
_AI_SCHEMA = (
    '{"missing_major_countries": [{"country": <string>, "currently": <"in Rest of <region>" or "absent">, '
    '"why_it_matters": <1 sentence>}], '
    '"double_counted_companies": [{"companies": [<names>], "issue": <1 sentence>}], '
    '"segments_vs_definition": [{"segment_or_dimension": <string>, "issue": <1 sentence>}], '
    '"questionable_regulators": [{"country": <string>, "named": <string>, "issue": <1 sentence>}], '
    '"recommendations_without_finding": [{"recommendation": <short>, "issue": <1 sentence>}], '
    '"other_relevance_issues": [<1 sentence each>]}'
)


def ai_review(market_dir: Path, m: dict, ag: dict, paras: list[str]) -> dict:
    from me_engine.generation.config import GenConfig
    from me_engine.generation.llm import GenLLM
    prose = json.loads((market_dir / "cmi_prose.json").read_text(encoding="utf-8")) \
        if (market_dir / "cmi_prose.json").exists() else {}
    traced = [f"{a.get('lead')} {a.get('text')}" for a in (prose.get("slots") or {}).get("priority_actions") or []
              if isinstance(a, dict)]
    facts = {
        "market": m["market"],
        "definition": m["definition"],
        "segment_dimensions": {m["dimensions"][d]["singular"]: list(m["segments"]["Global"][d])
                               for d in ("type", "demand") if m["sections"][d]["show"]},
        "countries_named_per_region": {r: [c for c in d] for r, d in m["countries"].items()},
        "region_shares_2026_pct": {r: round(v["v26"] / m["global"]["v26"] * 100, 1) for r, v in m["regions"].items()},
        "company_share_chart_global": [n for n, _ in m["vendor_chart_global"]],
        "company_share_chart_mea": [n for n, _ in m["vendor_chart_mea"]],
        "regulators_named": [(r.get("country"), r.get("regulator")) for r in (ag.get("regulation") or {}).get("snapshot") or []],
        "recommendations_with_their_findings": traced or (ag.get("strategy") or {}).get("priority_actions"),
        "stakeholder_moves": [s.get("priority_moves") for s in (ag.get("strategy") or {}).get("stakeholders") or []],
        "key_findings_text": [p for p in paras if len(p) > 80 and not re.match(r"^\s*\d+\s{2,}", p)][:12],
        "note": ("The global and MEA share charts are separate charts (MEA is a regional view), so the same company "
                 "appearing in both is not double counting; flag only a group listed twice WITHIN one chart."),
    }
    llm = GenLLM(GenConfig.load(ROOT), cache_path=ROOT / "output" / ".gen_cache.db")
    try:
        return llm.complete_json(_AI_SYSTEM, f"REPORT FACTS:\n{json.dumps(facts, ensure_ascii=False)}\n\n"
                                             f"Audit it. Return JSON exactly in this shape:\n{_AI_SCHEMA}")
    except Exception as exc:  # noqa: BLE001
        return {"error": str(exc)}


# --------------------------------------------------------------------------- #
def audit(slug: str, use_ai: bool = True) -> dict:
    mdir = FLEET / "output" / slug
    fleet = json.loads((mdir / "whitepaper_fields.json").read_text(encoding="utf-8"))
    cmi = json.loads((mdir / "cmi_fields.json").read_text(encoding="utf-8"))
    ag = cmi["fields"]
    m = model.build(fleet, cmi)
    paras, tables = _doc_texts(mdir / f"{slug}-cmi-whitepaper.docx")
    res = {
        "market": m["market"],
        "computed": {
            "math_sums_cagr_increments": check_math(m),
            "share_claims_vs_growth": check_share_claims(m, paras),
            # informational: ties in the data (the claim guard blocks any untied claim about them)
            "ranking_ties_info": check_ties(m),
            "blank_or_dash_cells": check_tables(tables),
            "announced_counts": check_counts(paras, tables, ag),
            "countries_without_profile": check_profiles(m, paras),
            "exec_summary_numbers_not_in_body": check_exec_vs_body(paras),
            "labels": check_labels(paras),
        },
    }
    if use_ai:
        res["ai_review"] = ai_review(mdir, m, ag, paras)
    (mdir / "audit_report.json").write_text(json.dumps(res, indent=2, ensure_ascii=False), encoding="utf-8")
    return res


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--market-slug", required=True)
    p.add_argument("--no-ai", action="store_true")
    a = p.parse_args()
    r = audit(a.market_slug, not a.no_ai)
    print(json.dumps(r, indent=1, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
