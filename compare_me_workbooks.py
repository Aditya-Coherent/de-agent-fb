"""Compare two already-built ME workbooks, cell-by-cell — no regeneration.

Manual accuracy-testing tool, separate from the main generate flow (Input
Sheet -> ME workbook, which never does this automatically). Use it whenever
you want to check a generated ME workbook against a known-good reference,
without re-running the agent fleet.

Matches sheets by geography name and rows by (band, label) — the same
(band, label) -> row scheme `SheetLayout.discover` uses, re-derived here from
an in-memory grid instead (see `_grid`/`_discover`) for speed — so it's robust
to the two workbooks listing geographies/segments in a different order, and
just skips whatever doesn't exist in both (reporting the overlap size, same
idea as compare_input_sheets.py).

Example:
    python compare_me_workbooks.py \
        --generated "agentGeneratedMEWorkbook/ME - Global Process Analyzer Market.xlsx" \
        --reference "my_reference/ME-Global Process Analyzer Market.xlsx" \
        --save-report "output/compare-me-process-analyzer.json"
"""
from __future__ import annotations

import argparse
import json
import statistics
import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent / "src"))

from openpyxl import load_workbook

from me_engine.domain.taxonomy import YEARS, Band
from me_engine.io.layout import COL_LABEL, COL_METRIC_START, metric_col

_BAND_BY_TITLE = {b.value: b for b in Band}


def _metrics(pairs: list[tuple[float, float]]) -> dict[str, Any]:
    """pairs = [(reference, generated), ...]"""
    n = len(pairs)
    if n == 0:
        return {"n": 0}
    rel_errs = [abs(gen - ref) / abs(ref) for ref, gen in pairs if ref != 0]
    if not rel_errs:
        return {"n": n, "note": "all reference values were 0 — relative error undefined"}
    sorted_errs = sorted(rel_errs)
    m = len(sorted_errs)
    return {
        "n": n,
        "mean_rel_error": statistics.mean(rel_errs),
        "median_rel_error": statistics.median(rel_errs),
        "p90_rel_error": sorted_errs[int(m * 0.9)] if m > 1 else sorted_errs[0],
        "worst_rel_error": max(rel_errs),
        "within_5pct": sum(1 for e in rel_errs if e <= 0.05) / m * 100,
    }


def _grid(ws) -> list[tuple]:
    """Read a whole sheet sequentially into memory once.

    openpyxl's read_only worksheets are only fast for sequential access —
    scattered/out-of-order `.cell(row, col)` calls (what a naive lookup-by-
    label loop does, and what `SheetLayout.discover` does internally too) each
    re-seek the underlying XML stream and are catastrophically slow (measured
    ~10-17ms/cell here — minutes for a full multi-sheet workbook). Reading
    once via `iter_rows` and indexing the resulting list in plain Python
    avoids that entirely; everything below works off this in-memory grid, no
    further openpyxl cell access.
    """
    return list(ws.iter_rows(values_only=True))


def _cell(grid: list[tuple], row: int, col: int):
    r, c = row - 1, col - 1
    if 0 <= r < len(grid) and 0 <= c < len(grid[r]):
        return grid[r][c]
    return None


def _discover(grid: list[tuple]) -> dict[tuple, int]:
    """Same (band, label) -> row mapping as `SheetLayout.discover`, but derived
    from an already-loaded in-memory grid instead of re-scanning via slow
    per-cell openpyxl calls."""
    rows: dict[tuple, int] = {}
    current: Band | None = None
    for r, row in enumerate(grid, start=1):
        marker = row[COL_METRIC_START - 1] if len(row) >= COL_METRIC_START else None
        if isinstance(marker, str) and marker in _BAND_BY_TITLE:
            current = _BAND_BY_TITLE[marker]
            continue
        label = row[COL_LABEL - 1] if len(row) >= COL_LABEL else None
        if current is not None and isinstance(label, str) and label.strip():
            rows[(current, label)] = r
    return rows


def compare_workbooks(generated_path: Path, reference_path: Path) -> dict[str, Any]:
    gen_wb = load_workbook(generated_path, data_only=True, read_only=True)
    ref_wb = load_workbook(reference_path, data_only=True, read_only=True)

    common_sheets = sorted(set(gen_wb.sheetnames) & set(ref_wb.sheetnames))
    all_pairs: list[tuple[float, float]] = []
    per_band: dict[str, list[tuple[float, float]]] = {b.value: [] for b in Band}
    per_geo: dict[str, list[tuple[float, float]]] = {}
    worst_cells: list[tuple[float, str, str, str, int, float, float]] = []

    for geo in common_sheets:
        gen_grid, ref_grid = _grid(gen_wb[geo]), _grid(ref_wb[geo])
        gen_rows, ref_rows = _discover(gen_grid), _discover(ref_grid)
        common_keys = sorted(set(gen_rows) & set(ref_rows))
        geo_pairs: list[tuple[float, float]] = []

        for band, label in common_keys:
            gen_row, ref_row = gen_rows[(band, label)], ref_rows[(band, label)]
            for year in YEARS:
                col = metric_col(year)
                gen_val, ref_val = _cell(gen_grid, gen_row, col), _cell(ref_grid, ref_row, col)
                if not isinstance(gen_val, (int, float)) or not isinstance(ref_val, (int, float)):
                    continue
                pair = (float(ref_val), float(gen_val))
                all_pairs.append(pair)
                per_band[band.value].append(pair)
                geo_pairs.append(pair)
                if ref_val != 0:
                    rel_err = abs(gen_val - ref_val) / abs(ref_val)
                    worst_cells.append((rel_err, geo, band.value, label, year, ref_val, gen_val))

        if geo_pairs:
            per_geo[geo] = geo_pairs

    worst_cells.sort(key=lambda c: c[0], reverse=True)

    return {
        "generated": str(generated_path),
        "reference": str(reference_path),
        "sheets_compared": len(common_sheets),
        "sheets_generated_only": sorted(set(gen_wb.sheetnames) - set(ref_wb.sheetnames)),
        "sheets_reference_only": sorted(set(ref_wb.sheetnames) - set(gen_wb.sheetnames)),
        "overall": _metrics(all_pairs),
        "per_band": {band: _metrics(pairs) for band, pairs in per_band.items()},
        "per_geography": {geo: _metrics(pairs) for geo, pairs in per_geo.items()},
        "worst_cells": [
            {"geography": g, "band": b, "segment": s, "year": y, "reference": ref, "generated": gen,
             "rel_error": err}
            for err, g, b, s, y, ref, gen in worst_cells[:15]
        ],
    }


def format_report(rep: dict[str, Any]) -> str:
    def fmt(m: dict[str, Any]) -> str:
        if not m or m.get("n", 0) == 0:
            return "  (no overlapping cells)"
        if "note" in m:
            return f"  n={m['n']:>5}  {m['note']}"
        return (
            f"  n={m['n']:>5}  mean={m['mean_rel_error']:.2%}  median={m['median_rel_error']:.2%}  "
            f"p90={m['p90_rel_error']:.2%}  worst={m['worst_rel_error']:.2%}  "
            f"within5%={m['within_5pct']:.0f}%"
        )

    lines = [
        "=" * 78,
        "GENERATED vs REFERENCE — ME workbook comparison",
        "=" * 78,
        f"generated : {rep['generated']}",
        f"reference : {rep['reference']}",
        f"sheets compared: {rep['sheets_compared']}",
    ]
    if rep["sheets_generated_only"]:
        lines.append(f"  only in generated: {', '.join(rep['sheets_generated_only'])}")
    if rep["sheets_reference_only"]:
        lines.append(f"  only in reference: {', '.join(rep['sheets_reference_only'])}")
    lines += ["", "OVERALL:", fmt(rep["overall"]), "", "PER BAND:"]
    for band, m in rep["per_band"].items():
        lines.append(f"  {band}:")
        lines.append("  " + fmt(m))
    if rep["worst_cells"]:
        lines += ["", f"Top {len(rep['worst_cells'])} largest deviations:"]
        for c in rep["worst_cells"]:
            lines.append(
                f"  {c['geography']:18s} | {c['band']:28s} | {c['segment']:35s} | {c['year']} "
                f"| ref={c['reference']:>12.2f} gen={c['generated']:>12.2f} ({c['rel_error']:+.1%})"
            )
    lines.append("=" * 78)
    return "\n".join(lines)


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Compare two ME workbooks cell-by-cell (no regeneration).")
    p.add_argument("--generated", required=True, help="The generated ME workbook to check")
    p.add_argument("--reference", required=True, help="The known-good reference ME workbook")
    p.add_argument("--save-report", default=None, help="Optional JSON report output path")
    return p.parse_args()


def main() -> int:
    args = parse_args()
    rep = compare_workbooks(Path(args.generated), Path(args.reference))
    print(format_report(rep))
    if args.save_report:
        report_path = Path(args.save_report)
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(rep, indent=2, ensure_ascii=False), encoding="utf-8")
        print(f"\nSaved report: {report_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
