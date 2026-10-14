import json
import math
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from me_engine.generation import usecase_agents as ua  # noqa: E402
from me_engine.generation import usecase_builder as ub  # noqa: E402

GEOS = ["U.S.", "Germany", "India"]
SPEC = {"geographies": GEOS, "regions": {"North America": ["U.S."], "Europe": ["Germany"], "Asia Pacific": ["India"]}}
WB = {"USA": {"name": "United States", "region": "North America", "gdp_usd": 2.7e13, "medhigh_tech_manufacturing_usd": 1.2e12},
      "DEU": {"name": "Germany", "region": "Europe & Central Asia", "gdp_usd": 4.4e12, "medhigh_tech_manufacturing_usd": 4e11},
      "IND": {"name": "India", "region": "South Asia", "gdp_usd": 3.5e12, "medhigh_tech_manufacturing_usd": 2e11}}
NEST = {"By Application": {"__root__": ["Turbines", "Pumps"]}}
CLASSES = [{"class": "small", "new_machines_per_year": {"value": 1000, "basis": "E1", "reason": "r"}},
           {"class": "large", "new_machines_per_year": {"value": 200, "basis": "assumption", "reason": "r"}}]


class FakeLLM:
    """Stateless (agents run in parallel). The price agent first breaks the critic's rule."""

    def __init__(self, oem=True):
        self.oem = oem

    def complete_json(self, system, user):
        if "decide WHAT a market counts" in system:
            return {"count_replacements": not self.oem, "demand_location": "machine_builder" if self.oem else "installation",
                    "product_unit": "one gearbox", "reason": "OEM market"}
        if "pick the ONE country-level driver" in system:
            return {u: {"driver": "medhigh_tech_manufacturing_usd", "reason": "machine builders"} for u in ("Turbines", "Pumps")}
        if "FLEET agent" in system:
            cl = json.loads(json.dumps(CLASSES))
            if not self.oem:
                for c in cl:
                    c["installed_machines"] = {"value": 20 * c["new_machines_per_year"]["value"], "basis": "assumption", "reason": "r"}
                    c["product_life_years"] = {"value": 20, "basis": "assumption", "reason": "r"}
            return {"machine": "one machine", "classes": cl}
        if "FIT agent" in system:
            return {"classes": [{"class": "small", "adoption_share": {"value": 0.1, "basis": "assumption", "reason": "r"},
                                 "units_per_machine": {"value": 1, "basis": "assumption", "reason": "r"}},
                                {"class": "large", "adoption_share": {"value": 0.5, "basis": "assumption", "reason": "r"},
                                 "units_per_machine": {"value": 2, "basis": "assumption", "reason": "r"}}]}
        if "PRICE agent using ONE method" in system:
            v = lambda x: {"value": x, "basis": "assumption", "reason": "r"}  # noqa: E731
            if "own price" in system:  # first answer 900k disagrees 9x with the other two methods -> retried
                small = 100_000 if "REVIEW OF YOUR PREVIOUS ANSWER" in user else 900_000
                return {"classes": [{"class": "small", "price_usd": v(small)}, {"class": "large", "price_usd": v(300_000)}]}
            if "engineering cost" in system:
                return {"classes": [{"class": "small", "rated_kw": v(1000), "usd_per_kw": v(100)},
                                    {"class": "large", "rated_kw": v(3000), "usd_per_kw": v(100)}]}
            return {"classes": [{"class": "small", "machine_price_usd": v(1_000_000), "product_share_of_machine": v(0.1)},
                                {"class": "large", "machine_price_usd": v(5_000_000), "product_share_of_machine": v(0.06)}]}
        raise AssertionError(system[:60])


def _build(tmp_path, oem=True):
    (tmp_path / ub.WB_CACHE).write_text(json.dumps(WB), encoding="utf-8")
    return ua.build(FakeLLM(oem), SPEC, NEST, "Global Widget Gearbox OEM Market", tmp_path,
                    lambda qs: [{"text": "1,000 small machines are built a year", "domain": "example.org"},
                                {"text": "Coherent Market Insights: the market is worth US$ 1 bn", "domain": "x.com"}],
                    {}, say=lambda *_: None, fao_cache={}, workers=4)


def test_first_fit_units_and_price_by_code(tmp_path):
    res = _build(tmp_path)
    r = res["use_cases"][0]
    # (1000 x 0.1 x 1) + (200 x 0.5 x 2) = 100 + 200 = 300 gearboxes; no replacements in an OEM market
    assert math.isclose(r["world_units"], 300.0)
    assert r["spec"]["count_replacements"] is False
    assert "installed_machines" not in r["spec"]["world"]
    # unit-weighted price after the critic's retry: (100 x 100k + 200 x 300k) / 300
    assert math.isclose(r["price_per_quantity_unit"], (100 * 100_000 + 200 * 300_000) / 300)
    assert any("price" in fb for s in r["agents"]["samples"] for fb in s["feedback"])
    assert math.isclose(res["global_value"], 2 * 300 * r["price_per_quantity_unit"] / 1e6)


def test_publisher_blocked_and_driver_spreads(tmp_path):
    res = _build(tmp_path)
    r = res["use_cases"][0]
    assert r["blocked"]["snippets"] == 2 and all("Coherent" not in e["text"] for e in r["evidence"])
    tot = sum(WB[i]["medhigh_tech_manufacturing_usd"] for i in WB)
    assert math.isclose(r["units"]["Germany"], 300 * 4e11 / tot)


def test_replacements_counted_when_scope_says_so(tmp_path):
    res = _build(tmp_path, oem=False)
    r = res["use_cases"][0]
    # flow per class = new + installed / life = 2 x new
    assert math.isclose(r["world_units"], 600.0)
    assert r["spec"]["count_replacements"] is True


def test_ledger_reconciles_first_fit(tmp_path):
    res = _build(tmp_path)
    led = ub.build_ledger(res, SPEC, run_id="a1")
    assert led.mismatches == []
    world = [n for n in led.rows if n.name.endswith("world quantity per year")][0]
    assert world.formula.count("×") == 2 and "/" not in world.formula


def test_price_methods_and_critic():
    names = ["a"]
    row = {"class": "a", "price_usd": 120, "rated_kw": 10, "usd_per_kw": 10, "machine_price_usd": 1000,
           "product_share_of_machine": 0.09}
    assert ua.price_methods(row) == {"direct": 120, "per_kw": 100, "machine_share": 90.0}
    assert math.isclose(ua.class_price(row), 100)
    assert not ua.critique_price({"classes": [row]}, names)
    spread = {**row, "price_usd": 1000}       # 10x away from the other methods
    assert any("disagree" in p for p in ua.critique_price({"classes": [spread]}, names))
    one = {"class": "a", "price_usd": 100}     # a single method is not enough
    assert ua.critique_price({"classes": [one]}, names)


def test_geography_menu_follows_scope(tmp_path):
    table = ub.DriverTable(WB, GEOS, {})

    class Geo:
        def complete_json(self, system, user):
            return {"Turbines": {"driver": "co2_power_mt"}, "Pumps": {"driver": "medhigh_tech_manufacturing_usd"}}

    ucs = [{"use_case": "Turbines"}, {"use_case": "Pumps"}]
    g = ua.geography(Geo(), "M", ucs, {"demand_location": "machine_builder"}, table)
    assert g["Turbines"]["driver"] in ua._OEM_DRIVERS  # an operation proxy is not admissible for builders
    assert g["Pumps"]["driver"] == "medhigh_tech_manufacturing_usd"
    g = ua.geography(Geo(), "M", ucs, {"demand_location": "installation"}, table)
    assert g["Turbines"]["driver"] == "co2_power_mt"


def test_explain_report(tmp_path):
    import importlib.util
    res = _build(tmp_path)
    spec = importlib.util.spec_from_file_location("explain", Path(__file__).resolve().parents[1] / "explain_usecases.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    md = mod.explain(json.loads(json.dumps(res, default=str)))
    assert "Sample 1" in md and "per_kw" in md and "Critic (code) corrections" in md


def test_indentation_sets_hierarchy(tmp_path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment

    from me_engine.tools.extract_schema import _extract_dimensions

    wb = Workbook()
    ws = wb.active
    for r, (lab, ind, v) in enumerate([("By X", 0, None), ("A", 0, 0.03), ("B", 0, 0.33), ("C", 0, 0.53),
                                       ("C1", 1, 0.13), ("C2", 1, 0.87)], start=4):
        ws.cell(r, 2, lab).alignment = Alignment(indent=ind)
        if v is not None:
            ws.cell(r, 3, v)
    dims = _extract_dimensions(ws)
    # numbers alone would read B, C, C1 as A's children (0.33 + 0.53 + 0.13 ~ 1); the indentation says otherwise
    assert {k: v for k, v in dims[0].parent.items() if v} == {"C1": "C", "C2": "C"}


def test_fleet_scope_and_forecast_checks():
    sc = {"count_replacements": False, "end_users": ["Oil & Gas", "Power Generation", "Cement"]}
    hvac = {"classes": [{"class": "small chillers", "serves": ["HVAC"], "new_machines_per_year": {"value": 480000}}]}
    assert any("end-user" in p for p in ua.critique_fleet(hvac, sc))
    ok = {"classes": [{"class": "process", "serves": ["Oil & Gas"], "new_machines_per_year": {"value": 4000, "reason": "E4"}}]}
    assert not ua.critique_fleet(ok, sc)
    fc = {"classes": [{"class": "x", "serves": ["Cement"],
                       "new_machines_per_year": {"value": 117, "reason": "shipments expected to reach 817 units by 2031"}}]}
    assert any("forecast" in p for p in ua.critique_fleet(fc, sc))
    assert ua.end_users({"By Application": {"__root__": ["A"]},
                         "By Industry Vertical": {"__root__": ["Oil & Gas", "Cement"], "Oil & Gas": ["Petrochemical"]}},
                        "By Application") == ["Oil & Gas", "Cement", "Petrochemical"]


def test_fit_speed_and_pattern_checks():
    names = ["s", "m", "l"]
    same_speed = {"classes": [{"class": "l", "machine_speed_rpm": 3000, "other_side_speed_rpm": 3000,
                               "adoption_share": 0.6, "units_per_machine": 1}]}
    assert any("same speed" in p for p in ua.critique_fit(same_speed, ["l"]))
    pattern = {"classes": [{"class": n, "adoption_share": a, "units_per_machine": 1}
                           for n, a in zip(names, (0.05, 0.15, 0.25))]}
    assert any("pattern" in p for p in ua.critique_fit(pattern, names))


def test_median_sample_keeps_inputs_together():
    sc = {"count_replacements": False, "product_unit": "u"}
    def s(k, new, a, price):
        red = {"new_machines_per_year": new, "adoption_share": a, "units_per_machine": 1.0, "units": new * a,
               "price": price, "classes": [{"class": "c"}]}
        return {"sample": k, "reduced": red, "basis": {"price": "assumption"}}
    samples = [s(1, 480000, 0.05, 25000), s(2, 4000, 0.4, 120000), s(3, 3000, 0.5, 100000)]
    spec, info = ua.combine(samples, sc, "d", "m")
    # values: 600M, 192M, 150M -> the median sample is #2, all of its inputs together
    assert info["median_sample"] == 2
    assert spec["world"]["new_machines_per_year"]["value"] == 4000 and spec["world"]["adoption_share"]["value"] == 0.4
    assert spec["price_usd_per_unit"]["value"] == 120000
